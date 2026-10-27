"""API-level tests: authentication, CSRF, RBAC incl. entity scoping, core workflows."""

import io

import pytest


def _assessment(client, code):
    rows = client.get("/api/assessments").json()
    return next(a for a in rows if a["entity"]["code"] == code)


def test_login_rejects_wrong_password(app_):
    from fastapi.testclient import TestClient

    c = TestClient(app_)
    r = c.post("/api/auth/login", json={"email": "isb@demo.local", "password": "wrong"})
    assert r.status_code == 401


def test_csrf_required_for_cookie_auth(login):
    isb = login("isb@demo.local")
    asm = _assessment(isb, "ATL")
    ans = isb.get(f"/api/assessments/{asm['id']}/answers").json()[0]
    r = isb.c.patch(f"/api/answers/{ans['id']}", json={"notes": "x"})  # no CSRF header
    assert r.status_code == 403


def test_answer_update_recalculates_score(login):
    isb = login("isb@demo.local")
    asm = _assessment(isb, "ATC")
    ans = isb.get(f"/api/assessments/{asm['id']}/answers").json()[0]
    r = isb.patch(
        f"/api/answers/{ans['id']}",
        json={"implemented": "Yes", "documented": "Yes", "effectiveness": "Not Tested"},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["score"] == pytest.approx(0.75)
    assert body["wp_required"] is True
    assert body["finding_id"] == f"ATC-{ans['number']}"
    detail = isb.get(f"/api/answers/{ans['id']}").json()
    assert detail["revisions"], "a revision must be recorded"


def test_entity_scoped_contributor(login):
    nl = login("nl.contact@demo.local")
    codes = {a["entity"]["code"] for a in nl.get("/api/assessments").json()}
    assert codes == {"ATN", "ATC"}
    isb = login("isb@demo.local")
    atl = _assessment(isb, "ATL")
    assert nl.get(f"/api/assessments/{atl['id']}").status_code == 403
    pg_answer = isb.get(f"/api/assessments/{atl['id']}/answers").json()[0]
    assert nl.patch(f"/api/answers/{pg_answer['id']}", json={"notes": "hack"}).status_code == 403
    atn = _assessment(nl, "ATN")
    own = nl.get(f"/api/assessments/{atn['id']}/answers").json()[0]
    assert nl.patch(f"/api/answers/{own['id']}", json={"evidence": "Interview"}).status_code == 200
    # search never returns ATL/ATS answers to the NL contact
    hits = nl.get(
        "/api/search", params={"q": "Monitoring and logging", "types": ["answer"], "limit": 50}
    ).json()
    ents = {h["entity_id"] for h in hits}
    allowed = {atn["entity"]["id"], _assessment(nl, "ATC")["entity"]["id"]}
    assert ents <= allowed


def test_viewer_is_read_only(login):
    v = login("viewer@demo.local")
    assert v.get("/api/dashboard/gap").status_code == 200
    assert (
        v.post("/api/action-items", json={"entity_id": 1, "requirement_id": 1, "title": "x"}).status_code
        == 403
    )
    assert (
        v.post(
            "/api/comments", params={"target_type": "requirement", "target_id": 1}, json={"body": "x"}
        ).status_code
        == 403
    )


def test_no_privilege_escalation(login):
    isb = login("isb@demo.local")
    roles = {r["key"]: r["id"] for r in isb.get("/api/roles").json()}
    r = isb.post(
        "/api/users",
        json={
            "email": "esc@demo.local",
            "display_name": "x",
            "password": "Escalate-Pass-2026!",
            "assignments": [{"role_id": roles["admin"], "entity_id": None}],
        },
    )
    assert r.status_code == 403
    admin_id = next(u["id"] for u in isb.get("/api/users").json() if u["email"] == "admin@example.local")
    assert isb.post(f"/api/users/{admin_id}/reset-password").status_code == 403


def test_action_plan_from_findings(login):
    isb = login("isb@demo.local")
    unplanned = isb.get("/api/action-items/unplanned").json()
    assert unplanned
    r = isb.post(
        "/api/action-items/from-findings",
        json={"answer_ids": [unplanned[0]["answer_id"]], "priority": "High"},
    )
    assert r.status_code == 200 and len(r.json()) == 1
    item = r.json()[0]
    assert item["code"].startswith("AP-")
    assert item["priority"] == "High"


def test_document_upload_download_and_restricted(login):
    isb = login("isb@demo.local")
    files = {"file": ("note.txt", io.BytesIO(b"Restricted crisis contact list"), "text/plain")}
    meta = '{"title":"Crisis contacts","doc_type":"Record","classification":"Restricted"}'
    r = isb.post("/api/documents", files=files, data={"metadata": meta})
    assert r.status_code == 200, r.text
    doc = r.json()
    v = doc["current_version"]
    d = isb.get(f"/api/documents/{doc['id']}/versions/{v['id']}/download")
    assert d.status_code == 200 and d.content == b"Restricted crisis contact list"
    assert "attachment" in d.headers["content-disposition"]
    viewer = login("viewer@demo.local")
    assert viewer.get(f"/api/documents/{doc['id']}").status_code == 404
    hits = viewer.get("/api/search", params={"q": "crisis contact list"}).json()
    assert all(not (h["source_type"] == "document" and h["source_id"] == doc["id"]) for h in hits)


def test_upload_rejects_spoofed_type(login):
    isb = login("isb@demo.local")
    files = {"file": ("evil.pdf", io.BytesIO(b"<html><script>alert(1)</script>"), "application/pdf")}
    r = isb.post("/api/documents", files=files, data={"metadata": '{"title":"x"}'})
    assert r.status_code == 400


def test_tasks_and_my_work(login):
    isb = login("isb@demo.local")
    it = login("it@demo.local")
    r = isb.post(
        "/api/tasks",
        json={"title": "Collect firewall config", "assignee_id": it.me["id"], "priority": "High"},
    )
    assert r.status_code == 200
    task = r.json()
    mine = it.get("/api/me/work").json()
    assert any(t["id"] == task["id"] for t in mine["tasks"])
    assert it.patch(f"/api/tasks/{task['id']}", json={"status": "done"}).json()["status"] == "done"
    notes = it.get("/api/notifications").json()
    assert any("Collect firewall config" in n["title"] for n in notes["items"])


def test_regulations_linked(login):
    isb = login("isb@demo.local")
    req = isb.get("/api/requirements").json()
    r321 = next(r for r in req if r["number"] == "3.2.1")
    detail = isb.get(f"/api/requirements/{r321['id']}").json()
    codes = {p["regulation_code"] for p in detail["provisions"]}
    assert {"IR2024/2690", "NIS2", "BSIG", "CBW", "ISO27001"} <= codes


def test_export_workbook(login):
    isb = login("isb@demo.local")
    r = isb.get("/api/export/gap-workbook")
    assert r.status_code == 200
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert {"ATL", "ATS", "ATN", "ATC", "Action Plan", "Dashboard"} <= set(wb.sheetnames)
    ws = wb["ATL"]
    assert ws["J6"].value in ("Yes", "Partial", "No", "N/A")
    check = wb["Lists"]["Q6"]
    assert check.data_type == "f" and str(check.value).startswith("=IFERROR(")


_TEMPLATE_NAMES = (
    "NIS2_GAP_Analysis_Questionnaire_v2.xlsx",
    "NIS2_Implementation_Project_Plan_v2.xlsx",
)


def _template_paths():
    from pathlib import Path

    repo = Path(__file__).resolve().parents[2]
    return [d / n for n in _TEMPLATE_NAMES for d in (repo / "templates", repo / "backend/app/seed/templates")]


def _open(path):
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path)


def test_template_has_no_entity_sheets():
    for path in [p for p in _template_paths() if "GAP" in p.name]:
        assert _open(path).sheetnames == ["Dashboard", "Lists", "Requirements", "Action Plan", "Instructions"]
    for name in _TEMPLATE_NAMES:
        copies = {p.read_bytes() for p in _template_paths() if p.name == name}
        assert len(copies) == 1, f"the two copies of {name} differ"


def test_templates_have_no_author_metadata():
    for path in _template_paths():
        props = _open(path).properties
        assert not props.creator and not props.lastModifiedBy, path


def test_project_template_has_no_entity_names():
    from tests.test_no_group_traces import PATTERNS, compiled

    if not PATTERNS:
        pytest.skip("the pattern file is local and not published")
    banned = compiled()
    for path in [p for p in _template_paths() if "Project_Plan" in p.name]:
        for ws in _open(path):
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        for rx in banned:
                            assert not rx.search(c.value), (ws.title, c.coordinate, rx.pattern)


def test_gap_export_roundtrip(login):
    from app.services.workbook import parse_gap_workbook

    isb = login("isb@demo.local")
    r = isb.get("/api/export/gap-workbook")
    assert r.status_code == 200
    parsed = parse_gap_workbook(io.BytesIO(r.content))
    assert [e["name"] for e in parsed["entities"]] == ["ATL", "ATS", "ATN", "ATC"]
    assert len(parsed["requirements"]) == 174
    for code in ("ATL", "ATS", "ATN", "ATC"):
        rows = {x["number"]: x for x in parsed["answers"][code]}
        api_rows = isb.get(f"/api/assessments/{_assessment(isb, code)['id']}/answers").json()
        assert len(rows) == len(api_rows) == 174
        for a in api_rows:
            got = rows[a["number"]]
            for f in ("scope", "implemented", "documented", "effectiveness", "evidence", "effort"):
                assert got[f] == a[f], (code, a["number"], f)
    r = isb.post(
        "/api/import/gap-workbook",
        params={"dry_run": True},
        files={"file": ("gap.xlsx", io.BytesIO(r.content), "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    assert r.json()["unknown_entities"] == [] and set(r.json()["entities"].values()) == {0}


def test_gap_export_supports_many_entities():
    import openpyxl
    from sqlalchemy.orm import Session

    from app.db import engine
    from app.models import Entity
    from app.services.export import export_gap

    with Session(engine) as db:
        before = db.query(Entity).count()
        for i in range(8):
            db.add(Entity(code=f"X{i}", name=f"Extra {i}", finding_prefix=f"X{i}", sort_order=100 + i))
        db.flush()
        data = export_gap(db, 1, None)
        db.rollback()
    wb = openpyxl.load_workbook(io.BytesIO(data))
    codes = [wb["Lists"][f"F{6 + i}"].value for i in range(before + 8)]
    assert codes[-8:] == [f"X{i}" for i in range(8)] and wb["Lists"][f"F{6 + before + 8}"].value is None
    assert all(c in wb.sheetnames for c in codes) and "Requirements" in wb.sheetnames
    last = wb["X7"]
    assert last["A2"].value == f"=Lists!$F${5 + before + 8}"
    assert last["U6"].value == f'=IF($O6<>"Yes","",Lists!$P${5 + before + 8}&"-"&$A6)'
    assert len(last.data_validations.dataValidation) == 6 and last.freeze_panes
    dash = wb["Dashboard"]
    assert (
        dash["A10"].value == "=Lists!$F$6"
        and dash[f"A{9 + before + 8}"].value == f"=Lists!$F${5 + before + 8}"
    )
