"""Supplier register: RBAC (incl. entity-scoped supplier managers), tier derivation, review cycle, contract
clauses, certificates, questionnaire Excel round-trip, findings -> action items, incident links."""

import io
from datetime import date, timedelta

import openpyxl
import pytest

from app.features.suppliers.logic import add_months, cert_state, due_state, suggest_tier

PW = "Demo-Password-2026!"


def _entities(client) -> dict[str, int]:
    return {e["code"]: e["id"] for e in client.get("/api/entities").json()}


@pytest.fixture(scope="module")
def pgn_manager(login):
    """A user holding the assessor role (incl. suppliers.manage and assets.manage) for ATN only."""
    from sqlalchemy import select

    from app.core.security import hash_password
    from app.db import SessionLocal
    from app.models import Entity, Role, RoleAssignment, User

    db = SessionLocal()
    try:
        if not db.scalar(select(User).where(User.email == "atn.manager@demo.local")):
            u = User(
                email="atn.manager@demo.local", display_name="ATN Manager", password_hash=hash_password(PW)
            )
            db.add(u)
            db.flush()
            role = db.scalar(select(Role).where(Role.key == "assessor"))
            atn = db.scalar(select(Entity).where(Entity.code == "ATN"))
            db.add(RoleAssignment(user_id=u.id, role_id=role.id, entity_id=atn.id))
            db.commit()
    finally:
        db.close()
    return login("atn.manager@demo.local")


def _supplier(client, name):
    return next(s for s in client.get("/api/suppliers").json() if s["name"] == name)


def _body(detail: dict, **over) -> dict:
    keys = [
        "name",
        "legal_name",
        "country",
        "website",
        "description",
        "services_products",
        "category",
        "access_type",
        "tier",
        "tier_overridden",
        "tier_rationale",
        "status",
        "contacts",
        "relationship_owner_id",
        "contract_reference",
        "contract_end",
        "notes",
        "entity_ids",
    ]
    b = {k: detail.get(k) for k in keys}
    b.update(over)
    return b


# --- pure rules ------------------------------------------------------------------------------------------


def test_tier_suggestion_rules():
    assert suggest_tier("it_remote", 4)[0] == "critical"
    assert suggest_tier("ot_remote", 3)[0] == "critical"
    assert suggest_tier("cloud_data", 2)[0] == "high"
    assert suggest_tier("none", 3)[0] == "high"
    assert suggest_tier("none", 2)[0] == "standard"
    assert suggest_tier("none", None)[0] == "standard"


def test_review_and_certificate_states():
    today = date(2026, 9, 29)
    assert due_state(None, False, "active", today) == "never"
    assert due_state(date(2026, 9, 1), True, "active", today) == "overdue"
    assert due_state(date(2026, 10, 20), True, "active", today) == "due_soon"
    assert due_state(date(2027, 10, 20), True, "active", today) == "ok"
    assert due_state(date(2020, 1, 1), True, "offboarded", today) == "n/a"
    assert cert_state(date(2026, 9, 28), today) == "expired"
    assert cert_state(date(2026, 11, 1), today) == "expiring"
    assert cert_state(date(2027, 11, 1), today) == "valid"
    assert add_months(date(2026, 1, 31), 1) == date(2026, 2, 28)


# --- register and RBAC -----------------------------------------------------------------------------------


def test_demo_register_and_derived_tiers(login):
    isb = login("isb@demo.local")
    rows = isb.get("/api/suppliers").json()
    assert len(rows) >= 12
    msp = _supplier(isb, "NordIT Managed Services")
    assert msp["tier"] == "critical" and msp["suggested_tier"] == "critical" and not msp["tier_overridden"]
    assert msp["dependent_services"] >= 1
    printer = _supplier(isb, "PrintFleet Managed Print")
    assert (
        printer["tier_overridden"] and printer["tier"] == "standard" and printer["suggested_tier"] == "high"
    )
    ot = _supplier(isb, "Rhein Automation Service")
    assert ot["due_state"] == "overdue"
    assert ot["certificates"]["expiring"] == 1
    detail = isb.get(f"/api/suppliers/{msp['id']}").json()
    assert [c["clause"] for c in detail["clause_list"]] == list("abcdefgh")
    assert "notify" in detail["clause_list"][3]["text"]
    assert detail["provision"]["url"].startswith("/regulations/")
    assert detail["can_edit"] is True


def test_read_only_roles_cannot_write(login):
    viewer = login("viewer@demo.local")
    nl = login("nl.contact@demo.local")
    s = _supplier(viewer, "ERPWorks Cloud")
    assert viewer.get(f"/api/suppliers/{s['id']}").json()["can_edit"] is False
    for c in (viewer, nl):
        assert c.post("/api/suppliers", json={"name": "X"}).status_code == 403
        assert c.put(f"/api/suppliers/{s['id']}/clauses/a", json={"status": "agreed"}).status_code == 403
        assert (
            c.post(
                f"/api/suppliers/{s['id']}/assessments", json={"assessed_on": "2026-09-01", "method": "audit"}
            ).status_code
            == 403
        )
        assert c.delete(f"/api/suppliers/{s['id']}").status_code == 403
    # dependent services of other entities are hidden
    msp = nl.get(f"/api/suppliers/{_supplier(nl, 'ERPWorks Cloud')['id']}").json()
    ents = _entities(nl)
    assert all(s["entity_id"] in (None, ents["ATN"], ents["ATC"]) for s in msp["services"])
    assert msp["hidden_dependencies"] > 0


def test_entity_scoped_supplier_manager(login, pgn_manager):
    isb = login("isb@demo.local")
    ents = _entities(isb)
    assert (
        pgn_manager.post("/api/suppliers", json={"name": "Scoped A", "entity_ids": [ents["ATL"]]}).status_code
        == 403
    )
    assert pgn_manager.post("/api/suppliers", json={"name": "Scoped B", "entity_ids": []}).status_code == 403
    r = pgn_manager.post(
        "/api/suppliers", json={"name": "Scoped C", "entity_ids": [ents["ATN"]], "access_type": "it_remote"}
    )
    assert r.status_code == 200, r.text
    created = r.json()
    assert created["code"].startswith("SUP-") and created["can_edit"]
    # cannot extend the supplier to an entity outside the scope
    assert (
        pgn_manager.put(
            f"/api/suppliers/{created['id']}", json=_body(created, entity_ids=[ents["ATN"], ents["ATL"]])
        ).status_code
        == 403
    )
    # supplier used only by ATL entities: not even visible
    erp_pg_only = _supplier(isb, "LogiShip Carrier Portal")  # ATL, ATS
    assert (
        pgn_manager.put(
            f"/api/suppliers/{erp_pg_only['id']}/clauses/a", json={"status": "agreed"}
        ).status_code
        == 404
    )
    # shared master data: a supplier also used by ATL / ATS / ATC is read-only for a ATN-only manager
    ms = _supplier(isb, "Microsoft 365")
    assert _supplier(pgn_manager, "Microsoft 365")["can_edit"] is False
    assert (
        pgn_manager.put(f"/api/suppliers/{ms['id']}/clauses/h", json={"status": "agreed"}).status_code == 403
    )
    assert pgn_manager.delete(f"/api/suppliers/{ms['id']}").status_code == 403
    assert (
        pgn_manager.put(f"/api/suppliers/{created['id']}/clauses/h", json={"status": "agreed"}).status_code
        == 200
    )
    assert pgn_manager.delete(f"/api/suppliers/{created['id']}").status_code == 200


def test_create_update_clauses_certificates_assessments(login):
    isb = login("isb@demo.local")
    ents = _entities(isb)
    r = isb.post(
        "/api/suppliers",
        json={
            "name": "Test Remote Care",
            "country": "de",
            "access_type": "none",
            "entity_ids": [ents["ATL"]],
            "contacts": [{"name": "Hotline", "email": "hotline@test.example"}, {}],
        },
    )
    assert r.status_code == 200, r.text
    s = r.json()
    assert s["country"] == "DE" and len(s["contacts"]) == 1
    assert s["tier"] == "standard" and s["due_state"] == "never"
    # tier override needs a tier
    assert (
        isb.put(f"/api/suppliers/{s['id']}", json=_body(s, tier_overridden=True, tier=None)).status_code
        == 400
    )
    s = isb.put(
        f"/api/suppliers/{s['id']}", json=_body(s, tier_overridden=True, tier="high", tier_rationale="Test")
    ).json()
    assert s["tier"] == "high" and s["suggested_tier"] == "standard"
    # clause checklist
    d = isb.put(
        f"/api/suppliers/{s['id']}/clauses/d", json={"status": "agreed", "reference": "MSA § 7"}
    ).json()
    assert (
        d["clauses"]["d"] == "agreed"
        and d["clause_counts"]["agreed"] == 1
        and d["clause_counts"]["missing"] == 7
    )
    assert isb.put(f"/api/suppliers/{s['id']}/clauses/z", json={"status": "agreed"}).status_code == 404
    # certificate expiry state
    c = isb.post(
        f"/api/suppliers/{s['id']}/certificates",
        json={"cert_type": "ISO/IEC 27001", "valid_until": (date.today() + timedelta(days=20)).isoformat()},
    )
    assert c.json()["state"] == "expiring"
    # assessment: next due follows the tier cycle (high = 24 months)
    a = isb.post(
        f"/api/suppliers/{s['id']}/assessments",
        json={
            "assessed_on": "2026-06-01",
            "method": "audit",
            "score": 55,
            "findings": [{"text": "No incident notification clause", "severity": "high"}],
        },
    ).json()
    assert a["next_due"] == "2028-06-01"
    detail = isb.get(f"/api/suppliers/{s['id']}").json()
    assert detail["last_assessment"]["id"] == a["id"] and detail["due_state"] == "ok"
    # history is in the audit log
    audit = isb.get("/api/audit", params={"target_type": "supplier"}).json()
    assert any(e["target_id"] == s["id"] for e in (audit["items"] if isinstance(audit, dict) else audit))
    # finding -> action plan item (IR 5.1.6 by default)
    viewer = login("viewer@demo.local")
    url = f"/api/suppliers/assessments/{a['id']}/findings/0/action-item"
    assert viewer.post(url, json={"entity_id": ents["ATL"]}).status_code == 403
    r = isb.post(url, json={"entity_id": ents["ATL"], "priority": "High"})
    assert r.status_code == 200, r.text
    f = r.json()["findings"][0]
    assert f["action_item_id"] and f["action_item_code"].startswith("AP-")
    item = isb.get(f"/api/action-items/{f['action_item_id']}").json()
    item = item.get("item", item)
    assert item["requirement"]["number"] == "5.1.6"
    assert isb.post(url, json={"entity_id": ents["ATL"]}).status_code == 400
    assert isb.delete(f"/api/suppliers/{s['id']}").status_code == 200
    assert isb.get(f"/api/suppliers/{s['id']}").status_code == 404


def _answer_workbook(data: bytes, answers: list) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Questionnaire"]
    for i, r in enumerate(range(2, ws.max_row + 1)):
        ws.cell(r, 5).value = answers[i % len(answers)]
        ws.cell(r, 6).value = "see policy"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_questionnaire_round_trip(login):
    isb = login("isb@demo.local")
    s = _supplier(isb, "Maas Industrial Controls")  # critical tier, never assessed
    r = isb.get(f"/api/suppliers/{s['id']}/questionnaire.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert {"Instructions", "Questionnaire", "_meta"} <= set(wb.sheetnames)
    n_questions = wb["Questionnaire"].max_row - 1
    filled = _answer_workbook(r.content, ["Yes", "ja", "No", "Partial", "N/A", "maybe"])
    files = {"file": ("answers.xlsx", filled, "application/octet-stream")}
    dry = isb.post(
        f"/api/suppliers/{s['id']}/questionnaire/import", params={"dry_run": True}, files=files
    ).json()
    assert dry["questions"] == n_questions and dry["template"]["tier"] == "critical"
    assert any(row["status"] == "invalid" for row in dry["rows"])
    assert dry["findings"] and 0 < dry["score"] < 100
    assert dry["assessment_id"] is None
    assert _supplier(isb, "Maas Industrial Controls")["last_assessment"] is None
    real = isb.post(
        f"/api/suppliers/{s['id']}/questionnaire/import",
        params={"dry_run": False},
        files={"file": ("answers.xlsx", filled, "application/octet-stream")},
    ).json()
    assert real["assessment_id"]
    after = _supplier(isb, "Maas Industrial Controls")
    assert after["last_assessment"]["method"] == "questionnaire"
    assert after["last_assessment"]["score"] == dry["score"]
    assert after["next_due"] == add_months(date.today(), 12).isoformat()
    # a workbook issued to one supplier cannot be imported for another
    other = _supplier(isb, "ERPWorks Cloud")
    bad = isb.post(
        f"/api/suppliers/{other['id']}/questionnaire/import",
        params={"dry_run": True},
        files={"file": ("answers.xlsx", filled, "application/octet-stream")},
    )
    assert bad.status_code == 400
    assert (
        isb.post(
            f"/api/suppliers/{s['id']}/questionnaire/import",
            params={"dry_run": True},
            files={"file": ("a.txt", b"hello", "text/plain")},
        ).status_code
        == 400
    )
    # the NL contact (read only) may not import
    nl = login("nl.contact@demo.local")
    assert (
        nl.post(
            f"/api/suppliers/{s['id']}/questionnaire/import",
            params={"dry_run": True},
            files={"file": ("answers.xlsx", filled, "application/octet-stream")},
        ).status_code
        == 403
    )


def test_questionnaire_template_versioning(login):
    isb = login("isb@demo.local")
    data = isb.get("/api/suppliers/questionnaires").json()
    std = data["latest"]["standard"]
    questions = [*std["questions"], {"section": "Exit", "text": "Do you provide an exit plan?", "weight": 1}]
    r = isb.post(
        "/api/suppliers/questionnaires",
        json={"tier": "standard", "title": std["title"], "questions": questions, "change_note": "Exit plan"},
    )
    assert r.status_code == 200, r.text
    new = r.json()
    assert new["version"] == std["version"] + 1
    assert new["questions"][-1]["id"] == f"Q{max(int(q['id'][1:]) for q in std['questions']) + 1:02d}"
    assert (
        login("viewer@demo.local")
        .post(
            "/api/suppliers/questionnaires", json={"tier": "standard", "title": "x", "questions": questions}
        )
        .status_code
        == 403
    )


def test_incident_mentions_and_links(login):
    isb = login("isb@demo.local")
    s = _supplier(isb, "Microsoft 365")
    inc = isb.get(f"/api/suppliers/{s['id']}/incidents").json()
    assert any(i["code"] == "INC-001" for i in inc["mentions"])
    iid = next(i["id"] for i in inc["mentions"] if i["code"] == "INC-001")
    assert isb.post(f"/api/suppliers/{s['id']}/incidents", json={"incident_id": iid}).status_code == 200
    inc = isb.get(f"/api/suppliers/{s['id']}/incidents").json()
    assert [i["id"] for i in inc["linked"]] == [iid] and all(i["id"] != iid for i in inc["mentions"])
    # INC-001 belongs to ATL: the NL contact does not see it
    nl = login("nl.contact@demo.local")
    nl_inc = nl.get(f"/api/suppliers/{s['id']}/incidents").json()
    assert not nl_inc["linked"] and not nl_inc["mentions"]
    assert isb.delete(f"/api/suppliers/{s['id']}/incidents/{iid}").status_code == 200


def test_search_comments_and_export(login):
    isb = login("isb@demo.local")
    hits = isb.get("/api/search", params={"q": "NordIT", "types": ["supplier"]}).json()
    assert any(h["url"].startswith("/suppliers/") for h in hits)
    s = _supplier(isb, "Unitel Carrier Services")
    r = isb.post(
        "/api/comments",
        params={"target_type": "supplier", "target_id": s["id"]},
        json={"body": "Contract renewal in Q1"},
    )
    assert r.status_code == 200, r.text
    assert (
        login("viewer@demo.local")
        .get("/api/comments", params={"target_type": "supplier", "target_id": s["id"]})
        .status_code
        == 200
    )
    # formula injection is neutralised in the register export
    ents = _entities(isb)
    evil = isb.post(
        "/api/suppliers", json={"name": '=HYPERLINK("http://x")', "entity_ids": [ents["ATL"]]}
    ).json()
    r = isb.get("/api/suppliers/export.xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Suppliers"]
    cell = next(row[1] for row in ws.iter_rows(min_row=2) if row[0].value == evil["code"])
    assert cell.data_type == "s"
    assert login("viewer@demo.local").get("/api/suppliers/export.xlsx").status_code == 403
    isb.delete(f"/api/suppliers/{evil['id']}")


def test_certificate_and_review_reminders(login):
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.features.suppliers.reminders import run_reminders
    from app.models import Notification, User

    isb = login("isb@demo.local")
    ents = _entities(isb)
    owner = next(s for s in isb.get("/api/stakeholders").json() if s["name"] == "IT Coordinator")
    s = isb.post(
        "/api/suppliers",
        json={"name": "Reminder Test Ltd", "entity_ids": [ents["ATS"]], "relationship_owner_id": owner["id"]},
    ).json()
    isb.post(
        f"/api/suppliers/{s['id']}/certificates",
        json={"cert_type": "TISAX", "valid_until": (date.today() + timedelta(days=5)).isoformat()},
    )
    isb.post(
        f"/api/suppliers/{s['id']}/assessments",
        json={"assessed_on": "2020-01-01", "method": "audit", "next_due": "2021-01-01"},
    )
    db = SessionLocal()
    try:
        assert run_reminders(db) >= 2
        uid = db.scalar(select(User.id).where(User.email == "it@demo.local"))
        titles = [n.title for n in db.scalars(select(Notification).where(Notification.user_id == uid))]
        assert any("Reminder Test Ltd: renew TISAX certificate" in t for t in titles)
        assert any(
            t.startswith("Overdue since 2021-01-01: Supplier review: Reminder Test Ltd") for t in titles
        )
        before = len(titles)
        run_reminders(db)  # deduplicated
        assert len(list(db.scalars(select(Notification).where(Notification.user_id == uid)))) == before
    finally:
        db.close()
    isb.delete(f"/api/suppliers/{s['id']}")


def test_supplier_visibility_follows_using_entities(login):
    """Scoped readers see group suppliers and suppliers used by one of their entities – in list, detail, search,
    incidents and exports."""
    isb = login("isb@demo.local")
    nl = login("nl.contact@demo.local")
    ents = _entities(isb)
    names = {s["name"] for s in nl.get("/api/suppliers").json()}
    assert {"NordIT Managed Services", "ERPWorks Cloud", "Microsoft 365"} <= names
    assert "LogiShip Carrier Portal" not in names and "SafeVault Backup" not in names
    pg_only = _supplier(isb, "LogiShip Carrier Portal")
    assert nl.get(f"/api/suppliers/{pg_only['id']}").status_code == 404
    assert nl.get(f"/api/suppliers/{pg_only['id']}/incidents").status_code == 404
    assert nl.post(
        "/api/comments", params={"target_type": "supplier", "target_id": pg_only["id"]}, json={"body": "x"}
    ).status_code in (403, 404)
    # a group supplier without using entities is visible to every supplier reader
    grp = isb.post("/api/suppliers", json={"name": "Group Auditors Ltd", "entity_ids": []}).json()
    assert nl.get(f"/api/suppliers/{grp['id']}").status_code == 200
    # search: chunks per using entity
    assert not [h for h in nl.get("/api/search", params={"q": "LogiShip", "types": ["supplier"]}).json()]
    assert [h for h in isb.get("/api/search", params={"q": "LogiShip", "types": ["supplier"]}).json()]
    hits = nl.get("/api/search", params={"q": "NordIT", "types": ["supplier"]}).json()
    assert len({h["source_id"] for h in hits}) == len(hits) == 1
    # the auditor (group-wide read + export) exports everything; a ATN-scoped exporter would only get ATN suppliers
    r = login("auditor@demo.local").get("/api/suppliers/export.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Suppliers"]
    assert pg_only["code"] in {row[0].value for row in ws.iter_rows(min_row=2)}
    isb.delete(f"/api/suppliers/{grp['id']}")
    assert ents


def test_finding_action_links_are_server_state(login):
    isb = login("isb@demo.local")
    ents = _entities(isb)
    s = isb.post("/api/suppliers", json={"name": "Finding Test GmbH", "entity_ids": [ents["ATL"]]}).json()
    a = isb.post(
        f"/api/suppliers/{s['id']}/assessments",
        json={
            "assessed_on": "2026-05-01",
            "method": "audit",
            "findings": [
                {"text": "Forged", "severity": "high", "action_item_id": 1, "action_item_code": "AP-001"}
            ],
        },
    ).json()
    assert a["findings"][0]["action_item_id"] is None and a["findings"][0]["action_item_code"] is None
    a = isb.post(
        f"/api/suppliers/assessments/{a['id']}/findings/0/action-item", json={"entity_id": ents["ATL"]}
    ).json()
    real = a["findings"][0]["action_item_id"]
    body = {
        "assessed_on": "2026-05-01",
        "method": "audit",
        "findings": [
            {"text": "Forged", "severity": "high", "action_item_id": 1},
            {"text": "New", "severity": "low", "action_item_id": 2},
        ],
    }
    a = isb.put(f"/api/suppliers/assessments/{a['id']}", json=body).json()
    assert a["findings"][0]["action_item_id"] == real and a["findings"][1]["action_item_id"] is None
    isb.delete(f"/api/suppliers/{s['id']}")


def test_questionnaire_rejects_foreign_template_and_huge_sheets(login):
    isb = login("isb@demo.local")
    s = _supplier(isb, "NordIT Managed Services")  # critical tier
    data = isb.get(f"/api/suppliers/{s['id']}/questionnaire.xlsx").content
    std_id = isb.get("/api/suppliers/questionnaires").json()["latest"]["standard"]["id"]
    wb = openpyxl.load_workbook(io.BytesIO(data))
    meta = wb["_meta"]
    for row in meta.iter_rows():
        if row[0].value == "template_id":
            row[1].value = std_id
    buf = io.BytesIO()
    wb.save(buf)
    r = isb.post(
        f"/api/suppliers/{s['id']}/questionnaire/import",
        params={"dry_run": True},
        files={"file": ("q.xlsx", buf.getvalue(), "application/octet-stream")},
    )
    assert r.status_code == 400 and "tier" in r.json()["detail"]
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Questionnaire"]
    for i in range(600):
        ws.cell(row=100 + i, column=1).value = f"X{i}"
    buf = io.BytesIO()
    wb.save(buf)
    r = isb.post(
        f"/api/suppliers/{s['id']}/questionnaire/import",
        params={"dry_run": True},
        files={"file": ("q.xlsx", buf.getvalue(), "application/octet-stream")},
    )
    assert r.status_code == 400 and "Too many rows" in r.json()["detail"]
