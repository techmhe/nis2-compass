"""Service and asset inventory, BIA, continuity plans and exercises: entity-scoped RBAC, BIA derivations and
approval, restore-test coverage, consistency flags, dependency tree, CSV/Excel import with dry run, export."""

import io
import json
from datetime import date, timedelta

import openpyxl
import pytest

from app.features.inventory.logic import derive, item_values

PW = "Demo-Password-2026!"


def _entities(client) -> dict[str, int]:
    return {e["code"]: e["id"] for e in client.get("/api/entities").json()}


@pytest.fixture(scope="module")
def pgn_manager(login):
    """A user holding the assessor role (incl. assets.manage) for ATN only."""
    from sqlalchemy import select

    from app.core.security import hash_password
    from app.db import SessionLocal
    from app.models import Entity, Role, RoleAssignment, User

    db = SessionLocal()
    try:
        if not db.scalar(select(User).where(User.email == "atn.manager@demo.local")):
            u = User(
                email="atn.manager@demo.local", display_name="ATN Manager", password_hash=hash_password(PW)
            )
            db.add(u)
            db.flush()
            role = db.scalar(select(Role).where(Role.key == "assessor"))
            atn = db.scalar(select(Entity).where(Entity.code == "ATN"))
            db.add(RoleAssignment(user_id=u.id, role_id=role.id, entity_id=atn.id))
            db.commit()
    finally:
        db.close()
    return login("atn.manager@demo.local")


def _service(client, name):
    return next(s for s in client.get("/api/inventory/services").json() if s["name"] == name)


def _asset(client, name):
    return next(a for a in client.get("/api/inventory/assets").json() if a["name"] == name)


def _svc_body(s: dict, **over) -> dict:
    keys = [
        "entity_id",
        "name",
        "description",
        "kind",
        "criticality",
        "nis2_relevant",
        "owner_id",
        "rto_hours",
        "rpo_hours",
        "status",
        "tags",
        "notes",
    ]
    b = {k: s.get(k) for k in keys}
    b.update(over)
    return b


# --- pure BIA rules --------------------------------------------------------------------------------------


def test_bia_derivation():
    imp = {"operational": {"4h": 1, "24h": 3, "72h": 4, "1w": 4}, "financial": {"4h": 0, "24h": 2}}
    d = derive(imp)
    assert d["mtpd_hours"] == 24 and d["rto_hours"] == 12 and d["rpo_hours"] == 12
    assert d["criticality"] == 3
    assert derive({"safety": {"4h": 4}})["mtpd_hours"] == 4
    assert derive({"safety": {"4h": 4}})["criticality"] == 4
    assert derive({})["mtpd_hours"] is None
    assert derive({"reputation": {"1w": 3}})["mtpd_hours"] == 168
    assert derive({"operational": {"4h": "9"}})["windows"][0]["max"] == 4  # clamped

    class Item:
        impacts = imp
        mtpd_hours = None
        rto_hours = 48.0
        rpo_hours = None

    v = item_values(Item())
    assert v["mtpd_hours"] == 24 and v["rto_hours"] == 48 and v["rpo_hours"] == 24
    assert any("longer than the maximum tolerable" in w for w in v["warnings"])


# --- RBAC ------------------------------------------------------------------------------------------------


def test_entity_scoping_for_contributor_and_viewer(login):
    nl = login("nl.contact@demo.local")
    isb = login("isb@demo.local")
    ents = _entities(isb)
    allowed = {None, ents["ATN"], ents["ATC"]}
    services = nl.get("/api/inventory/services").json()
    assets = nl.get("/api/inventory/assets").json()
    assert services and {s["entity_id"] for s in services} <= allowed
    assert {a["entity_id"] for a in assets} <= allowed
    assert all(not s["can_edit"] for s in services)
    pg_service = _service(isb, "ERP and order management")
    assert nl.get(f"/api/inventory/services/{pg_service['id']}").status_code == 403
    assert nl.get(f"/api/inventory/assets/{_asset(isb, 'PLC line 1')['id']}").status_code == 403
    assert {b["entity"]["code"] for b in nl.get("/api/inventory/bias").json()} <= {"ATN", "ATC"}
    assert {e["entity_id"] for e in nl.get("/api/inventory/exercises").json()} <= allowed
    flags = nl.get("/api/inventory/consistency").json()
    assert {f["entity_id"] for f in flags} <= allowed
    tree = nl.get("/api/inventory/dependencies").json()
    assert {s["entity"]["id"] if s["entity"] else None for s in tree["services"]} <= allowed
    for c in (nl, login("viewer@demo.local")):
        assert (
            c.post("/api/inventory/services", json={"name": "x", "entity_id": ents["ATN"]}).status_code == 403
        )
        assert (
            c.post("/api/inventory/assets", json={"name": "x", "entity_id": ents["ATN"]}).status_code == 403
        )
        assert (
            c.post("/api/inventory/bias", json={"entity_id": ents["ATN"], "period": "2027"}).status_code
            == 403
        )
    assert login("viewer@demo.local").get("/api/inventory/overview").status_code == 200


def test_entity_scoped_manager(login, pgn_manager):
    ents = _entities(login("isb@demo.local"))
    assert (
        pgn_manager.post(
            "/api/inventory/services", json={"name": "Scoped", "entity_id": ents["ATL"]}
        ).status_code
        == 403
    )
    assert pgn_manager.post("/api/inventory/services", json={"name": "Scoped group"}).status_code == 403
    r = pgn_manager.post(
        "/api/inventory/services", json={"name": "Scoped ATN", "entity_id": ents["ATN"], "criticality": 2}
    )
    assert r.status_code == 200, r.text
    s = r.json()
    # moving it to another entity is not allowed
    assert (
        pgn_manager.put(
            f"/api/inventory/services/{s['id']}", json=_svc_body(s, entity_id=ents["ATC"])
        ).status_code
        == 403
    )
    assert pgn_manager.delete(f"/api/inventory/services/{s['id']}").status_code == 200


# --- services, assets, dependencies ----------------------------------------------------------------------


def test_service_asset_crud_dependencies_and_history(login):
    isb = login("isb@demo.local")
    ents = _entities(isb)
    ms = next(s for s in isb.get("/api/suppliers").json() if s["name"] == "Microsoft 365")
    s = isb.post(
        "/api/inventory/services",
        json={"name": "Test quality lab", "entity_id": ents["ATS"], "criticality": 3, "tags": ["lab", " "]},
    ).json()
    assert s["code"].startswith("SVC-") and s["tags"] == ["lab"]
    a = isb.post(
        "/api/inventory/assets",
        json={
            "name": "LIMS server",
            "entity_id": ents["ATS"],
            "asset_type": "IT",
            "confidentiality": 3,
            "integrity": 4,
            "availability": 2,
            "supplier_id": ms["id"],
        },
    ).json()
    assert a["supplier"]["name"] == "Microsoft 365" and a["owner"] is None
    assert isb.post("/api/inventory/assets", json={"name": "x", "asset_type": "Mainframe"}).status_code == 422
    assert isb.post("/api/inventory/assets", json={"name": "x", "confidentiality": 5}).status_code == 422
    d = isb.put(
        f"/api/inventory/services/{s['id']}/dependencies",
        json={"asset_ids": [a["id"]], "supplier_ids": [ms["id"]]},
    ).json()
    assert d["asset_ids"] == [a["id"]] and d["suppliers"][0]["id"] == ms["id"]
    s2 = isb.put(f"/api/inventory/services/{s['id']}", json=_svc_body(d, criticality=4, rto_hours=8)).json()
    assert s2["criticality"] == 4
    assert any(h["action"] == "update" and "criticality" in (h["details"] or {}) for h in s2["history"])
    # the tree shows service -> asset -> supplier
    tree = isb.get("/api/inventory/dependencies", params={"entity_id": ents["ATS"]}).json()
    node = next(x for x in tree["services"] if x["id"] == s["id"])
    assert node["assets"][0]["supplier"]["id"] == ms["id"] and node["suppliers"][0]["tier"]
    # the supplier sees the new dependent service
    sup = isb.get(f"/api/suppliers/{ms['id']}").json()
    assert any(x["id"] == s["id"] for x in sup["services"])
    # asset-side dependency editing
    a2 = isb.put(f"/api/inventory/assets/{a['id']}/services", json={"service_ids": []}).json()
    assert a2["service_ids"] == []
    # consistency flags for the new objects: no owner, critical without BIA / plan
    codes = {(f["object_code"], f["code"]) for f in isb.get("/api/inventory/consistency").json()}
    assert (s["code"], "critical_no_bia") in codes and (s["code"], "critical_no_plan") in codes
    assert (a["code"], "asset_no_owner") in codes and (a["code"], "asset_orphan") in codes
    # comments on inventory objects use the generic comment API
    assert (
        isb.post(
            "/api/comments",
            params={"target_type": "asset", "target_id": a["id"]},
            json={"body": "Check backup scope"},
        ).status_code
        == 200
    )
    assert isb.delete(f"/api/inventory/assets/{a['id']}").status_code == 200
    assert isb.delete(f"/api/inventory/services/{s['id']}").status_code == 200
    hits = isb.get("/api/search", params={"q": "Manufacturing execution", "types": ["service"]}).json()
    assert any(h["url"].startswith("/inventory/services") for h in hits)


# --- BIA workflow ----------------------------------------------------------------------------------------


def test_bia_workflow_adopts_targets(login):
    isb = login("isb@demo.local")
    ents = _entities(isb)
    b = isb.post("/api/inventory/bias", json={"entity_id": ents["ATS"], "period": "2027-test"}).json()
    assert b["status"] == "draft" and b["items"], "entity services are pre-filled"
    item = next(i for i in b["items"] if i["label"] == "OT line control plant 2")
    body = {
        "service_id": item["service_id"],
        "impacts": {"operational": {"4h": 4, "24h": 4}, "safety": {"4h": 1}},
        "min_staffing": 3,
    }
    r = isb.put(f"/api/inventory/bia-items/{item['id']}", json=body)
    assert r.status_code == 200, r.text
    it = next(i for i in r.json()["items"] if i["id"] == item["id"])
    assert it["mtpd_hours"] == 4 and it["rto_hours"] == 2 and it["rpo_hours"] == 2
    assert it["dependent_assets"] and it["dependent_suppliers"]
    # process without service record
    r = isb.post(f"/api/inventory/bias/{b['id']}/items", json={"process_name": "Tool shop", "impacts": {}})
    assert r.status_code == 200
    assert isb.post(f"/api/inventory/bias/{b['id']}/items", json={"impacts": {}}).status_code == 400
    # contributors cannot approve
    assert login("nl.contact@demo.local").post(f"/api/inventory/bias/{b['id']}/approve").status_code == 403
    # four-eyes: the creator / last editor cannot approve
    assert isb.post(f"/api/inventory/bias/{b['id']}/approve").status_code == 403
    assert isb.get(f"/api/inventory/bias/{b['id']}").json()["can_approve"] is False
    assessor = login("assessor@demo.local")
    ap = assessor.post(f"/api/inventory/bias/{b['id']}/approve").json()
    assert ap["status"] == "approved" and ap["approved_by"]
    # approved BIAs are locked, including their details
    assert isb.put(f"/api/inventory/bias/{b['id']}", json={"period": "x"}).status_code == 400
    svc = isb.get(f"/api/inventory/services/{item['service_id']}").json()
    assert svc["rto_hours"] == 2 and svc["rpo_hours"] == 2
    assert svc["bia"]["status"] == "approved"
    # approved BIAs are locked
    assert isb.put(f"/api/inventory/bia-items/{item['id']}", json=body).status_code == 400
    assert isb.delete(f"/api/inventory/bias/{b['id']}").status_code == 400
    reopened = isb.post(f"/api/inventory/bias/{b['id']}/reopen").json()
    assert reopened["status"] == "draft" and reopened["approved_by"] is None
    assert isb.delete(f"/api/inventory/bias/{b['id']}").status_code == 200


def test_demo_bias_and_coverage(login):
    isb = login("isb@demo.local")
    bias = {b["entity"]["code"]: b for b in isb.get("/api/inventory/bias").json()}
    assert bias["ATL"]["status"] == "approved" and bias["ATN"]["status"] == "draft"
    ov = isb.get("/api/inventory/overview").json()
    cov = {c["name"]: c for c in ov["coverage"]}
    assert cov["Production control (MES) ATC"]["state"] == "failed"
    assert cov["ERP and order management"]["state"] == "ok"
    assert cov["Identity and access management"]["state"] == "no_test"
    assert ov["counts"]["assets"] >= 30 and ov["flags"]["high"] > 0


# --- plans and exercises ---------------------------------------------------------------------------------


def test_exercises_update_plans_coverage_and_actions(login):
    isb = login("isb@demo.local")
    ents = _entities(isb)
    mes = _service(isb, "Production control (MES) ATC")
    plan = isb.post(
        "/api/inventory/plans",
        json={
            "entity_id": ents["ATC"],
            "title": "DRP ATC",
            "plan_type": "DRP",
            "service_ids": [mes["id"]],
            "test_interval_months": 6,
        },
    ).json()
    assert plan["code"].startswith("PLAN-") and plan["test_state"] == "never"
    day = date.today() - timedelta(days=2)
    ex = {
        "entity_id": ents["ATC"],
        "plan_id": plan["id"],
        "title": "Repeat restore MES",
        "exercise_type": "restore_test",
        "exercise_date": day.isoformat(),
        "result": "passed",
        "achieved_rto_hours": 6,
        "achieved_rpo_hours": 2,
        "service_ids": [mes["id"]],
        "asset_ids": [],
    }
    r = isb.post("/api/inventory/exercises", json=ex)
    assert r.status_code == 200, r.text
    e = r.json()
    p = isb.get(f"/api/inventory/plans/{plan['id']}").json()
    assert p["last_tested"] == day.isoformat() and p["test_state"] == "ok"
    cov = {c["service_id"]: c for c in isb.get("/api/inventory/overview").json()["coverage"]}
    assert cov[mes["id"]]["state"] == "ok"
    # measured restore time above the RTO target
    e = isb.put(f"/api/inventory/exercises/{e['id']}", json={**ex, "achieved_rto_hours": 12}).json()
    cov = {c["service_id"]: c for c in isb.get("/api/inventory/overview").json()["coverage"]}
    assert cov[mes["id"]]["state"] == "exceeds"
    # follow-up actions
    r = isb.post(
        f"/api/inventory/exercises/{e['id']}/actions",
        json={"kind": "action_item", "title": "Shorten MES restore time"},
    )
    assert r.status_code == 200, r.text
    acts = r.json()["actions"]
    assert acts[0]["kind"] == "action_item" and acts[0]["code"].startswith("AP-")
    item = isb.get(f"/api/action-items/{acts[0]['id']}").json()["item"]
    assert item["requirement"]["number"] == "4.2.6" and item["entity"]["code"] == "ATC"
    it_user = next(u for u in isb.get("/api/users/directory").json() if u["email"] == "it@demo.local")
    r = isb.post(
        f"/api/inventory/exercises/{e['id']}/actions",
        json={"kind": "task", "title": "Update runbook", "assignee_id": it_user["id"]},
    )
    assert r.json()["actions"][-1]["code"].startswith("T-")
    assert (
        login("viewer@demo.local")
        .post(f"/api/inventory/exercises/{e['id']}/actions", json={"kind": "task", "title": "x"})
        .status_code
        == 403
    )
    assert isb.delete(f"/api/inventory/exercises/{e['id']}").status_code == 200
    assert isb.delete(f"/api/inventory/plans/{plan['id']}").status_code == 200


# --- import / export -------------------------------------------------------------------------------------


def test_import_dry_run_then_commit_and_export(login):
    isb = login("isb@demo.local")
    csv_text = (
        "Hostname;Asset name;Company;Type;Owner;C;I;A;Supplier;Status;Services\n"
        "IMP-1;Import test PLC;ATS;PLC;OT;low;high;very high;Rhein Automation Service;active;"
        "OT line control plant 2\n"
        "IMP-2;=cmd|' /C calc'!A0;ATL;server;Central IT;2;2;2;;;\n"
        "IMP-3;Broken;XX;blah;Nobody;5;;;;;\n"
    )
    files = {"file": ("assets.csv", csv_text.encode(), "text/csv")}
    prev = isb.post("/api/inventory/import/preview", files=files).json()
    assert prev["row_count"] == 3
    mapping = prev["suggested"]["assets"]
    assert (
        mapping["name"] == "Asset name"
        and mapping["entity"] == "Company"
        and mapping["confidentiality"] == "C"
    )
    form = {"kind": "assets", "mapping": json.dumps(mapping), "dry_run": "true"}
    before = len(isb.get("/api/inventory/assets").json())
    dry = isb.post(
        "/api/inventory/import", data=form, files={"file": ("assets.csv", csv_text.encode(), "text/csv")}
    )
    assert dry.status_code == 200, dry.text
    d = dry.json()
    assert d["created"] == 2 and d["errors"] == 1
    bad = next(r for r in d["rows"] if r["row"] == 4)
    assert bad["action"] == "error" and any("Unknown entity" in m for m in bad["messages"])
    assert len(isb.get("/api/inventory/assets").json()) == before
    # the NL contact has no inventory write rights
    assert (
        login("nl.contact@demo.local")
        .post(
            "/api/inventory/import", data=form, files={"file": ("assets.csv", csv_text.encode(), "text/csv")}
        )
        .status_code
        == 403
    )
    form["dry_run"] = "false"
    real = isb.post(
        "/api/inventory/import", data=form, files={"file": ("assets.csv", csv_text.encode(), "text/csv")}
    ).json()
    assert real["created"] == 2
    plc = _asset(isb, "Import test PLC")
    assert plc["asset_type"] == "OT" and (plc["confidentiality"], plc["integrity"], plc["availability"]) == (
        1,
        3,
        4,
    )
    assert plc["supplier"]["name"] == "Rhein Automation Service" and plc["service_ids"]
    # re-import is idempotent (matched by external ID)
    again = isb.post(
        "/api/inventory/import", data=form, files={"file": ("assets.csv", csv_text.encode(), "text/csv")}
    ).json()
    assert again["created"] == 0 and again["unchanged"] == 2
    # the export neutralises formulas
    r = isb.get("/api/inventory/export.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert {"Services", "Assets", "Dependencies"} <= set(wb.sheetnames)
    cell = next(row[2] for row in wb["Assets"].iter_rows(min_row=2) if row[1].value == "IMP-2")
    assert cell.value.startswith("=") and cell.data_type == "s"
    # an xlsx export can be read back by the importer (column labels map automatically)
    prev = isb.post(
        "/api/inventory/import/preview",
        files={"file": ("inv.xlsx", r.content, "application/octet-stream")},
        data={"sheet": "Services"},
    ).json()
    assert prev["sheet"] == "Services" and prev["suggested"]["services"]["rto_hours"] == "RTO (hours)"
    for name in ("Import test PLC", "=cmd|' /C calc'!A0"):
        isb.delete(f"/api/inventory/assets/{_asset(isb, name)['id']}")


def test_plan_test_reminders(login):
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.features.inventory.reminders import run_reminders
    from app.models import Notification, User

    isb = login("isb@demo.local")
    ents = _entities(isb)
    owner = next(s for s in isb.get("/api/stakeholders").json() if s["name"] == "IT Coordinator")
    p = isb.post(
        "/api/inventory/plans",
        json={
            "entity_id": ents["ATS"],
            "title": "Crisis plan ATS",
            "plan_type": "crisis",
            "owner_id": owner["id"],
            "next_test_due": (date.today() + timedelta(days=3)).isoformat(),
        },
    ).json()
    db = SessionLocal()
    try:
        run_reminders(db)
        uid = db.scalar(select(User.id).where(User.email == "it@demo.local"))
        assert any(
            p["code"] in n.title for n in db.scalars(select(Notification).where(Notification.user_id == uid))
        )
    finally:
        db.close()
    isb.delete(f"/api/inventory/plans/{p['id']}")


def test_exercise_links_need_write_access_to_foreign_objects(login, pgn_manager):
    """M8: a ATN-only manager must not update ATL plans or ATL service coverage through an exercise."""
    isb = login("isb@demo.local")
    ents = _entities(isb)
    pg_plan = next(p for p in isb.get("/api/inventory/plans").json() if p["code"] == "PLAN-001")
    erp = _service(isb, "ERP and order management")
    pgn_svc = _service(isb, "Warehouse management")
    base = {
        "entity_id": ents["ATN"],
        "title": "Scoped test",
        "exercise_type": "restore_test",
        "exercise_date": date.today().isoformat(),
        "result": "passed",
        "achieved_rto_hours": 1,
        "service_ids": [pgn_svc["id"]],
        "asset_ids": [],
    }
    assert pgn_manager.post(
        "/api/inventory/exercises", json={**base, "plan_id": pg_plan["id"]}
    ).status_code in (400, 403)
    assert pgn_manager.post(
        "/api/inventory/exercises", json={**base, "service_ids": [erp["id"]]}
    ).status_code in (400, 403)
    assert isb.get(f"/api/inventory/plans/{pg_plan['id']}").json()["last_tested"] == pg_plan["last_tested"]
    r = pgn_manager.post("/api/inventory/exercises", json=base)
    assert r.status_code == 200, r.text
    # plans: linking a foreign service is not allowed either
    r2 = pgn_manager.post(
        "/api/inventory/plans", json={"entity_id": ents["ATN"], "title": "x", "service_ids": [erp["id"]]}
    )
    assert r2.status_code in (400, 403)
    assert pgn_manager.delete(f"/api/inventory/exercises/{r.json()['id']}").status_code == 200


def test_bia_copy_is_scoped(login, pgn_manager):
    """M9: the copy source must be readable and services of other entities are not copied."""
    isb = login("isb@demo.local")
    ents = _entities(isb)
    pg_bia = next(b for b in isb.get("/api/inventory/bias").json() if b["entity"]["code"] == "ATL")
    assert (
        pgn_manager.post(
            "/api/inventory/bias",
            json={"entity_id": ents["ATN"], "period": "copy", "copy_from_id": pg_bia["id"]},
        ).status_code
        == 403
    )
    b = isb.post(
        "/api/inventory/bias", json={"entity_id": ents["ATN"], "period": "copy", "copy_from_id": pg_bia["id"]}
    ).json()
    pg_ids = {s["id"] for s in isb.get("/api/inventory/services", params={"entity_id": ents["ATL"]}).json()}
    assert b["items"] and not [i for i in b["items"] if i["service_id"] in pg_ids]
    assert any(i["process_name"] for i in b["items"])
    isb.delete(f"/api/inventory/bias/{b['id']}")


def test_import_limits_and_scoped_matching(login, pgn_manager):
    """M10 caps and import matching only within the caller's scope."""
    isb = login("isb@demo.local")
    wide = ";".join(f"c{i}" for i in range(70)) + "\n" + ";".join("x" for _ in range(70)) + "\n"
    r = isb.post("/api/inventory/import/preview", files={"file": ("w.csv", wide.encode(), "text/csv")})
    assert r.status_code == 400 and "columns" in r.json()["detail"]
    long = "Name;Entity\n" + "".join(f"A{i};ATL\n" for i in range(5100))
    r = isb.post("/api/inventory/import/preview", files={"file": ("l.csv", long.encode(), "text/csv")})
    assert r.status_code == 400 and "rows" in r.json()["detail"]
    # a ATL asset name / code is not matched for a ATN manager
    csv_text = "Code;Name;Entity\nAST-008;ERP cloud tenant;ATN\n"
    form = {
        "kind": "assets",
        "mapping": json.dumps({"code": "Code", "name": "Name", "entity": "Entity"}),
        "dry_run": "true",
    }
    d = pgn_manager.post(
        "/api/inventory/import", data=form, files={"file": ("a.csv", csv_text.encode(), "text/csv")}
    ).json()
    assert d["rows"][0]["action"] == "create"


def test_nested_data_of_other_entities_is_hidden(login):
    """L3: nested lists and search text do not reveal other entities' objects."""
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import SearchChunk

    isb = login("isb@demo.local")
    nl = login("nl.contact@demo.local")
    backup = _service(isb, "Backup and recovery")  # group service, tested in the ATL exercise EX-001
    assert any(
        e["code"] == "EX-001" for e in isb.get(f"/api/inventory/services/{backup['id']}").json()["exercises"]
    )
    d = nl.get(f"/api/inventory/services/{backup['id']}").json()
    assert not d["exercises"] and (d["coverage"] or {}).get("last_success") is None
    # a group service linked to a ATL asset: the asset stays out of the NL contact's view and of the search text
    ents = _entities(isb)
    svc = isb.post("/api/inventory/services", json={"name": "Group reporting", "criticality": 2}).json()
    pg_asset = isb.post(
        "/api/inventory/assets", json={"name": "Secret ATL reporting box", "entity_id": ents["ATL"]}
    ).json()
    isb.put(
        f"/api/inventory/services/{svc['id']}/dependencies",
        json={"asset_ids": [pg_asset["id"]], "supplier_ids": []},
    )
    d = nl.get(f"/api/inventory/services/{svc['id']}").json()
    assert d["asset_ids"] == [] and d["assets"] == []
    tree = nl.get("/api/inventory/dependencies").json()
    node = next(x for x in tree["services"] if x["id"] == svc["id"])
    assert node["assets"][0]["hidden"] and not node["assets"][0]["code"] and not node["assets"][0]["name"]
    db = SessionLocal()
    try:
        text = " ".join(
            c.content
            for c in db.scalars(
                select(SearchChunk).where(
                    SearchChunk.source_type == "service", SearchChunk.source_id == svc["id"]
                )
            )
        )
        assert "Secret ATL reporting box" not in text
    finally:
        db.close()
    isb.delete(f"/api/inventory/assets/{pg_asset['id']}")
    isb.delete(f"/api/inventory/services/{svc['id']}")
