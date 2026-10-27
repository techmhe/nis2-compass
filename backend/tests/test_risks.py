"""Risk register: calculations, RBAC incl. entity scoping, gap bridge, acceptance via approval, heat map, export."""

import io
from datetime import date

import openpyxl

from app.features.risks.service import Method, add_months


def _risk(client, code):
    return next(r for r in client.get("/api/risks").json() if r["code"] == code)


def _entity(client, code):
    return next(e for e in client.get("/api/entities").json() if e["code"] == code)


def test_levels_and_appetite():
    m = Method()  # defaults: medium >= 5, high >= 10, very high >= 16, appetite medium
    assert m.level(1, 4) == "low"
    assert m.level(1, 5) == "medium"
    assert m.level(3, 3) == "medium"
    assert m.level(2, 5) == "high"
    assert m.level(4, 4) == "very_high"
    assert m.level(None, 3) is None
    assert not m.above_appetite("medium") and m.above_appetite("high")
    assert m.matrix()[4][4] == "very_high" and m.matrix()[0][0] == "low"
    strict = Method(threshold_medium=4, threshold_high=8, threshold_very_high=15, appetite_level="low")
    assert strict.level(2, 4) == "high" and strict.above_appetite("medium")
    assert add_months(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert add_months(date(2026, 11, 15), 12) == date(2027, 11, 15)


def test_entity_scoping_and_read_only(login):
    isb, nl, viewer = login("isb@demo.local"), login("nl.contact@demo.local"), login("viewer@demo.local")
    all_codes = {r["entity"]["code"] if r["entity"] else None for r in isb.get("/api/risks").json()}
    assert {"ATL", "ATS", "ATN", "ATC", None} <= all_codes
    nl_codes = {r["entity"]["code"] if r["entity"] else None for r in nl.get("/api/risks").json()}
    assert nl_codes <= {"ATN", "ATC", None} and "ATN" in nl_codes
    pg_risk = _risk(isb, "RSK-001")
    assert nl.get(f"/api/risks/{pg_risk['id']}").status_code == 404
    pgn_risk = _risk(nl, "RSK-006")
    detail = nl.get(f"/api/risks/{pgn_risk['id']}").json()
    assert detail["can_edit"] is False  # contributor: risks.read only
    assert nl.patch(f"/api/risks/{pgn_risk['id']}", json={"title": "x"}).status_code == 403
    assert viewer.post("/api/risks", json={"title": "x"}).status_code == 403
    assert viewer.get(f"/api/risks/{pgn_risk['id']}").status_code == 200
    assert viewer.put("/api/risks/methods", json={}).status_code in (403, 422)
    # heat map only counts visible risks
    assert nl.get("/api/risks/heatmap").json()["total"] == len(nl.get("/api/risks").json())


def test_create_and_status_rules(login):
    assessor = login("assessor@demo.local")
    atn = _entity(assessor, "ATN")
    r = assessor.post(
        "/api/risks",
        json={
            "entity_id": atn["id"],
            "title": "USB malware in the packaging line",
            "inherent_likelihood": 3,
            "inherent_impact": 4,
            "residual_likelihood": 2,
            "residual_impact": 4,
            "treatments": ["reduce", "reduce"],
        },
    )
    assert r.status_code == 200, r.text
    risk = r.json()["risk"]
    assert risk["code"].startswith("RSK-") and risk["status"] == "assessed"
    assert (
        risk["inherent_level"] == "high"
        and risk["residual_level"] == "medium"
        and risk["treatments"] == ["reduce"]
    )
    assert risk["next_review_date"]  # default: one review cycle ahead
    rid = risk["id"]
    # accepted only through an approval
    assert assessor.patch(f"/api/risks/{rid}", json={"status": "accepted"}).status_code == 400
    # transfer as the only treatment is flagged
    r = assessor.patch(f"/api/risks/{rid}", json={"treatments": ["transfer"]})
    assert r.json()["risk"]["transfer_only"] is True
    r = assessor.patch(f"/api/risks/{rid}", json={"status": "closed"})
    assert r.json()["risk"]["status"] == "closed"
    hist = assessor.get(f"/api/risks/{rid}/history").json()
    assert [h["action"] for h in hist][-1] == "create" and any(
        h["details"] and "status" in h["details"] for h in hist
    )
    assert assessor.delete(f"/api/risks/{rid}").status_code == 200
    assert assessor.get(f"/api/risks/{rid}").status_code == 404


def test_create_from_answer_links_gap(login):
    isb = login("isb@demo.local")
    asm = next(a for a in isb.get("/api/assessments").json() if a["entity"]["code"] == "ATL")
    rows = isb.get(f"/api/assessments/{asm['id']}/answers").json()
    ans = next(a for a in rows if a["implemented"] in ("No", "Partial"))
    draft = isb.get(f"/api/risks/from-answer/{ans['id']}").json()
    assert draft["title"].startswith("Gap in") and ans["number"] in draft["title"] and draft["can_create"]
    r = isb.post(
        "/api/risks",
        json={
            "entity_id": draft["entity"]["id"],
            "title": draft["title"],
            "scenario": draft["scenario"],
            "source_answer_id": ans["id"],
        },
    )
    assert r.status_code == 200, r.text
    d = r.json()
    assert any(lk["link_type"] == "gap" and lk["target_id"] == ans["id"] for lk in d["links"])
    linked = isb.get("/api/risks/linked", params={"target_type": "answer", "target_id": ans["id"]}).json()
    assert any(x["id"] == d["risk"]["id"] for x in linked)
    # an answer of another entity cannot be linked as gap of an entity risk
    other = next(a for a in isb.get("/api/assessments").json() if a["entity"]["code"] == "ATS")
    other_ans = isb.get(f"/api/assessments/{other['id']}/answers").json()[0]
    r = isb.post(
        f"/api/risks/{d['risk']['id']}/links",
        json={"link_type": "control", "target_type": "answer", "target_id": other_ans["id"]},
    )
    assert r.status_code == 400
    # treatment must be an action item or work package
    r = isb.post(
        f"/api/risks/{d['risk']['id']}/links",
        json={"link_type": "treatment", "target_type": "answer", "target_id": ans["id"]},
    )
    assert r.status_code == 400
    wp = isb.get("/api/work-packages").json()[0]
    r = isb.post(
        f"/api/risks/{d['risk']['id']}/links",
        json={"link_type": "treatment", "target_type": "work_package", "target_id": wp["id"]},
    )
    assert r.status_code == 200 and r.json()["label"].startswith(wp["code"])
    assert isb.get(f"/api/risks/{d['risk']['id']}").json()["risk"]["status"] == "in_treatment"
    # comments on risks are supported (registered target)
    r = isb.post(
        "/api/comments",
        json={"body": "Discussed in the ISMS meeting"},
        params={"target_type": "risk", "target_id": d["risk"]["id"]},
    )
    assert r.status_code == 200


def _accept_ready(client, entity_id, title, residual=(3, 4)):
    rid = client.post(
        "/api/risks",
        json={
            "entity_id": entity_id,
            "title": title,
            "inherent_likelihood": 3,
            "inherent_impact": 4,
            "residual_likelihood": residual[0],
            "residual_impact": residual[1],
            "treatments": ["reduce"],
        },
    ).json()["risk"]["id"]
    return rid


def _make_acceptable(client, rid):
    r = client.patch(
        f"/api/risks/{rid}",
        json={
            "treatments": ["accept"],
            "acceptance_category": "Treatment cost disproportionate to the risk",
            "treatment_justification": "Replacement planned with the 2028 budget",
        },
    )
    assert r.status_code == 200, r.text


def _accept(requester, approver, rid):
    a = requester.post(
        "/api/approvals",
        json={"target_type": "risk_acceptance", "target_id": rid, "approver_user_id": approver.me["id"]},
    )
    assert a.status_code == 200, a.text
    r = approver.post(f"/api/approvals/{a.json()['id']}/decide", json={"decision": "approved", "note": "ok"})
    assert r.status_code == 200, r.text
    return a.json()


def test_acceptance_via_approval(login):
    isb, assessor = login("isb@demo.local"), login("assessor@demo.local")
    atl = _entity(isb, "ATL")
    rid = _accept_ready(assessor, atl["id"], "Legacy CNC controller without vendor support")
    req = {"target_type": "risk_acceptance", "target_id": rid, "approver_user_id": isb.me["id"]}
    r = assessor.post("/api/approvals", json=req)
    assert r.status_code == 400 and "Accept" in r.json()["detail"]
    _make_acceptable(assessor, rid)
    assert (
        assessor.post("/api/approvals", json={**req, "approver_user_id": assessor.me["id"]}).status_code
        == 400
    )
    r = assessor.post("/api/approvals", json=req)
    assert r.status_code == 200, r.text
    a = r.json()
    assert a["meta"]["above_appetite"] and a["meta"]["requires_management_body"]  # 3 x 4 = high > medium
    assert _risk(isb, _code(isb, rid))["acceptance_state"] == "pending"
    r = isb.post(
        f"/api/approvals/{a['id']}/decide", json={"decision": "approved", "note": "Board decision 2026-09"}
    )
    assert r.status_code == 200
    risk = isb.get(f"/api/risks/{rid}").json()["risk"]
    assert risk["status"] == "accepted" and risk["accepted_by"] == isb.me["display_name"]
    assert risk["acceptance_state"] == "approved" and risk["next_review_date"] >= date.today().isoformat()
    assert isb.delete(f"/api/risks/{rid}").status_code == 400  # part of the acceptance record
    # changing the residual rating lets the acceptance lapse
    risk = isb.patch(f"/api/risks/{rid}", json={"residual_impact": 5}).json()["risk"]
    assert risk["status"] == "assessed" and risk["acceptance_state"] == "changed_after_approval"
    st = isb.get("/api/approvals/target", params={"target_type": "risk_acceptance", "target_id": rid}).json()
    assert st["state"] == "changed_after_approval" and st["can_request"]


def test_acceptance_approver_rules(login):
    """H2: read-only users cannot accept risks; above appetite needs the management body (or governance rights)."""
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.features.governance.models import BoardMember

    isb, assessor, viewer, pm = (
        login(e) for e in ("isb@demo.local", "assessor@demo.local", "viewer@demo.local", "pm@demo.local")
    )
    atn = _entity(isb, "ATN")
    above = _accept_ready(isb, atn["id"], "Single point of failure in the MES database")
    _make_acceptable(isb, above)
    body = {"target_type": "risk_acceptance", "target_id": above}
    for u in (viewer, pm, assessor):  # assessor has risks.manage but no governance.manage
        r = isb.post("/api/approvals", json={**body, "approver_user_id": u.me["id"]})
        assert r.status_code == 400 and "management body" in r.json()["detail"], r.text
    ids = {u["id"] for u in isb.get("/api/approvals/approvers", params=body).json()}
    assert not ids & {viewer.me["id"], pm.me["id"], assessor.me["id"]}
    # a management body member of ATN (linked user) may decide, even with read-only rights
    db = SessionLocal()
    try:
        bm = db.scalar(select(BoardMember).where(BoardMember.entity_id == atn["id"]))
        old_user = bm.user_id
        bm.user_id = viewer.me["id"]
        db.commit()
        ids = {u["id"] for u in isb.get("/api/approvals/approvers", params=body).json()}
        assert viewer.me["id"] in ids
        _accept(isb, viewer, above)
        assert isb.get(f"/api/risks/{above}").json()["risk"]["status"] == "accepted"
    finally:
        bm = db.get(BoardMember, bm.id)
        bm.user_id = old_user
        db.commit()
        db.close()
    # within appetite: risks.manage suffices, read-only still not
    within = _accept_ready(isb, atn["id"], "Outdated label printer firmware", residual=(2, 2))
    _make_acceptable(isb, within)
    r = isb.post(
        "/api/approvals",
        json={"target_type": "risk_acceptance", "target_id": within, "approver_user_id": viewer.me["id"]},
    )
    assert r.status_code == 400
    a = _accept(isb, assessor, within)
    assert not a["meta"]["requires_management_body"]


def test_acceptance_lapses_on_justification_entity_close_and_method(login):
    """M2: the acceptance covers entity, justification, method and the accepted state."""
    isb, assessor = login("isb@demo.local"), login("assessor@demo.local")
    atl, ats = _entity(isb, "ATL"), _entity(isb, "ATS")

    def accepted():
        rid = _accept_ready(assessor, atl["id"], "Accepted risk", residual=(1, 2))
        _make_acceptable(assessor, rid)
        _accept(assessor, isb, rid)
        assert isb.get(f"/api/risks/{rid}").json()["risk"]["status"] == "accepted"
        return rid

    rid = accepted()
    r = isb.patch(f"/api/risks/{rid}", json={"treatment_justification": "Something else"}).json()["risk"]
    assert r["status"] == "assessed" and r["acceptance_state"] == "changed_after_approval"
    rid = accepted()
    r = isb.patch(f"/api/risks/{rid}", json={"entity_id": ats["id"]}).json()["risk"]
    assert r["status"] == "assessed" and r["acceptance_state"] == "changed_after_approval"
    rid = accepted()
    r = isb.patch(f"/api/risks/{rid}", json={"status": "closed"}).json()["risk"]
    assert r["acceptance_state"] == "changed_after_approval"
    r = isb.patch(f"/api/risks/{rid}", json={"status": "assessed"}).json()[
        "risk"
    ]  # reopening does not revive it
    assert r["acceptance_state"] == "changed_after_approval"
    isb.patch(f"/api/risks/{rid}", json={"status": "closed"})
    r = assessor.post(
        "/api/approvals",
        json={"target_type": "risk_acceptance", "target_id": rid, "approver_user_id": isb.me["id"]},
    )
    assert r.status_code == 400 and "Closed" in r.json()["detail"]
    # a pending acceptance cannot be approved after the risk was closed
    rid = _accept_ready(assessor, atl["id"], "Pending then closed", residual=(1, 2))
    _make_acceptable(assessor, rid)
    a = assessor.post(
        "/api/approvals",
        json={"target_type": "risk_acceptance", "target_id": rid, "approver_user_id": isb.me["id"]},
    ).json()
    isb.patch(f"/api/risks/{rid}", json={"status": "closed"})
    assert isb.post(f"/api/approvals/{a['id']}/decide", json={"decision": "approved"}).status_code == 409
    # changing the method (appetite / levels) of the entity ends acceptances made under the old method
    rid = accepted()
    base = isb.get("/api/risks/methods/effective", params={"entity_id": atl["id"]}).json()
    body = {
        k: base[k]
        for k in (
            "name",
            "likelihood_scale",
            "impact_scale",
            "threshold_medium",
            "threshold_high",
            "threshold_very_high",
            "appetite_level",
            "acceptance_criteria",
            "review_cycle_months",
        )
    }
    assert (
        isb.put(
            "/api/risks/methods", json={**body, "appetite_level": "low"}, params={"entity_id": atl["id"]}
        ).status_code
        == 200
    )
    try:
        r = isb.get(f"/api/risks/{rid}").json()["risk"]
        assert r["status"] == "assessed" and r["acceptance_state"] == "changed_after_approval"
    finally:
        isb.delete("/api/risks/methods", params={"entity_id": atl["id"]})


def test_owner_must_see_the_risk_and_heatmap_scope(login):
    """L1 / L2: owners must be able to read the risk; the heat map of a foreign entity is not exposed."""
    isb, nl = login("isb@demo.local"), login("nl.contact@demo.local")
    atl, atn = _entity(isb, "ATL"), _entity(isb, "ATN")
    r = isb.post("/api/risks", json={"entity_id": atl["id"], "title": "x", "owner_user_id": nl.me["id"]})
    assert r.status_code == 400 and "cannot see" in r.json()["detail"]
    rid = isb.post(
        "/api/risks", json={"entity_id": atn["id"], "title": "Owned by NL", "owner_user_id": nl.me["id"]}
    ).json()["risk"]["id"]
    assert isb.patch(f"/api/risks/{rid}", json={"entity_id": atl["id"]}).status_code == 400
    assert (
        isb.patch(
            f"/api/risks/{rid}", json={"entity_id": atl["id"], "owner_user_id": isb.me["id"]}
        ).status_code
        == 200
    )
    assert nl.get("/api/risks/heatmap", params={"entity_id": atl["id"]}).status_code == 403
    assert nl.get("/api/risks/heatmap", params={"entity_id": atn["id"]}).status_code == 200


def test_reminders_skip_owners_without_access(login):
    from app.db import SessionLocal
    from app.features.risks.models import Risk
    from app.features.risks.service import run_risk_reminders
    from app.models import Notification

    isb, nl = login("isb@demo.local"), login("nl.contact@demo.local")
    atl = _entity(isb, "ATL")
    rid = isb.post(
        "/api/risks", json={"entity_id": atl["id"], "title": "Stale owner", "next_review_date": "2020-01-01"}
    ).json()["risk"]["id"]
    db = SessionLocal()
    try:
        db.get(Risk, rid).owner_user_id = nl.me["id"]  # e.g. ownership from before a role change
        db.commit()
        run_risk_reminders(db)
        assert (
            db.query(Notification)
            .filter(Notification.user_id == nl.me["id"], Notification.link == f"/risks/{rid}")
            .count()
            == 0
        )
    finally:
        db.close()


def _code(client, rid):
    return client.get(f"/api/risks/{rid}").json()["risk"]["code"]


def test_register_filters_and_heatmap(login):
    isb = login("isb@demo.local")
    rows = isb.get("/api/risks").json()
    hm = isb.get("/api/risks/heatmap").json()
    assert hm["total"] == len(rows)
    rated = sum(1 for r in rows if r["residual_level"])
    assert sum(sum(row) for row in hm["cells"]["residual"]) == rated
    assert sum(hm["by_level"]["residual"].values()) == rated
    r0 = next(r for r in rows if r["residual_likelihood"])
    cell = isb.get(
        "/api/risks",
        params={
            "basis": "residual",
            "likelihood": r0["residual_likelihood"],
            "impact": r0["residual_impact"],
        },
    ).json()
    assert len(cell) == hm["cells"]["residual"][r0["residual_likelihood"] - 1][r0["residual_impact"] - 1]
    atn = _entity(isb, "ATN")
    assert all(
        r["entity"]["id"] == atn["id"] for r in isb.get("/api/risks", params={"entity_id": atn["id"]}).json()
    )
    assert all(r["entity"] is None for r in isb.get("/api/risks", params={"group": True}).json())
    overdue = isb.get("/api/risks", params={"review_overdue": True}).json()
    assert overdue and all(r["review_overdue"] for r in overdue)
    above = isb.get("/api/risks", params={"above_appetite": True}).json()
    assert all(r["above_appetite"] for r in above)
    # ATC uses its own, stricter method
    atc = _entity(isb, "ATC")
    m = isb.get("/api/risks/methods/effective", params={"entity_id": atc["id"]}).json()
    assert m["appetite_level"] == "low" and not m["inherited"]


def test_method_override_permissions(login):
    assessor, nl = login("assessor@demo.local"), login("nl.contact@demo.local")
    ats = _entity(assessor, "ATS")
    base = assessor.get("/api/risks/methods/effective", params={"entity_id": ats["id"]}).json()
    assert base["inherited"]
    body = {
        k: base[k]
        for k in (
            "name",
            "likelihood_scale",
            "impact_scale",
            "threshold_medium",
            "threshold_high",
            "threshold_very_high",
            "appetite_level",
            "acceptance_criteria",
            "review_cycle_months",
        )
    }
    assert nl.put("/api/risks/methods", json=body, params={"entity_id": ats["id"]}).status_code == 403
    bad = {**body, "threshold_high": 4}
    assert assessor.put("/api/risks/methods", json=bad, params={"entity_id": ats["id"]}).status_code == 422
    r = assessor.put(
        "/api/risks/methods",
        json={**body, "name": "ATS method", "review_cycle_months": 6},
        params={"entity_id": ats["id"]},
    )
    assert r.status_code == 200 and r.json()["inherited"] is False
    assert assessor.delete("/api/risks/methods", params={"entity_id": ats["id"]}).status_code == 200
    assert assessor.get("/api/risks/methods/effective", params={"entity_id": ats["id"]}).json()["inherited"]


def test_export_neutralises_formulas(login):
    isb = login("isb@demo.local")
    isb.post(
        "/api/risks",
        json={"title": '=HYPERLINK("http://evil","click")', "residual_likelihood": 1, "residual_impact": 1},
    )
    r = isb.get("/api/risks/export")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Risk register", "Treatment plan", "Method"]
    ws = wb["Risk register"]
    cells = [c for c in ws["C"] if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")]
    assert cells and all(c.data_type == "s" for c in cells)
    assert login("viewer@demo.local").get("/api/risks/export").status_code == 403  # no reports.export


def test_threats_and_link_options(login):
    isb, nl = login("isb@demo.local"), login("nl.contact@demo.local")
    threats = isb.get("/api/risks/threats").json()
    assert (
        len(threats) >= 30
        and any(t["code"] == "G 0.39" for t in threats)
        and any(t["code"].startswith("OT-") for t in threats)
    )
    assert (
        nl.post("/api/risks/threats", json={"code": "X-1", "title": "x", "category": "c"}).status_code == 403
    )
    opts = nl.get("/api/risks/link-options", params={"target_type": "action_item"}).json()
    assert all(o["entity_code"] in ("ATN", "ATC", None) for o in opts)


def test_review_and_reminders(login):
    from app.db import SessionLocal
    from app.features.risks.service import run_risk_reminders
    from app.models import Notification

    isb, nl = login("isb@demo.local"), login("nl.contact@demo.local")
    overdue = next(r for r in nl.get("/api/risks").json() if r["review_overdue"])
    db = SessionLocal()
    try:
        assert run_risk_reminders(db) >= 1
        assert (
            db.query(Notification)
            .filter(Notification.user_id == nl.me["id"], Notification.link == f"/risks/{overdue['id']}")
            .count()
            >= 1
        )
    finally:
        db.close()
    r = isb.post(f"/api/risks/{overdue['id']}/review", json={"note": "Reviewed in the ISMS meeting"})
    assert r.status_code == 200
    risk = r.json()["risk"]
    assert not risk["review_overdue"] and risk["last_reviewed_at"] == date.today().isoformat()
