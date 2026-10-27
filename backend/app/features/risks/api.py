"""Risk register and treatment plan (§ 30(2) no. 1 BSIG, Cbb Art. 7, IR 2024/2690 2.1).

Risk = likelihood x impact on 5-point scales of the effective method (entity override or group default).
Acceptance of a residual risk runs through the approval workflow (target type "risk_acceptance"); acceptance
above the risk appetite is flagged "requires management body decision". "Transfer" as the only treatment
triggers a warning: the BSI does not allow blanket risk transfer or acceptance for services in scope."""

import io
import re
from datetime import date, datetime
from typing import Literal

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.api.util import (
    TARGETS,
    can_access_target,
    check_target,
    cleanup_target,
    doc_links_for,
    get_or_404,
    next_code,
)
from app.core.deps import Principal, build_principal, get_principal, require
from app.db import get_db, utcnow
from app.features.approvals import service as approvals
from app.features.approvals.models import ApprovalRequest
from app.features.risks import hooks  # noqa: F401  (registers targets, search source and approval adapter)
from app.features.risks.models import (
    ACCEPTANCE_CATEGORIES,
    LEVELS,
    LINK_TYPES,
    RISK_CATEGORIES,
    RISK_STATUS,
    TREATMENTS,
    Risk,
    RiskLink,
    RiskMethod,
    Threat,
)
from app.features.risks.service import (
    Method,
    MethodCache,
    acceptance_version,
    add_months,
    effective_method,
    transfer_only,
)
from app.models import (
    ActionItem,
    Answer,
    Assessment,
    AssessmentRound,
    AuditEvent,
    Entity,
    Requirement,
    Stakeholder,
    User,
    WorkPackage,
)
from app.schemas.common import EntityBrief, Ok, StakeholderBrief, UserBrief
from app.services.audit import apply_changes, audit
from app.services.indexing import index_object, remove_object
from app.services.scoring import due_status

router = APIRouter(prefix="/risks", tags=["risks"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
Treatment = Literal["reduce", "avoid", "accept", "transfer"]
Status = Literal["identified", "assessed", "in_treatment", "accepted", "closed"]
Level = Literal["low", "medium", "high", "very_high"]


# --- schemas ---------------------------------------------------------------------------------------------------


class ScaleStep(BaseModel):
    label: str = Field(min_length=1, max_length=60)
    description: str | None = Field(default=None, max_length=500)


class MethodOut(BaseModel):
    id: int | None
    entity: EntityBrief | None
    inherited: bool
    name: str
    likelihood_scale: list[ScaleStep]
    impact_scale: list[ScaleStep]
    threshold_medium: int
    threshold_high: int
    threshold_very_high: int
    appetite_level: str
    acceptance_criteria: str | None
    review_cycle_months: int
    matrix: list[list[str]]


def method_out(db: Session, m: Method, entity_id: int | None = None) -> MethodOut:
    eid = entity_id if entity_id is not None else m.entity_id
    ent = db.get(Entity, eid) if eid else None
    return MethodOut(
        id=m.id,
        entity=EntityBrief.model_validate(ent) if ent else None,
        inherited=m.inherited,
        name=m.name,
        likelihood_scale=m.likelihood_scale,
        impact_scale=m.impact_scale,
        threshold_medium=m.threshold_medium,
        threshold_high=m.threshold_high,
        threshold_very_high=m.threshold_very_high,
        appetite_level=m.appetite_level,
        acceptance_criteria=m.acceptance_criteria,
        review_cycle_months=m.review_cycle_months,
        matrix=m.matrix(),
    )


class ThreatOut(BaseModel):
    id: int
    code: str
    title: str
    category: str
    source: str | None
    is_active: bool


class RiskOut(BaseModel):
    id: int
    code: str
    entity: EntityBrief | None
    title: str
    scenario: str | None
    threat: ThreatOut | None
    category: str | None
    owner_stakeholder: StakeholderBrief | None
    owner_user: UserBrief | None
    inherent_likelihood: int | None
    inherent_impact: int | None
    inherent_score: int | None
    inherent_level: str | None
    existing_controls: str | None
    residual_likelihood: int | None
    residual_impact: int | None
    residual_score: int | None
    residual_level: str | None
    above_appetite: bool
    appetite_level: str
    treatments: list[str]
    transfer_only: bool
    treatment_justification: str | None
    acceptance_category: str | None
    status: str
    next_review_date: date | None
    last_reviewed_at: date | None
    review_overdue: bool
    accepted_at: datetime | None
    accepted_by: str | None
    acceptance_state: str  # none | pending | approved | rejected | changes_requested | changed_after_approval
    link_counts: dict[str, int]
    created_at: datetime
    updated_at: datetime


def _score(a: int | None, b: int | None) -> int | None:
    return a * b if a and b else None


def risk_out(
    db: Session,
    r: Risk,
    methods: MethodCache,
    acc: dict[int, ApprovalRequest] | None = None,
    counts: dict[int, dict[str, int]] | None = None,
) -> RiskOut:
    m = methods.get(r.entity_id)
    inh = m.level(r.inherent_likelihood, r.inherent_impact)
    res = m.level(r.residual_likelihood, r.residual_impact)
    if acc is None:
        acc = _acceptance_map(db, [r.id])
    if counts is None:
        counts = _link_counts(db, [r.id])
    req = acc.get(r.id)
    state = "none"
    if req:
        state = req.status
        if req.status == "approved" and req.target_version != acceptance_version(r, m)[0]:
            state = "changed_after_approval"
    return RiskOut(
        id=r.id,
        code=r.code,
        entity=EntityBrief.model_validate(r.entity) if r.entity else None,
        title=r.title,
        scenario=r.scenario,
        threat=ThreatOut.model_validate(r.threat, from_attributes=True) if r.threat else None,
        category=r.category,
        owner_stakeholder=StakeholderBrief.model_validate(r.owner_stakeholder)
        if r.owner_stakeholder
        else None,
        owner_user=UserBrief.model_validate(r.owner_user) if r.owner_user else None,
        inherent_likelihood=r.inherent_likelihood,
        inherent_impact=r.inherent_impact,
        inherent_score=_score(r.inherent_likelihood, r.inherent_impact),
        inherent_level=inh,
        existing_controls=r.existing_controls,
        residual_likelihood=r.residual_likelihood,
        residual_impact=r.residual_impact,
        residual_score=_score(r.residual_likelihood, r.residual_impact),
        residual_level=res,
        above_appetite=m.above_appetite(res),
        appetite_level=m.appetite_level,
        treatments=r.treatments or [],
        transfer_only=transfer_only(r),
        treatment_justification=r.treatment_justification,
        acceptance_category=r.acceptance_category,
        status=r.status,
        next_review_date=r.next_review_date,
        last_reviewed_at=r.last_reviewed_at,
        review_overdue=r.status != "closed"
        and r.next_review_date is not None
        and r.next_review_date < date.today(),
        accepted_at=r.accepted_at,
        accepted_by=r.accepted_by,
        acceptance_state=state,
        link_counts=counts.get(r.id, {}),
        created_at=r.created_at,
        updated_at=r.updated_at,
    )


def _acceptance_map(db: Session, ids: list[int]) -> dict[int, ApprovalRequest]:
    """Latest non-withdrawn acceptance request per risk."""
    out: dict[int, ApprovalRequest] = {}
    if not ids:
        return out
    rows = db.scalars(
        select(ApprovalRequest)
        .where(
            ApprovalRequest.target_type == "risk_acceptance",
            ApprovalRequest.target_id.in_(ids),
            ApprovalRequest.status != "withdrawn",
        )
        .order_by(ApprovalRequest.created_at, ApprovalRequest.id)
    )
    for r in rows:
        out[r.target_id] = r
    return out


def _link_counts(db: Session, ids: list[int]) -> dict[int, dict[str, int]]:
    out: dict[int, dict[str, int]] = {}
    if not ids:
        return out
    for rid, lt in db.execute(select(RiskLink.risk_id, RiskLink.link_type).where(RiskLink.risk_id.in_(ids))):
        d = out.setdefault(rid, {})
        d[lt] = d.get(lt, 0) + 1
    return out


def _visible_risks(db: Session, pr: Principal) -> list[Risk]:
    return [r for r in db.scalars(select(Risk).order_by(Risk.code)) if pr.can_read("risks.read", r.entity_id)]


# --- reference data ------------------------------------------------------------------------------------------


@router.get("/meta")
def meta(pr: Principal = Depends(require("risks.read"))):
    return {
        "treatments": TREATMENTS,
        "statuses": RISK_STATUS,
        "levels": LEVELS,
        "categories": RISK_CATEGORIES,
        "acceptance_categories": ACCEPTANCE_CATEGORIES,
        "link_types": LINK_TYPES,
        "link_targets": sorted(t for t in TARGETS if t != "risk"),
    }


@router.get("/methods", response_model=list[MethodOut])
def list_methods(pr: Principal = Depends(require("risks.read")), db: Session = Depends(get_db)):
    """Group default first, then the effective method of every entity the caller can see."""
    mc = MethodCache(db)
    res = [method_out(db, mc.get(None))]
    for e in db.scalars(
        select(Entity).where(Entity.is_active.is_(True)).order_by(Entity.sort_order, Entity.code)
    ):
        if pr.can_read("risks.read", e.id):
            res.append(method_out(db, mc.get(e.id), e.id))
    return res


@router.get("/methods/effective", response_model=MethodOut)
def get_effective_method(
    entity_id: int | None = None,
    pr: Principal = Depends(require("risks.read")),
    db: Session = Depends(get_db),
):
    pr.ensure("risks.read", entity_id)
    return method_out(db, effective_method(db, entity_id), entity_id)


class MethodIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    likelihood_scale: list[ScaleStep] = Field(min_length=5, max_length=5)
    impact_scale: list[ScaleStep] = Field(min_length=5, max_length=5)
    threshold_medium: int = Field(ge=2, le=25)
    threshold_high: int = Field(ge=3, le=25)
    threshold_very_high: int = Field(ge=4, le=25)
    appetite_level: Level = "medium"
    acceptance_criteria: str | None = Field(default=None, max_length=10000)
    review_cycle_months: int = Field(default=12, ge=1, le=36)

    def model_post_init(self, _ctx) -> None:
        if not (self.threshold_medium < self.threshold_high < self.threshold_very_high):
            raise ValueError("thresholds must increase: medium < high < very high")


@router.put("/methods", response_model=MethodOut)
def save_method(
    body: MethodIn,
    entity_id: int | None = None,
    pr: Principal = Depends(get_principal),
    db: Session = Depends(get_db),
):
    """Create or update the group default (no entity_id) or an entity override."""
    pr.ensure("risks.manage", entity_id, write=True)
    if entity_id is not None:
        get_or_404(db, Entity, entity_id, "Entity")
    m = db.scalar(
        select(RiskMethod).where(
            RiskMethod.entity_id == entity_id if entity_id is not None else RiskMethod.entity_id.is_(None)
        )
    )
    data = body.model_dump()
    data["likelihood_scale"] = [s.model_dump() for s in body.likelihood_scale]
    data["impact_scale"] = [s.model_dump() for s in body.impact_scale]
    if m is None:
        m = RiskMethod(entity_id=entity_id, updated_by_id=pr.id, **data)
        db.add(m)
        db.flush()
        audit(db, pr, "create", "risk_method", m.id, m.name, entity_id=entity_id)
    else:
        diff = apply_changes(m, data)
        m.updated_by_id = pr.id
        audit(db, pr, "update", "risk_method", m.id, m.name, entity_id=entity_id, details=diff)
    db.flush()
    _lapse_stale_acceptances(db, pr)
    db.commit()
    return method_out(db, Method.of(m), entity_id)


def _lapse_stale_acceptances(db: Session, pr: Principal) -> None:
    """A changed method (levels, appetite) invalidates acceptances made under the old one."""
    accepted = list(db.scalars(select(Risk).where(Risk.status == "accepted")))
    acc = _acceptance_map(db, [r.id for r in accepted])
    mc = MethodCache(db)
    for r in accepted:
        req = acc.get(r.id)
        if (
            req is None
            or req.status != "approved"
            or req.target_version != acceptance_version(r, mc.get(r.entity_id))[0]
        ):
            r.status = "assessed"
            audit(
                db,
                pr,
                "update",
                "risk",
                r.id,
                f"{r.code} {r.title}",
                entity_id=r.entity_id,
                details={
                    "status": ["accepted", "assessed"],
                    "acceptance_lapsed": [False, True],
                    "reason": "risk method changed",
                },
            )
            index_object(db, "risk", r.id, commit=False)


@router.delete("/methods", response_model=Ok)
def delete_method_override(
    entity_id: int, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)
):
    """Remove an entity override; the entity falls back to the group default."""
    pr.ensure("risks.manage", entity_id, write=True)
    m = db.scalar(select(RiskMethod).where(RiskMethod.entity_id == entity_id))
    if m is None:
        raise HTTPException(404, "This entity uses the group default method")
    audit(db, pr, "delete", "risk_method", m.id, m.name, entity_id=entity_id)
    db.delete(m)
    db.flush()
    _lapse_stale_acceptances(db, pr)
    db.commit()
    return Ok()


@router.get("/threats", response_model=list[ThreatOut])
def list_threats(
    include_inactive: bool = False,
    pr: Principal = Depends(require("risks.read")),
    db: Session = Depends(get_db),
):
    stmt = select(Threat).order_by(Threat.sort_order, Threat.id)
    if not include_inactive:
        stmt = stmt.where(Threat.is_active.is_(True))
    return [ThreatOut.model_validate(t, from_attributes=True) for t in db.scalars(stmt)]


class ThreatIn(BaseModel):
    code: str = Field(min_length=1, max_length=20)
    title: str = Field(min_length=1, max_length=300)
    category: str = Field(min_length=1, max_length=60)
    source: str | None = Field(default=None, max_length=120)
    is_active: bool = True


@router.post("/threats", response_model=ThreatOut)
def create_threat(body: ThreatIn, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    pr.ensure("risks.manage", None, write=True)
    if db.scalar(select(Threat.id).where(Threat.code == body.code.strip())):
        raise HTTPException(409, "A threat with this code already exists")
    last = db.scalar(select(Threat.sort_order).order_by(Threat.sort_order.desc()).limit(1)) or 0
    t = Threat(**{**body.model_dump(), "code": body.code.strip()}, sort_order=last + 1)
    db.add(t)
    db.flush()
    audit(db, pr, "create", "risk_threat", t.id, f"{t.code} {t.title}")
    db.commit()
    return ThreatOut.model_validate(t, from_attributes=True)


@router.put("/threats/{threat_id}", response_model=ThreatOut)
def update_threat(
    threat_id: int, body: ThreatIn, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)
):
    pr.ensure("risks.manage", None, write=True)
    t = get_or_404(db, Threat, threat_id, "Threat")
    if db.scalar(select(Threat.id).where(Threat.code == body.code.strip(), Threat.id != t.id)):
        raise HTTPException(409, "A threat with this code already exists")
    diff = apply_changes(t, {**body.model_dump(), "code": body.code.strip()})
    audit(db, pr, "update", "risk_threat", t.id, f"{t.code} {t.title}", details=diff)
    db.commit()
    return ThreatOut.model_validate(t, from_attributes=True)


# --- register --------------------------------------------------------------------------------------------------


def _mine(db: Session, pr: Principal, r: Risk) -> bool:
    if r.owner_user_id == pr.id:
        return True
    if r.owner_stakeholder_id:
        st = db.get(Stakeholder, r.owner_stakeholder_id)
        return bool(st and st.user_id == pr.id)
    return False


def _filtered(
    db: Session,
    pr: Principal,
    *,
    entity_id: int | None,
    group: bool,
    status: list[str] | None,
    owner_user_id: int | None,
    mine: bool,
    review_overdue: bool,
    q: str | None,
    threat_id: int | None,
    category: str | None,
    treatment: str | None,
    above_appetite: bool,
) -> tuple[list[Risk], MethodCache]:
    mc = MethodCache(db)
    rows = _visible_risks(db, pr)
    today = date.today()
    ql = (q or "").strip().lower()
    out = []
    for r in rows:
        if group and r.entity_id is not None:
            continue
        if entity_id is not None and r.entity_id != entity_id:
            continue
        if status and r.status not in status:
            continue
        if owner_user_id and r.owner_user_id != owner_user_id:
            continue
        if mine and not _mine(db, pr, r):
            continue
        if review_overdue and not (
            r.status != "closed" and r.next_review_date and r.next_review_date < today
        ):
            continue
        if threat_id and r.threat_id != threat_id:
            continue
        if category and r.category != category:
            continue
        if treatment and treatment not in (r.treatments or []):
            continue
        if above_appetite:
            m = mc.get(r.entity_id)
            if not m.above_appetite(m.level(r.residual_likelihood, r.residual_impact)):
                continue
        if (
            ql
            and ql
            not in " ".join(
                filter(
                    None,
                    [r.code, r.title, r.scenario, r.existing_controls, r.threat.title if r.threat else None],
                )
            ).lower()
        ):
            continue
        out.append(r)
    return out, mc


def _filter_params(
    entity_id: int | None = None,
    group: bool = False,
    status: list[Status] | None = Query(default=None),
    owner_user_id: int | None = None,
    mine: bool = False,
    review_overdue: bool = False,
    q: str | None = Query(default=None, max_length=200),
    threat_id: int | None = None,
    category: str | None = Query(default=None, max_length=60),
    treatment: Treatment | None = None,
    above_appetite: bool = False,
) -> dict:
    return {
        "entity_id": entity_id,
        "group": group,
        "status": status,
        "owner_user_id": owner_user_id,
        "mine": mine,
        "review_overdue": review_overdue,
        "q": q,
        "threat_id": threat_id,
        "category": category,
        "treatment": treatment,
        "above_appetite": above_appetite,
    }


@router.get("", response_model=list[RiskOut])
def list_risks(
    f: dict = Depends(_filter_params),
    level: Level | None = None,
    basis: Literal["inherent", "residual"] = "residual",
    likelihood: int | None = Query(default=None, ge=1, le=5),
    impact: int | None = Query(default=None, ge=1, le=5),
    pr: Principal = Depends(require("risks.read")),
    db: Session = Depends(get_db),
):
    rows, mc = _filtered(db, pr, **f)
    if level or likelihood or impact:
        keep = []
        for r in rows:
            ll = r.inherent_likelihood if basis == "inherent" else r.residual_likelihood
            ii = r.inherent_impact if basis == "inherent" else r.residual_impact
            if likelihood and ll != likelihood:
                continue
            if impact and ii != impact:
                continue
            if level and mc.get(r.entity_id).level(ll, ii) != level:
                continue
            keep.append(r)
        rows = keep
    ids = [r.id for r in rows]
    acc, counts = _acceptance_map(db, ids), _link_counts(db, ids)
    return [risk_out(db, r, mc, acc, counts) for r in rows]


@router.get("/heatmap")
def heatmap(
    f: dict = Depends(_filter_params),
    pr: Principal = Depends(require("risks.read")),
    db: Session = Depends(get_db),
):
    """Counts per matrix cell (likelihood x impact) for inherent and residual ratings. Cells are coloured with the
    effective method of the selected entity (group default when several entities are shown)."""
    if f["entity_id"] is not None:
        pr.ensure("risks.read", f["entity_id"])  # the entity's method is only shown to readers of its risks
    rows, mc = _filtered(db, pr, **f)
    method = mc.get(f["entity_id"])
    grid = {"inherent": [[0] * 5 for _ in range(5)], "residual": [[0] * 5 for _ in range(5)]}
    unrated = {"inherent": 0, "residual": 0}
    by_level = {b: {lv: 0 for lv in LEVELS} for b in grid}
    per_entity: dict[int | None, dict] = {}
    for r in rows:
        m = mc.get(r.entity_id)
        pe = per_entity.setdefault(
            r.entity_id,
            {
                "entity": EntityBrief.model_validate(r.entity) if r.entity else None,
                "total": 0,
                "residual": {lv: 0 for lv in LEVELS},
                "above_appetite": 0,
            },
        )
        pe["total"] += 1
        for basis, ll, ii in (
            ("inherent", r.inherent_likelihood, r.inherent_impact),
            ("residual", r.residual_likelihood, r.residual_impact),
        ):
            if ll and ii:
                grid[basis][ll - 1][ii - 1] += 1
                lv = m.level(ll, ii)
                by_level[basis][lv] += 1
                if basis == "residual":
                    pe["residual"][lv] += 1
                    if m.above_appetite(lv):
                        pe["above_appetite"] += 1
            else:
                unrated[basis] += 1
    return {
        "method": method_out(db, method, f["entity_id"]),
        "total": len(rows),
        "cells": grid,
        "unrated": unrated,
        "by_level": by_level,
        "entities": sorted(
            per_entity.values(),
            key=lambda x: (x["entity"] is not None, x["entity"].code if x["entity"] else ""),
        ),
    }


# --- export ---------------------------------------------------------------------------------------------------

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_HEAD_FILL = PatternFill("solid", fgColor="0D6770")
_LEVEL_LABEL = {"low": "Low", "medium": "Medium", "high": "High", "very_high": "Very high", None: ""}


def _put(ws, row: int, col: int, value) -> None:
    """Write a cell; user-supplied strings are never interpreted as formulas (CSV/Excel injection)."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, datetime):
        value = value.replace(tzinfo=None)
    cell.value = value
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str) and len(value) > 40)


def _sheet(wb, title: str, headers: list[tuple[str, int]], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    for c, (h, w) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    for r_i, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            _put(ws, r_i, c, v)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"


def _owner(r: Risk) -> str | None:
    return (
        r.owner_user.display_name
        if r.owner_user
        else (r.owner_stakeholder.name if r.owner_stakeholder else None)
    )


@router.get("/export")
def export_register(
    f: dict = Depends(_filter_params),
    pr: Principal = Depends(require("risks.read", "reports.export")),
    db: Session = Depends(get_db),
):
    rows, mc = _filtered(db, pr, **f)
    # risks.read alone does not allow exporting: also limit to the entities reports.export covers
    rows = [r for r in rows if pr.can_read("reports.export", r.entity_id)]
    ids = [r.id for r in rows]
    acc = _acceptance_map(db, ids)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    reg = []
    for r in rows:
        m = mc.get(r.entity_id)
        inh = m.level(r.inherent_likelihood, r.inherent_impact)
        res = m.level(r.residual_likelihood, r.residual_impact)
        req = acc.get(r.id)
        reg.append(
            [
                r.code,
                r.entity.code if r.entity else "Group",
                r.title,
                r.scenario,
                f"{r.threat.code} {r.threat.title}" if r.threat else None,
                r.category,
                _owner(r),
                r.inherent_likelihood,
                r.inherent_impact,
                _score(r.inherent_likelihood, r.inherent_impact),
                _LEVEL_LABEL[inh],
                r.existing_controls,
                r.residual_likelihood,
                r.residual_impact,
                _score(r.residual_likelihood, r.residual_impact),
                _LEVEL_LABEL[res],
                "Yes" if m.above_appetite(res) else "No",
                ", ".join(t.capitalize() for t in (r.treatments or [])),
                r.treatment_justification,
                r.acceptance_category,
                r.status.replace("_", " ").capitalize(),
                (req.status.replace("_", " ") if req else None),
                r.accepted_by,
                r.accepted_at,
                r.last_reviewed_at,
                r.next_review_date,
            ]
        )
    _sheet(
        wb,
        "Risk register",
        [
            ("Code", 10),
            ("Entity", 9),
            ("Risk", 40),
            ("Scenario", 50),
            ("Threat", 30),
            ("Category", 20),
            ("Owner", 22),
            ("Inherent likelihood", 11),
            ("Inherent impact", 11),
            ("Inherent score", 10),
            ("Inherent level", 11),
            ("Existing controls", 45),
            ("Residual likelihood", 11),
            ("Residual impact", 11),
            ("Residual score", 10),
            ("Residual level", 11),
            ("Above appetite", 10),
            ("Treatment", 18),
            ("Treatment justification", 45),
            ("Acceptance category", 28),
            ("Status", 13),
            ("Acceptance approval", 16),
            ("Accepted by", 20),
            ("Accepted at", 16),
            ("Last reviewed", 13),
            ("Next review", 13),
        ],
        reg,
    )

    plan = []
    links = (
        db.scalars(
            select(RiskLink)
            .where(RiskLink.risk_id.in_(ids), RiskLink.link_type == "treatment")
            .order_by(RiskLink.risk_id, RiskLink.id)
        ).all()
        if ids
        else []
    )
    by_risk: dict[int, list[RiskLink]] = {}
    for lk in links:
        by_risk.setdefault(lk.risk_id, []).append(lk)
    for r in rows:
        m = mc.get(r.entity_id)
        base = [
            r.code,
            r.title,
            r.entity.code if r.entity else "Group",
            _LEVEL_LABEL[m.level(r.residual_likelihood, r.residual_impact)],
            ", ".join(t.capitalize() for t in (r.treatments or [])),
            _owner(r),
        ]
        acts = [lk for lk in by_risk.get(r.id, []) if can_access_target(db, pr, lk.target_type, lk.target_id)]
        if not acts:
            plan.append([*base, None, None, None, None, None, None, None])
        for lk in acts:
            if lk.target_type == "action_item" and (ai := db.get(ActionItem, lk.target_id)):
                plan.append(
                    [
                        *base,
                        "Action plan item",
                        ai.code,
                        ai.title,
                        ai.responsible.name if ai.responsible else None,
                        ai.target_date,
                        ai.status,
                        None,
                    ]
                )
            elif lk.target_type == "work_package" and (wp := db.get(WorkPackage, lk.target_id)):
                plan.append(
                    [
                        *base,
                        "Work package",
                        wp.code,
                        wp.title,
                        wp.owner.display_name if wp.owner else None,
                        wp.planned_finish,
                        wp.status,
                        round((wp.progress or 0) * 100),
                    ]
                )
    _sheet(
        wb,
        "Treatment plan",
        [
            ("Risk", 10),
            ("Risk title", 40),
            ("Entity", 9),
            ("Residual level", 12),
            ("Treatment", 18),
            ("Risk owner", 22),
            ("Measure type", 16),
            ("Measure", 10),
            ("Measure title", 45),
            ("Responsible", 22),
            ("Target date", 13),
            ("Status", 13),
            ("Progress %", 10),
        ],
        plan,
    )

    mrows = []
    ents = [None] + [
        e.id
        for e in db.scalars(select(Entity).order_by(Entity.sort_order, Entity.code))
        if pr.can_read("risks.read", e.id)
    ]
    for eid in ents:
        m = mc.get(eid)
        ent = db.get(Entity, eid) if eid else None
        mrows.append(
            [
                ent.code if ent else "Group default",
                "group default" if m.inherited else m.name,
                f"{m.threshold_medium}",
                f"{m.threshold_high}",
                f"{m.threshold_very_high}",
                _LEVEL_LABEL[m.appetite_level],
                m.review_cycle_months,
                "; ".join(f"{i + 1} {s['label']}" for i, s in enumerate(m.likelihood_scale)),
                "; ".join(f"{i + 1} {s['label']}" for i, s in enumerate(m.impact_scale)),
                m.acceptance_criteria,
            ]
        )
    _sheet(
        wb,
        "Method",
        [
            ("Scope", 16),
            ("Method", 24),
            ("Medium from score", 10),
            ("High from score", 10),
            ("Very high from score", 10),
            ("Risk appetite", 12),
            ("Review cycle (months)", 12),
            ("Likelihood scale", 50),
            ("Impact scale", 50),
            ("Acceptance criteria", 70),
        ],
        mrows,
    )

    buf = io.BytesIO()
    wb.save(buf)
    audit(db, pr, "export", "risk", None, f"Risk register ({len(rows)} risks)")
    db.commit()
    name = f"NIS2_Risk_Register_{date.today().isoformat()}.xlsx"
    return Response(
        buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{name}"'}
    )


# --- link helpers --------------------------------------------------------------------------------------------


def _section(req: Requirement) -> str:
    """Requirement section without its leading chapter number ("4.2 Backup management" -> "Backup management")."""
    return re.sub(r"^[\d.]+\s*", "", req.section or "") or req.number


def _answer_info(a: Answer) -> dict:
    return {
        "entity_code": a.assessment.entity.code,
        "requirement": a.requirement.number,
        "section": _section(a.requirement),
        "implemented": a.implemented,
        "score": a.score,
        "gap": a.gap_description,
    }


def _link_out(db: Session, pr: Principal, lk: RiskLink) -> dict:
    base = {
        "id": lk.id,
        "link_type": lk.link_type,
        "target_type": lk.target_type,
        "target_id": lk.target_id,
        "note": lk.note,
        "created_at": lk.created_at,
        "label": None,
        "url": None,
        "status": None,
        "due": None,
        "owner": None,
        "progress": None,
        "info": None,
        "exists": True,
        "accessible": True,
    }
    try:
        ad = approvals.get_adapter(lk.target_type)
    except HTTPException:
        return {**base, "label": lk.target_type, "exists": False}
    obj = db.get(ad.model, lk.target_id)
    if obj is None:
        return {**base, "label": "Deleted object", "exists": False}
    if not can_access_target(db, pr, lk.target_type, lk.target_id):
        return {**base, "label": "No access", "accessible": False}
    base.update(label=ad.label(obj), url=ad.url(obj) if ad.url else None, status=getattr(obj, "status", None))
    if isinstance(obj, ActionItem):
        base.update(
            due=obj.target_date,
            owner=obj.responsible.name if obj.responsible else None,
            info={"due_status": due_status(obj.target_date, obj.status), "entity_code": obj.entity.code},
        )
    elif isinstance(obj, WorkPackage):
        base.update(
            due=obj.planned_finish, owner=obj.owner.display_name if obj.owner else None, progress=obj.progress
        )
    elif isinstance(obj, Answer):
        base.update(status=obj.implemented, info=_answer_info(obj))
    return base


def _risk_links(db: Session, pr: Principal, r: Risk) -> list[dict]:
    rows = db.scalars(select(RiskLink).where(RiskLink.risk_id == r.id).order_by(RiskLink.id))
    return [_link_out(db, pr, lk) for lk in rows]


class LinkIn(BaseModel):
    link_type: Literal["gap", "control", "affected", "treatment"]
    target_type: str = Field(min_length=1, max_length=30)
    target_id: int
    note: str | None = Field(default=None, max_length=400)


def _add_link(db: Session, pr: Principal, r: Risk, body: LinkIn) -> RiskLink:
    if body.link_type in ("gap", "control") and body.target_type != "answer":
        raise HTTPException(400, "Gaps and existing controls are linked as assessment answers")
    if body.link_type == "treatment" and body.target_type not in ("action_item", "work_package"):
        raise HTTPException(400, "Treatment measures are action plan items or work packages")
    if body.target_type == "risk" and body.target_id == r.id:
        raise HTTPException(400, "A risk cannot be linked to itself")
    obj = check_target(db, pr, body.target_type, body.target_id)
    if isinstance(obj, Answer) and r.entity_id is not None and obj.assessment.entity_id != r.entity_id:
        raise HTTPException(400, "The answer belongs to another entity than the risk")
    if db.scalar(
        select(RiskLink.id).where(
            RiskLink.risk_id == r.id,
            RiskLink.link_type == body.link_type,
            RiskLink.target_type == body.target_type,
            RiskLink.target_id == body.target_id,
        )
    ):
        raise HTTPException(409, "This object is already linked")
    lk = RiskLink(
        risk_id=r.id,
        link_type=body.link_type,
        target_type=body.target_type,
        target_id=body.target_id,
        note=(body.note or "").strip() or None,
        created_by_id=pr.id,
    )
    db.add(lk)
    db.flush()
    return lk


@router.get("/link-options")
def link_options(
    target_type: str = Query(max_length=30),
    q: str | None = Query(default=None, max_length=200),
    entity_id: int | None = None,
    pr: Principal = Depends(require("risks.read")),
    db: Session = Depends(get_db),
):
    """Candidates for linking (answers of the current round, action items, work packages, documents, ...)."""
    if target_type not in TARGETS:
        raise HTTPException(400, "Unknown target type")
    ql = f"%{(q or '').strip()}%"
    out = []
    if target_type == "answer":
        stmt = (
            select(Answer)
            .join(Assessment, Assessment.id == Answer.assessment_id)
            .join(AssessmentRound, AssessmentRound.id == Assessment.round_id)
            .join(Requirement, Requirement.id == Answer.requirement_id)
            .where(AssessmentRound.is_current.is_(True))
            .order_by(Requirement.sort_order, Requirement.id)
        )
        if entity_id is not None:
            stmt = stmt.where(Assessment.entity_id == entity_id)
        if q:
            stmt = stmt.where(
                or_(
                    Requirement.number.ilike(ql),
                    Requirement.section.ilike(ql),
                    Answer.gap_description.ilike(ql),
                )
            )
        for a in db.scalars(stmt.limit(400)):
            if not pr.can_read("assessments.read", a.assessment.entity_id):
                continue
            out.append(
                {
                    "id": a.id,
                    "label": f"{a.requirement.number} {_section(a.requirement)}",
                    "hint": f"{a.assessment.entity.code}: {a.implemented or 'not answered'}",
                    "entity_code": a.assessment.entity.code,
                }
            )
            if len(out) >= 60:
                break
        return out
    ad = approvals.get_adapter(target_type)
    model = ad.model
    stmt = select(model)
    cols = [
        getattr(model, c) for c in ("code", "key", "number", "title", "name", "section") if hasattr(model, c)
    ]
    if q and cols:
        stmt = stmt.where(or_(*[c.ilike(ql) for c in cols]))
    if entity_id is not None and hasattr(model, "entity_id"):
        stmt = stmt.where(or_(model.entity_id == entity_id, model.entity_id.is_(None)))
    order = model.code if hasattr(model, "code") else model.id
    for obj in db.scalars(stmt.order_by(order).limit(300)):
        if not can_access_target(db, pr, target_type, obj.id):
            continue
        ent = getattr(obj, "entity", None)
        out.append(
            {
                "id": obj.id,
                "label": ad.label(obj),
                "hint": getattr(obj, "status", None),
                "entity_code": ent.code if ent is not None and hasattr(ent, "code") else None,
            }
        )
        if len(out) >= 60:
            break
    return out


@router.get("/linked")
def linked_risks(
    target_type: str = Query(max_length=30),
    target_id: int = Query(),
    pr: Principal = Depends(require("risks.read")),
    db: Session = Depends(get_db),
):
    """Risks linked to an object (e.g. the risks created from an assessment answer)."""
    check_target(db, pr, target_type, target_id)
    rows = db.execute(
        select(Risk, RiskLink.link_type)
        .join(RiskLink, RiskLink.risk_id == Risk.id)
        .where(RiskLink.target_type == target_type, RiskLink.target_id == target_id)
        .order_by(Risk.code)
    ).all()
    mc = MethodCache(db)
    seen, out = set(), []
    for r, lt in rows:
        if r.id in seen or not pr.can_read("risks.read", r.entity_id):
            continue
        seen.add(r.id)
        o = risk_out(db, r, mc).model_dump()
        o["link_type"] = lt
        out.append(o)
    return out


@router.get("/from-incident/{incident_id}")
def draft_from_incident(
    incident_id: int, pr: Principal = Depends(require("risks.read")), db: Session = Depends(get_db)
):
    """Prefill for a new risk created from an incident's root cause and lessons learned."""
    i = check_target(db, pr, "incident", incident_id)
    scenario = "\n\n".join(
        filter(
            None,
            [
                f"Incident {i.code} ({i.detected_at:%d %b %Y}): {i.title}."
                if i.detected_at
                else f"Incident {i.code}: {i.title}.",
                f"Root cause: {i.root_cause}" if i.root_cause else None,
                f"Impact: {i.impact}" if i.impact else None,
                f"Lessons learned: {i.lessons_learned}" if i.lessons_learned else None,
            ],
        )
    )
    return {
        "incident_id": i.id,
        "code": i.code,
        "entity": EntityBrief.model_validate(i.entity) if i.entity else None,
        "title": f"Recurrence of: {i.title}"[:400],
        "scenario": scenario,
        "existing_controls": f"Measures taken: {i.measures}" if i.measures else None,
        "can_create": pr.can_write("risks.manage", i.entity_id),
        "existing_risks": linked_risks("incident", i.id, pr, db),
    }


@router.get("/from-answer/{answer_id}")
def draft_from_answer(
    answer_id: int, pr: Principal = Depends(require("risks.read")), db: Session = Depends(get_db)
):
    """Prefill for a new risk created from a gap (No / Partial answer)."""
    a = check_target(db, pr, "answer", answer_id)
    req, ent = a.requirement, a.assessment.entity
    section = _section(req)
    state = {"No": "not implemented", "Partial": "only partially implemented"}.get(
        a.implemented or "", "not confirmed as implemented"
    )
    scenario = "\n\n".join(
        filter(
            None,
            [
                f"Requirement {req.number} ({section}) is {state} at {ent.name}.",
                f"Identified gap: {a.gap_description}" if a.gap_description else None,
                "Without this measure, a threat could exploit the weakness and disrupt production, delivery to customers "
                "or the confidentiality of information.",
            ],
        )
    )
    controls = (
        "\n".join(
            filter(
                None,
                [
                    f"Documented: {a.documented}" if a.documented else None,
                    f"Effectiveness: {a.effectiveness}" if a.effectiveness else None,
                    f"Evidence: {a.evidence}" if a.evidence else None,
                ],
            )
        )
        or None
    )
    return {
        "answer_id": a.id,
        "entity": EntityBrief.model_validate(ent),
        "requirement": {"id": req.id, "number": req.number, "section": section},
        "implemented": a.implemented,
        "gap_description": a.gap_description,
        "recommended_action": a.recommended_action,
        "title": f"Gap in {section[0].lower() + section[1:]} ({req.number})"[:400],
        "scenario": scenario,
        "existing_controls": controls,
        "can_create": pr.can_write("risks.manage", ent.id),
        "existing_risks": linked_risks("answer", a.id, pr, db),
    }


# --- CRUD ----------------------------------------------------------------------------------------------------


class _RiskValidators(BaseModel):
    """Validators shared by RiskIn and RiskPatch: treatments in canonical order, known acceptance category."""

    @field_validator("treatments", check_fields=False)
    @classmethod
    def _uniq(cls, v: list[str] | None) -> list[str] | None:
        return None if v is None else [t for t in TREATMENTS if t in set(v)]

    @field_validator("acceptance_category", check_fields=False)
    @classmethod
    def _cat(cls, v: str | None) -> str | None:
        if v and v not in ACCEPTANCE_CATEGORIES:
            raise ValueError("unknown acceptance category")
        return v or None


class RiskIn(_RiskValidators):
    entity_id: int | None = None
    title: str = Field(min_length=1, max_length=400)
    scenario: str | None = Field(default=None, max_length=20000)
    threat_id: int | None = None
    category: str | None = Field(default=None, max_length=60)
    owner_stakeholder_id: int | None = None
    owner_user_id: int | None = None
    inherent_likelihood: int | None = Field(default=None, ge=1, le=5)
    inherent_impact: int | None = Field(default=None, ge=1, le=5)
    existing_controls: str | None = Field(default=None, max_length=20000)
    residual_likelihood: int | None = Field(default=None, ge=1, le=5)
    residual_impact: int | None = Field(default=None, ge=1, le=5)
    treatments: list[Treatment] = []
    treatment_justification: str | None = Field(default=None, max_length=20000)
    acceptance_category: str | None = Field(default=None, max_length=120)
    status: Status | None = None
    next_review_date: date | None = None
    source_answer_id: int | None = None  # create from a gap: links the answer
    source_incident_id: int | None = (
        None  # create from an incident (lessons learned): links it as affected object
    )


class RiskPatch(_RiskValidators):
    entity_id: int | None = None
    title: str | None = Field(default=None, min_length=1, max_length=400)
    scenario: str | None = Field(default=None, max_length=20000)
    threat_id: int | None = None
    category: str | None = Field(default=None, max_length=60)
    owner_stakeholder_id: int | None = None
    owner_user_id: int | None = None
    inherent_likelihood: int | None = Field(default=None, ge=1, le=5)
    inherent_impact: int | None = Field(default=None, ge=1, le=5)
    existing_controls: str | None = Field(default=None, max_length=20000)
    residual_likelihood: int | None = Field(default=None, ge=1, le=5)
    residual_impact: int | None = Field(default=None, ge=1, le=5)
    treatments: list[Treatment] | None = None
    treatment_justification: str | None = Field(default=None, max_length=20000)
    acceptance_category: str | None = Field(default=None, max_length=120)
    status: Status | None = None
    next_review_date: date | None = None


def _check_refs(db: Session, data: dict) -> None:
    if data.get("threat_id") and not db.get(Threat, data["threat_id"]):
        raise HTTPException(400, "Unknown threat")
    if data.get("owner_stakeholder_id") and not db.get(Stakeholder, data["owner_stakeholder_id"]):
        raise HTTPException(400, "Unknown stakeholder")
    if data.get("owner_user_id") and not db.get(User, data["owner_user_id"]):
        raise HTTPException(400, "Unknown user")
    if data.get("entity_id") and not db.get(Entity, data["entity_id"]):
        raise HTTPException(400, "Unknown entity")


def _check_owner(
    db: Session, entity_id: int | None, owner_user_id: int | None, owner_stakeholder_id: int | None
) -> None:
    """The risk owner (user, or the user linked to the owner role) must be able to see the risk."""
    uids = []
    if owner_user_id:
        uids.append(owner_user_id)
    if owner_stakeholder_id and (st := db.get(Stakeholder, owner_stakeholder_id)) and st.user_id:
        uids.append(st.user_id)
    for uid in uids:
        u = db.get(User, uid)
        if u is None or not u.is_active or not build_principal(u).can_read("risks.read", entity_id):
            name = u.display_name if u else "The owner"
            raise HTTPException(
                400,
                f"{name} cannot see risks of this {'entity' if entity_id else 'scope'} "
                "and cannot own this risk",
            )


def _auto_status(r: Risk) -> None:
    if (
        r.status == "identified"
        and r.inherent_likelihood
        and r.inherent_impact
        and r.residual_likelihood
        and r.residual_impact
    ):
        r.status = "assessed"


def _detail(db: Session, pr: Principal, r: Risk) -> dict:
    mc = MethodCache(db)
    return {
        "risk": risk_out(db, r, mc),
        "method": method_out(db, mc.get(r.entity_id), r.entity_id),
        "links": _risk_links(db, pr, r),
        "documents": doc_links_for(db, pr, "risk", r.id),
        "can_edit": pr.can_write("risks.manage", r.entity_id),
    }


@router.post("")
def create_risk(body: RiskIn, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    pr.ensure("risks.manage", body.entity_id, write=True)
    data = body.model_dump(exclude={"source_answer_id", "source_incident_id"})
    _check_refs(db, data)
    _check_owner(db, data["entity_id"], data["owner_user_id"], data["owner_stakeholder_id"])
    if data["status"] == "accepted":
        raise HTTPException(
            400, "A risk is accepted through an approval. Create it first, then request acceptance."
        )
    m = effective_method(db, body.entity_id)
    data["status"] = data["status"] or "identified"
    if not data["next_review_date"]:
        data["next_review_date"] = add_months(date.today(), m.review_cycle_months or 12)
    r = Risk(code=next_code(db, Risk.code, "RSK-"), created_by_id=pr.id, **data)
    if r.status == "closed":
        r.closed_at = utcnow()
    _auto_status(r)
    db.add(r)
    db.flush()
    if body.source_answer_id:
        _add_link(db, pr, r, LinkIn(link_type="gap", target_type="answer", target_id=body.source_answer_id))
    if body.source_incident_id:
        _add_link(
            db, pr, r, LinkIn(link_type="affected", target_type="incident", target_id=body.source_incident_id)
        )
    source = {
        k: v
        for k, v in (
            ("source_answer_id", body.source_answer_id),
            ("source_incident_id", body.source_incident_id),
        )
        if v
    }
    audit(
        db, pr, "create", "risk", r.id, f"{r.code} {r.title}", entity_id=r.entity_id, details=source or None
    )
    db.commit()
    index_object(db, "risk", r.id)
    return _detail(db, pr, r)


def _get_risk(db: Session, pr: Principal, risk_id: int) -> Risk:
    r = get_or_404(db, Risk, risk_id, "Risk")
    if not pr.can_read("risks.read", r.entity_id):
        raise HTTPException(404, "Risk not found")
    return r


@router.get("/{risk_id}")
def get_risk(risk_id: int, pr: Principal = Depends(require("risks.read")), db: Session = Depends(get_db)):
    return _detail(db, pr, _get_risk(db, pr, risk_id))


# changing one of these ends an acceptance (they are part of the pinned acceptance version)
_ACCEPT_FIELDS = {
    "residual_likelihood",
    "residual_impact",
    "treatments",
    "acceptance_category",
    "treatment_justification",
    "entity_id",
}


@router.patch("/{risk_id}")
def update_risk(
    risk_id: int, body: RiskPatch, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)
):
    r = _get_risk(db, pr, risk_id)
    pr.ensure("risks.manage", r.entity_id, write=True)
    data = body.model_dump(exclude_unset=True)
    if "entity_id" in data and data["entity_id"] != r.entity_id:
        pr.ensure("risks.manage", data["entity_id"], write=True)
    if "title" in data and not data["title"]:
        raise HTTPException(400, "The title cannot be empty")
    if "treatments" in data and data["treatments"] is None:
        data["treatments"] = []
    _check_refs(db, data)
    if {"entity_id", "owner_user_id", "owner_stakeholder_id"} & data.keys():
        _check_owner(
            db,
            data.get("entity_id", r.entity_id),
            data.get("owner_user_id", r.owner_user_id),
            data.get("owner_stakeholder_id", r.owner_stakeholder_id),
        )
    if data.get("status") == "accepted" and r.status != "accepted":
        raise HTTPException(400, "A risk is accepted through an approval: use Request acceptance")
    old_status = r.status
    diff = apply_changes(r, data)
    if old_status == "accepted" and r.status == "accepted" and _ACCEPT_FIELDS & diff.keys():
        # the acceptance covered the previous rating / treatment / justification / entity: it lapses
        r.status = "assessed"
    if old_status == "accepted" and r.status != "accepted":
        diff["acceptance_lapsed"] = [False, True]
    _auto_status(r)
    if r.status != old_status:
        diff["status"] = [old_status, r.status]
        r.closed_at = utcnow() if r.status == "closed" else None
    if diff:
        audit(db, pr, "update", "risk", r.id, f"{r.code} {r.title}", entity_id=r.entity_id, details=diff)
    db.commit()
    index_object(db, "risk", r.id)
    return _detail(db, pr, r)


class ReviewIn(BaseModel):
    note: str | None = Field(default=None, max_length=4000)


@router.post("/{risk_id}/review")
def review_risk(
    risk_id: int, body: ReviewIn, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)
):
    """Record the periodic review (IR 2.1.4): sets the next review date one review cycle ahead."""
    r = _get_risk(db, pr, risk_id)
    pr.ensure("risks.manage", r.entity_id, write=True)
    m = effective_method(db, r.entity_id)
    old = r.next_review_date
    r.last_reviewed_at = date.today()
    r.next_review_date = add_months(date.today(), m.review_cycle_months or 12)
    audit(
        db,
        pr,
        "review",
        "risk",
        r.id,
        f"{r.code} {r.title}",
        entity_id=r.entity_id,
        details={
            "next_review_date": [old.isoformat() if old else None, r.next_review_date.isoformat()],
            "note": (body.note or "").strip() or None,
        },
    )
    db.commit()
    return _detail(db, pr, r)


@router.delete("/{risk_id}", response_model=Ok)
def delete_risk(risk_id: int, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)):
    r = _get_risk(db, pr, risk_id)
    pr.ensure("risks.manage", r.entity_id, write=True)
    if r.status == "accepted":
        raise HTTPException(400, "Accepted risks are part of the acceptance record. Close the risk instead.")
    audit(db, pr, "delete", "risk", r.id, f"{r.code} {r.title}", entity_id=r.entity_id)
    approvals.withdraw_for_deleted_target(db, "risk_acceptance", r.id)
    cleanup_target(db, "risk", r.id)
    db.delete(r)
    db.commit()
    remove_object(db, "risk", risk_id)
    return Ok()


@router.post("/{risk_id}/links")
def add_link(
    risk_id: int, body: LinkIn, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)
):
    r = _get_risk(db, pr, risk_id)
    pr.ensure("risks.manage", r.entity_id, write=True)
    lk = _add_link(db, pr, r, body)
    audit(
        db,
        pr,
        "link",
        "risk",
        r.id,
        f"{r.code} {r.title}",
        entity_id=r.entity_id,
        details={"link_type": lk.link_type, "target_type": lk.target_type, "target_id": lk.target_id},
    )
    if body.link_type == "treatment" and r.status in ("identified", "assessed"):
        r.status = "in_treatment"
        audit(
            db,
            pr,
            "update",
            "risk",
            r.id,
            f"{r.code} {r.title}",
            entity_id=r.entity_id,
            details={"status": ["assessed", "in_treatment"]},
        )
    db.commit()
    return _link_out(db, pr, lk)


@router.delete("/{risk_id}/links/{link_id}", response_model=Ok)
def remove_link(
    risk_id: int, link_id: int, pr: Principal = Depends(get_principal), db: Session = Depends(get_db)
):
    r = _get_risk(db, pr, risk_id)
    pr.ensure("risks.manage", r.entity_id, write=True)
    lk = get_or_404(db, RiskLink, link_id, "Link")
    if lk.risk_id != r.id:
        raise HTTPException(404, "Link not found")
    audit(
        db,
        pr,
        "unlink",
        "risk",
        r.id,
        f"{r.code} {r.title}",
        entity_id=r.entity_id,
        details={"link_type": lk.link_type, "target_type": lk.target_type, "target_id": lk.target_id},
    )
    db.delete(lk)
    db.commit()
    return Ok()


@router.get("/{risk_id}/history")
def risk_history(risk_id: int, pr: Principal = Depends(require("risks.read")), db: Session = Depends(get_db)):
    r = _get_risk(db, pr, risk_id)
    rows = db.scalars(
        select(AuditEvent)
        .where(
            or_(AuditEvent.target_type == "risk", AuditEvent.target_type == "risk_acceptance"),
            AuditEvent.target_id == r.id,
        )
        .order_by(AuditEvent.at.desc(), AuditEvent.id.desc())
        .limit(300)
    )
    users = {u.id: u.display_name for u in db.scalars(select(User))}
    return [
        {
            "id": e.id,
            "at": e.at,
            "action": e.action,
            "target_type": e.target_type,
            "actor": users.get(e.actor_id) or e.actor_label,
            "details": e.details,
        }
        for e in rows
    ]
