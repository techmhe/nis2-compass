"""Supplier security questionnaires: default question sets per tier and the Excel round-trip
(export a workbook to send to the supplier, import the returned answers)."""

import io
import warnings
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.features.suppliers.logic import normalise_answer, rating_for, score_answers

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
MAX_ROWS = 500
TIER_ORDER = {"standard": 0, "high": 1, "critical": 2}

# (id, section, question, reference, weight, minimum tier) – based on IR 2024/2690 Annex point 5.1 and
# ENISA "Good practices for supply chain cybersecurity" (2023). Paraphrased; no copyrighted standard text.
DEFAULT_QUESTIONS: list[tuple[str, str, str, str, int, str]] = [
    (
        "Q01",
        "Governance",
        "Do you operate an information security management system with a documented security "
        "policy approved by your management?",
        "IR 5.1.2(a)",
        2,
        "standard",
    ),
    (
        "Q02",
        "Governance",
        "Have you named a contact point for information security and incident matters for our organisation?",
        "IR 5.2(a)",
        1,
        "standard",
    ),
    (
        "Q03",
        "Certification",
        "Do you hold a valid certification covering the services supplied to us (e.g. ISO/IEC "
        "27001, TISAX, SOC 2, IEC 62443-2-4)? Please attach the certificate and its scope.",
        "IR 5.1.4(b)",
        2,
        "standard",
    ),
    (
        "Q04",
        "People",
        "Are employees working on our services trained in information security at least annually?",
        "IR 5.1.4(b)",
        1,
        "standard",
    ),
    (
        "Q05",
        "People",
        "Do you verify the background of employees with access to our systems or data, where "
        "legally permitted?",
        "IR 5.1.4(c)",
        1,
        "high",
    ),
    (
        "Q06",
        "Access",
        "Is remote access to our systems performed only through our approved access channel, with "
        "individual accounts and multi-factor authentication?",
        "IR 5.1.4(a), 11.7",
        2,
        "standard",
    ),
    (
        "Q07",
        "Access",
        "Do you log privileged activities performed on our systems and make the logs available to "
        "us on request?",
        "IR 5.1.4(a), 3.2",
        1,
        "high",
    ),
    (
        "Q08",
        "Access",
        "Are access rights of your staff to our systems and data revoked without delay when no longer "
        "needed?",
        "IR 11.3",
        1,
        "high",
    ),
    (
        "Q09",
        "Incidents",
        "Will you notify us without undue delay (target within 24 hours) of incidents that may "
        "affect the security of our network and information systems?",
        "IR 5.1.4(d)",
        3,
        "standard",
    ),
    (
        "Q10",
        "Incidents",
        "Do you have a documented incident response process that is tested at least annually?",
        "IR 3.1",
        1,
        "high",
    ),
    (
        "Q11",
        "Vulnerabilities",
        "Do you monitor vulnerabilities in the products and services supplied to us and "
        "inform us about relevant vulnerabilities and available fixes?",
        "IR 5.1.4(f), 6.10",
        2,
        "standard",
    ),
    (
        "Q12",
        "Vulnerabilities",
        "Do you provide security updates for the supplied products for a defined support period?",
        "IR 6.1",
        1,
        "high",
    ),
    (
        "Q13",
        "Development",
        "Do you follow a secure development lifecycle for software or firmware supplied to us?",
        "IR 5.1.2(a), 6.2",
        1,
        "critical",
    ),
    (
        "Q14",
        "Development",
        "Can you provide a software bill of materials (SBOM) for supplied software or devices?",
        "ENISA SCS good practice",
        1,
        "critical",
    ),
    (
        "Q15",
        "Continuity",
        "Do you maintain business continuity and disaster recovery plans for the services "
        "supplied and test them regularly?",
        "IR 4.1",
        2,
        "high",
    ),
    (
        "Q16",
        "Continuity",
        "Are backups of data you process for us protected and restore-tested regularly?",
        "IR 4.2",
        1,
        "high",
    ),
    ("Q17", "Data protection", "Is our data encrypted in transit and at rest?", "IR 9.1", 1, "standard"),
    (
        "Q18",
        "Data protection",
        "Do you process and store our data only in agreed locations and jurisdictions?",
        "IR 5.1.2(c)",
        1,
        "high",
    ),
    (
        "Q19",
        "Subcontracting",
        "Do you inform us before engaging subcontractors with access to our systems or data "
        "and pass on equivalent security requirements to them?",
        "IR 5.1.4(g)",
        2,
        "high",
    ),
    (
        "Q20",
        "Assurance",
        "Do you grant us the right to audit or provide independent audit reports on request?",
        "IR 5.1.4(e)",
        2,
        "standard",
    ),
    (
        "Q21",
        "Exit",
        "At the end of the contract, will you return our information, securely delete it and confirm "
        "the deletion in writing?",
        "IR 5.1.4(h)",
        1,
        "standard",
    ),
    (
        "Q22",
        "Assurance",
        "Do you have the services supplied to us penetration-tested at least annually and share "
        "a summary of the results?",
        "ENISA SCS good practice",
        1,
        "critical",
    ),
    (
        "Q23",
        "Operations",
        "Are the networks used to deliver our services segmented and monitored for security "
        "events (e.g. SOC / SIEM)?",
        "IR 3.2, 12.1",
        1,
        "critical",
    ),
    (
        "Q24",
        "OT remote maintenance",
        "Are OT remote maintenance sessions started only on our request, time-limited, "
        "supervised and recorded?",
        "IEC 62443-2-4 (title only)",
        2,
        "critical",
    ),
    (
        "Q25",
        "Supply chain",
        "Do you assess the cybersecurity of your own critical suppliers?",
        "IR 5.1.6, ENISA SCS good practice",
        1,
        "critical",
    ),
]


def default_questions(tier: str) -> list[dict]:
    lvl = TIER_ORDER[tier]
    return [
        {"id": i, "section": s, "text": t, "reference": r, "weight": w}
        for i, s, t, r, w, min_tier in DEFAULT_QUESTIONS
        if TIER_ORDER[min_tier] <= lvl
    ]


def _safe(ws, row: int, col: int, value):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        cell.data_type = "s"
    return cell


def build_questionnaire_xlsx(supplier, template, answers: dict[str, dict] | None = None) -> bytes:
    """Workbook to send to a supplier. `answers` pre-fills known answers {id: {answer, comment}}."""
    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="0D6770")
    white_bold = Font(bold=True, color="FFFFFF")
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    rows = [
        ("Supplier security questionnaire", None),
        (None, None),
        ("Supplier", supplier.name),
        ("Supplier reference", supplier.code),
        ("Questionnaire", f"{template.title} (tier {template.tier}, version {template.version})"),
        ("Issued on", date.today().isoformat()),
        (None, None),
        (
            "How to answer",
            "Please answer every question on the sheet 'Questionnaire' in column E with Yes, Partial, "
            "No or N/A, and explain or reference evidence in column F. Do not change the IDs in "
            "column A or the sheet names. Return the completed workbook to your contact.",
        ),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        _safe(ws, r, 1, a)
        _safe(ws, r, 2, b)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=14)

    q = wb.create_sheet("Questionnaire")
    headers = ["ID", "Section", "Question", "Reference", "Answer", "Comment / evidence"]
    widths = [8, 20, 80, 22, 12, 60]
    for c, (h, w) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = _safe(q, 1, c, h)
        cell.font = white_bold
        cell.fill = head_fill
        q.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    dv = DataValidation(type="list", formula1='"Yes,Partial,No,N/A"', allow_blank=True)
    q.add_data_validation(dv)
    answers = answers or {}
    for r, item in enumerate(template.questions, start=2):
        prev = answers.get(item["id"], {})
        for c, v in enumerate(
            [
                item["id"],
                item.get("section"),
                item["text"],
                item.get("reference"),
                prev.get("answer"),
                prev.get("comment"),
            ],
            start=1,
        ):
            cell = _safe(q, r, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        dv.add(f"E{r}")
    q.freeze_panes = "A2"

    meta = wb.create_sheet("_meta")
    for r, (k, v) in enumerate(
        [
            ("template_id", template.id),
            ("template_version", template.version),
            ("tier", template.tier),
            ("supplier_id", supplier.id),
            ("supplier_code", supplier.code),
        ],
        start=1,
    ):
        _safe(meta, r, 1, k)
        _safe(meta, r, 2, v)
    meta.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_meta(data: bytes) -> dict:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if "_meta" not in wb.sheetnames:
            return {}
        return {
            str(r[0]): r[1]
            for r in wb["_meta"].iter_rows(min_row=1, max_row=20, max_col=2, values_only=True)
            if r[0]
        }
    finally:
        wb.close()


def parse_questionnaire_xlsx(data: bytes, template) -> dict:
    """Read answers for the template's questions. Returns a per-row report, score, rating and findings."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if "Questionnaire" not in wb.sheetnames:
            raise ValueError("The workbook has no sheet named 'Questionnaire'")
        ws = wb["Questionnaire"]
        rows = []
        for r in ws.iter_rows(min_row=1, max_col=6, values_only=True):
            rows.append(r)
            if len(rows) > MAX_ROWS:
                raise ValueError(f"Too many rows (a questionnaire has at most {MAX_ROWS} rows)")
    finally:
        wb.close()
    header_idx = next((i for i, r in enumerate(rows) if r and str(r[0] or "").strip().lower() == "id"), None)
    if header_idx is None:
        raise ValueError("Header row with 'ID' in column A not found")
    by_id = {q["id"]: q for q in template.questions}
    seen: set[str] = set()
    report, answers, found = [], {}, []
    for offset, r in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        qid = str(r[0]).strip() if r and r[0] is not None else ""
        if not qid:
            continue
        raw = r[4] if len(r) > 4 else None
        comment = str(r[5]).strip() if len(r) > 5 and r[5] is not None else None
        entry = {
            "row": offset,
            "id": qid,
            "answer": None,
            "comment": comment,
            "status": "ok",
            "message": None,
        }
        if qid not in by_id:
            entry.update(status="unknown", message="Question ID is not part of this questionnaire version")
        elif qid in seen:
            entry.update(status="duplicate", message="Question answered twice; the first answer is used")
        else:
            seen.add(qid)
            ans = normalise_answer(raw)
            if ans is None:
                entry.update(status="unanswered", message="No answer given")
            elif ans == "?":
                entry.update(
                    status="invalid", answer=str(raw)[:40], message="Answer must be Yes, Partial, No or N/A"
                )
            else:
                entry["answer"] = ans
                answers[qid] = ans
                q = by_id[qid]
                if ans in ("No", "Partial"):
                    found.append(
                        {
                            "text": f"{qid} {q['text']} Answer: {ans}."
                            + (f" Supplier comment: {comment}" if comment else ""),
                            "severity": "high"
                            if ans == "No" and (q.get("weight") or 1) >= 2
                            else "medium"
                            if ans == "No"
                            else "low",
                            "action_item_id": None,
                            "question_id": qid,
                        }
                    )
        report.append(entry)
    missing = [q["id"] for q in template.questions if q["id"] not in seen]
    score = score_answers(template.questions, answers)
    stored = [
        {
            "id": q["id"],
            "question": q["text"],
            "answer": answers.get(q["id"]),
            "comment": next((e["comment"] for e in report if e["id"] == q["id"]), None),
        }
        for q in template.questions
    ]
    return {
        "rows": report,
        "missing": missing,
        "answered": len(answers),
        "questions": len(template.questions),
        "score": score,
        "rating": rating_for(score),
        "findings": found,
        "answers": stored,
    }
