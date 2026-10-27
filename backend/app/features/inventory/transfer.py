"""Excel export and CSV / Excel import (column mapping, dry run, per-row report) for services and assets."""

import csv
import io
import re
import warnings

import openpyxl
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.util import next_code
from app.core.deps import Principal
from app.features.inventory.logic import (
    ASSET_TYPES,
    CIA_LEVELS,
    CRITICALITY,
    LIFECYCLE,
    SERVICE_KINDS,
    SERVICE_STATUS,
)
from app.features.inventory.models import Asset, Service, service_assets, service_suppliers
from app.features.suppliers.api import visible_supplier_ids
from app.features.suppliers.models import Supplier
from app.models import Entity, Stakeholder
from app.services.audit import audit

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
MAX_ROWS = 5000
MAX_COLS = 60
MAX_CELL = 4000

# field -> (column label used in the export, synonyms for auto-mapping, required)
ASSET_FIELDS: dict[str, tuple[str, list[str], bool]] = {
    "code": ("Code", ["asset code", "asset id", "id"], False),
    "external_id": (
        "External ID",
        ["cmdb id", "ci id", "inventory id", "serial", "asset tag", "hostname"],
        False,
    ),
    "name": ("Name", ["asset", "asset name", "system", "ci name", "device"], True),
    "entity": ("Entity", ["company", "legal entity", "entity code"], False),
    "site": ("Site", ["plant", "standort", "werk"], False),
    "asset_type": ("Type", ["asset type", "category", "class"], False),
    "description": ("Description", ["purpose", "beschreibung"], False),
    "owner": ("Owner", ["asset owner", "responsible", "verantwortlich"], False),
    "confidentiality": ("Confidentiality", ["c", "vertraulichkeit"], False),
    "integrity": ("Integrity", ["i", "integritaet", "integrität"], False),
    "availability": ("Availability", ["a", "verfuegbarkeit", "verfügbarkeit"], False),
    "criticality": ("Criticality", ["kritikalitaet", "kritikalität", "priority"], False),
    "ot_zone": ("OT zone", ["zone", "security zone"], False),
    "ot_conduit": ("OT conduit", ["conduit", "conduits"], False),
    "supplier": (
        "Supplier",
        ["vendor support", "service provider", "maintenance provider", "lieferant"],
        False,
    ),
    "lifecycle": ("Lifecycle", ["status", "lifecycle state", "state"], False),
    "location": ("Location", ["room", "rack", "region"], False),
    "vendor_product": (
        "Vendor / product",
        ["manufacturer", "vendor", "product", "model", "hersteller"],
        False,
    ),
    "crypto_algorithms": ("Crypto algorithms", ["encryption", "crypto"], False),
    "key_management_owner": ("Key management owner", ["key owner"], False),
    "tags": ("Tags", ["labels"], False),
    "services": ("Services", ["used by", "supports", "service codes"], False),
}
SERVICE_FIELDS: dict[str, tuple[str, list[str], bool]] = {
    "code": ("Code", ["service code", "service id", "id"], False),
    "name": ("Name", ["service", "service name", "process", "prozess"], True),
    "entity": ("Entity", ["company", "legal entity", "entity code"], False),
    "kind": ("Kind", ["service or process"], False),
    "description": ("Description", ["beschreibung"], False),
    "criticality": ("Criticality", ["kritikalitaet", "kritikalität"], False),
    "nis2_relevant": ("NIS2 relevant", ["nis2", "in scope"], False),
    "owner": ("Owner", ["service owner", "process owner", "responsible"], False),
    "rto_hours": ("RTO (hours)", ["rto", "recovery time objective"], False),
    "rpo_hours": ("RPO (hours)", ["rpo", "recovery point objective"], False),
    "status": ("Status", ["state"], False),
    "tags": ("Tags", ["labels"], False),
    "suppliers": ("Suppliers", ["providers", "vendors"], False),
}
FIELDS = {"assets": ASSET_FIELDS, "services": SERVICE_FIELDS}


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9äöü]+", " ", str(s or "").lower()).strip()


def _safe(ws, row: int, col: int, value):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        cell.data_type = "s"
    return cell


# --- reading files ---------------------------------------------------------------------------------------


def read_table(data: bytes, filename: str, sheet: str | None = None) -> dict:
    """Return {sheets, sheet, columns, rows} where rows are lists of cell strings (header excluded)."""
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("Could not decode the CSV file")
        sample = text[:5000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            delim = dialect.delimiter
        except csv.Error:
            delim = ";" if sample.count(";") > sample.count(",") else ","
        raw = []
        for r in csv.reader(io.StringIO(text), delimiter=delim):
            if len(r) > MAX_COLS:
                raise ValueError(f"Too many columns (maximum {MAX_COLS})")
            raw.append([c[:MAX_CELL] for c in r])
            if len(raw) > MAX_ROWS + 1:
                raise ValueError(f"Too many rows (maximum {MAX_ROWS})")
        sheets, current = ["CSV"], "CSV"
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheets = [s for s in wb.sheetnames if not s.startswith("_")]
            if not sheets:
                raise ValueError("The workbook has no sheets")
            current = sheet if sheet in sheets else sheets[0]
            ws = wb[current]
            if (ws.max_column or 0) > MAX_COLS:
                raise ValueError(f"Too many columns (maximum {MAX_COLS})")
            raw = []
            for r in ws.iter_rows(max_col=MAX_COLS, values_only=True):
                raw.append(
                    [
                        "" if v is None else (v.isoformat() if hasattr(v, "isoformat") else str(v))[:MAX_CELL]
                        for v in r
                    ]
                )
                if len(raw) > MAX_ROWS + 1:
                    raise ValueError(f"Too many rows (maximum {MAX_ROWS})")
        finally:
            wb.close()
    else:
        raise ValueError("Upload a .csv or .xlsx file")
    raw = [[c.strip() if isinstance(c, str) else c for c in r] for r in raw]
    while raw and not any(raw[0]):
        raw.pop(0)
    if not raw:
        raise ValueError("The file is empty")
    header = raw[0]
    width = max(len(header), max((len(r) for r in raw[1:]), default=0))
    columns = []
    for i in range(width):
        h = header[i] if i < len(header) and header[i] else f"Column {i + 1}"
        while h in columns:
            h = f"{h} ({i + 1})"
        columns.append(h)
    rows = [r + [""] * (width - len(r)) for r in raw[1:] if any(r)]
    if len(rows) > MAX_ROWS:
        raise ValueError(f"Too many rows (maximum {MAX_ROWS})")
    return {"sheets": sheets, "sheet": current, "columns": columns, "rows": rows}


def suggest_mapping(kind: str, columns: list[str]) -> dict[str, str | None]:
    fields = FIELDS[kind]
    normcols = {_norm(c): c for c in columns}
    out: dict[str, str | None] = {}
    taken: set[str] = set()
    for field, (label, syn, _req) in fields.items():
        hit = None
        for cand in [label, field.replace("_", " "), *syn]:
            c = normcols.get(_norm(cand))
            if c and c not in taken:
                hit = c
                break
        out[field] = hit
        if hit:
            taken.add(hit)
    return out


# --- value parsing ---------------------------------------------------------------------------------------


def _level(v: str, labels: dict[int, str]) -> int | None:
    s = _norm(v)
    if not s:
        return None
    if s.isdigit() and 1 <= int(s) <= 4:
        return int(s)
    extra = {
        "very high": 4,
        "sehr hoch": 4,
        "critical": 4,
        "kritisch": 4,
        "high": 3,
        "hoch": 3,
        "medium": 2,
        "mittel": 2,
        "normal": 2,
        "moderate": 2,
        "low": 1,
        "niedrig": 1,
        "gering": 1,
    }
    for k, lab in labels.items():
        if _norm(lab) == s:
            return k
    if s in extra:
        return extra[s]
    raise ValueError(f"'{v}' is not a level (1–4 or {', '.join(labels.values())})")


def _choice(v: str, options: dict[str, str], extra: dict[str, str] | None = None) -> str | None:
    s = _norm(v)
    if not s:
        return None
    for k, lab in options.items():
        if _norm(k) == s or _norm(lab) == s:
            return k
    if extra and s in extra:
        return extra[s]
    raise ValueError(f"'{v}' is not one of: {', '.join(options)}")


def _bool(v: str) -> bool | None:
    s = _norm(v)
    if not s:
        return None
    if s in ("yes", "y", "true", "1", "x", "ja", "j"):
        return True
    if s in ("no", "n", "false", "0", "nein"):
        return False
    raise ValueError(f"'{v}' is not yes or no")


def _hours(v: str) -> float | None:
    s = str(v or "").strip().lower().replace(",", ".")
    if not s:
        return None
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(h|hours?|std|d|days?|tage?|w|weeks?|min|minutes?)?", s)
    if not m:
        raise ValueError(f"'{v}' is not a duration in hours (e.g. 4, 24h, 3d)")
    n, unit = float(m.group(1)), (m.group(2) or "h")
    if unit.startswith("d") or unit.startswith("tag"):
        n *= 24
    elif unit.startswith("w"):
        n *= 168
    elif unit.startswith("min"):
        n /= 60
    return round(n, 2)


ASSET_TYPE_EXTRA = {
    "server": "IT",
    "client": "IT",
    "application": "IT",
    "software": "IT",
    "it system": "IT",
    "ot system": "OT",
    "plc": "OT",
    "scada": "OT",
    "ics": "OT",
    "saas": "cloud",
    "iaas": "cloud",
    "paas": "cloud",
    "cloud service": "cloud",
    "database": "data",
    "information": "data",
    "building": "facility",
    "room": "facility",
    "firewall": "network",
    "switch": "network",
    "router": "network",
}
LIFECYCLE_EXTRA = {
    "operational": "in_operation",
    "active": "in_operation",
    "in use": "in_operation",
    "production": "in_operation",
    "eol": "end_of_life",
    "decommissioned": "retired",
    "disposed": "retired",
    "planning": "planned",
}


# --- import ----------------------------------------------------------------------------------------------


def run_import(
    db: Session, pr: Principal, kind: str, table: dict, mapping: dict[str, str | None], dry_run: bool
) -> dict:
    fields = FIELDS[kind]
    col_idx = {c: i for i, c in enumerate(table["columns"])}
    for f, col in mapping.items():
        if f not in fields:
            raise ValueError(f"Unknown field '{f}'")
        if col and col not in col_idx:
            raise ValueError(f"Column '{col}' not found in the file")
    for f, (label, _s, req) in fields.items():
        if req and not mapping.get(f):
            raise ValueError(f"Map a column to the required field '{label}'")

    ents = {e.code.lower(): e for e in db.scalars(select(Entity))}
    ents.update({e.name.lower(): e for e in ents.values()})
    stakeholders = {s.name.lower(): s for s in db.scalars(select(Stakeholder))}
    suppliers = {}
    # only suppliers the caller may see, otherwise the import would confirm names of other entities
    visible = visible_supplier_ids(db, pr)
    for s in db.scalars(select(Supplier).where(Supplier.id.in_(visible))):
        suppliers[s.name.lower()] = s
        suppliers[s.code.lower()] = s
    model = Asset if kind == "assets" else Service
    # match only records the caller may see, so the report does not reveal other entities' records
    existing = [o for o in db.scalars(select(model)) if pr.can_read("assets.read", o.entity_id)]
    by_code = {o.code.lower(): o for o in existing}
    by_ext = {o.external_id.lower(): o for o in existing if kind == "assets" and o.external_id}
    by_name = {(o.entity_id, o.name.lower()): o for o in existing}
    services_lookup = {}
    if kind == "assets":
        for s in db.scalars(select(Service)):
            if not pr.can_read("assets.read", s.entity_id):
                continue
            services_lookup[s.code.lower()] = s
            services_lookup.setdefault(s.name.lower(), s)
    perm = "assets.manage"
    counts = {"created": 0, "updated": 0, "unchanged": 0, "errors": 0, "warnings": 0}
    report = []
    seen_keys: set = set()
    prefix = "AST-" if kind == "assets" else "SVC-"

    def cell(row, field):
        col = mapping.get(field)
        return str(row[col_idx[col]] or "").strip() if col else ""

    names = {
        "owner_id": {s.id: s.name for s in stakeholders.values()},
        "supplier_id": {s.id: s.name for s in suppliers.values()},
        "entity_id": {e.id: e.code for e in ents.values()},
    }
    labels = {"owner_id": "Owner", "supplier_id": "Supplier", "entity_id": "Entity"}

    def show(k, v):
        return _fmt(names[k].get(v, v) if k in names and v is not None else v)

    for n, row in enumerate(table["rows"], start=2):
        errors, warns = [], []
        vals: dict = {}
        name = cell(row, "name")
        if not name:
            report.append({"row": n, "action": "error", "key": None, "messages": ["Name is empty"]})
            counts["errors"] += 1
            continue
        vals["name"] = name[:300]
        # entity
        ent_raw = cell(row, "entity")
        entity_id = None
        if ent_raw:
            e = ents.get(ent_raw.lower())
            if not e:
                errors.append(f"Unknown entity '{ent_raw}'")
            else:
                entity_id = e.id
        # find existing
        code = cell(row, "code")
        ext = cell(row, "external_id") if kind == "assets" else ""
        obj = None
        if code:
            obj = by_code.get(code.lower())
            if not obj:
                warns.append(f"Code {code} not found; a new record gets a new code")
        if not obj and ext:
            obj = by_ext.get(ext.lower())
        if not obj:
            obj = by_name.get((entity_id if ent_raw else None, name.lower()))
            if not obj and not ent_raw:
                obj = next((o for (eid, nm), o in by_name.items() if nm == name.lower()), None)
        if ent_raw or not obj:
            vals["entity_id"] = entity_id
        target_entity = vals.get("entity_id", obj.entity_id if obj else None)
        if not pr.can_write(perm, target_entity) or (obj and not pr.can_write(perm, obj.entity_id)):
            errors.append(
                "No permission to maintain inventory records for this entity"
                if target_entity
                else "Group-wide records need group-wide inventory rights"
            )
        key = (obj.id if obj else None, target_entity, name.lower())
        if key in seen_keys:
            errors.append("Duplicate row for the same record in this file")
        seen_keys.add(key)

        def parse(field, fn, row=row, errors=errors, vals=vals):
            raw = cell(row, field)
            if not mapping.get(field):
                return
            try:
                v = fn(raw)
            except ValueError as ex:
                errors.append(f"{fields[field][0]}: {ex}")
                return
            if v is None and not raw:
                return
            vals[field] = v

        simple = (
            [
                "description",
                "site",
                "ot_zone",
                "ot_conduit",
                "location",
                "vendor_product",
                "crypto_algorithms",
                "key_management_owner",
                "external_id",
            ]
            if kind == "assets"
            else ["description"]
        )
        for f in simple:
            if mapping.get(f) and cell(row, f):
                vals[f] = cell(row, f)[:300] if f != "description" else cell(row, f)
        if mapping.get("tags") and cell(row, "tags"):
            vals["tags"] = [t.strip()[:40] for t in re.split(r"[;,]", cell(row, "tags")) if t.strip()][:20]
        if mapping.get("owner") and cell(row, "owner"):
            st = stakeholders.get(cell(row, "owner").lower())
            if st:
                vals["owner_id"] = st.id
            else:
                warns.append(f"Owner '{cell(row, 'owner')}' is not a known stakeholder; owner not set")
        if kind == "assets":
            parse("asset_type", lambda v: _choice(v, ASSET_TYPES, ASSET_TYPE_EXTRA))
            parse("lifecycle", lambda v: _choice(v, LIFECYCLE, LIFECYCLE_EXTRA))
            for f in ("confidentiality", "integrity", "availability"):
                parse(f, lambda v: _level(v, CIA_LEVELS))
            parse("criticality", lambda v: _level(v, CRITICALITY))
            if mapping.get("supplier") and cell(row, "supplier"):
                sp = suppliers.get(cell(row, "supplier").lower())
                if sp:
                    vals["supplier_id"] = sp.id
                else:
                    warns.append(f"Supplier '{cell(row, 'supplier')}' not found in the supplier register")
            link_services = []
            if mapping.get("services") and cell(row, "services"):
                for token in re.split(r"[;,\n]", cell(row, "services")):
                    t = token.strip()
                    if not t:
                        continue
                    s = services_lookup.get(t.lower())
                    if s and pr.can_write(perm, s.entity_id):
                        link_services.append(s)
                    else:
                        warns.append(f"Service '{t}' not found or not editable; dependency skipped")
        else:
            parse("kind", lambda v: _choice(v, SERVICE_KINDS))
            parse("status", lambda v: _choice(v, SERVICE_STATUS, {"operational": "active", "live": "active"}))
            parse("criticality", lambda v: _level(v, CRITICALITY))
            parse("nis2_relevant", _bool)
            parse("rto_hours", _hours)
            parse("rpo_hours", _hours)
            link_suppliers = []
            if mapping.get("suppliers") and cell(row, "suppliers"):
                for token in re.split(r"[;,\n]", cell(row, "suppliers")):
                    t = token.strip()
                    if not t:
                        continue
                    sp = suppliers.get(t.lower())
                    if sp:
                        link_suppliers.append(sp)
                    else:
                        warns.append(f"Supplier '{t}' not found; dependency skipped")

        if errors:
            counts["errors"] += 1
            report.append(
                {"row": n, "action": "error", "key": obj.code if obj else name, "messages": errors + warns}
            )
            continue
        if warns:
            counts["warnings"] += 1
        if obj:
            changes = {k: [getattr(obj, k), v] for k, v in vals.items() if getattr(obj, k) != v}
            links_new = []
            if kind == "assets" and link_services:
                have = set(
                    db.scalars(select(service_assets.c.service_id).where(service_assets.c.asset_id == obj.id))
                )
                links_new = [s for s in link_services if s.id not in have]
            elif kind == "services" and link_suppliers:
                have = set(
                    db.scalars(
                        select(service_suppliers.c.supplier_id).where(
                            service_suppliers.c.service_id == obj.id
                        )
                    )
                )
                links_new = [s for s in link_suppliers if s.id not in have]
            if not changes and not links_new:
                counts["unchanged"] += 1
                report.append({"row": n, "action": "unchanged", "key": obj.code, "messages": warns})
                continue
            counts["updated"] += 1
            msg = [
                f"{labels.get(k) or (fields[k][0] if k in fields else k)}: {show(k, v[0])} → {show(k, v[1])}"
                for k, v in changes.items()
            ]
            msg += [f"New dependency on {s.code}" for s in links_new]
            report.append({"row": n, "action": "update", "key": obj.code, "messages": msg + warns})
            if not dry_run:
                for k, v in vals.items():
                    setattr(obj, k, v)
                _link(db, kind, obj, links_new)
                audit(
                    db,
                    pr,
                    "import",
                    kind[:-1],
                    obj.id,
                    f"{obj.code} {obj.name}",
                    entity_id=obj.entity_id,
                    details={k: [_plain(v[0]), _plain(v[1])] for k, v in changes.items()},
                )
        else:
            counts["created"] += 1
            if dry_run:
                code_new = "(new)"
            else:
                code_new = next_code(db, model.code, prefix)
            links = link_services if kind == "assets" else link_suppliers
            report.append(
                {
                    "row": n,
                    "action": "create",
                    "key": code_new,
                    "messages": [f"New {kind[:-1]} '{name}'"]
                    + [f"Dependency on {x.code}" for x in links]
                    + warns,
                }
            )
            if not dry_run:
                obj = model(code=code_new, **vals)
                if kind == "assets":
                    obj.asset_type = vals.get("asset_type") or "IT"
                    obj.lifecycle = vals.get("lifecycle") or "in_operation"
                db.add(obj)
                db.flush()
                by_name[(obj.entity_id, obj.name.lower())] = obj
                _link(db, kind, obj, links)
                audit(db, pr, "import", kind[:-1], obj.id, f"{obj.code} {obj.name}", entity_id=obj.entity_id)
    if not dry_run:
        db.commit()
    return {"dry_run": dry_run, "kind": kind, "total": len(table["rows"]), **counts, "rows": report}


def _link(db: Session, kind: str, obj, targets: list) -> None:
    from sqlalchemy import insert

    if not targets:
        return
    if kind == "assets":
        db.execute(insert(service_assets), [{"service_id": s.id, "asset_id": obj.id} for s in targets])
    else:
        db.execute(insert(service_suppliers), [{"service_id": obj.id, "supplier_id": s.id} for s in targets])


def _fmt(v) -> str:
    if v is None or v == "" or v == []:
        return "empty"
    if isinstance(v, list):
        return ", ".join(map(str, v))
    return str(v)


def _plain(v):
    return v if isinstance(v, (str, int, float, bool, list)) or v is None else str(v)


# --- export ----------------------------------------------------------------------------------------------


def export_workbook(db: Session, pr: Principal, services: list[Service], assets: list[Asset]) -> bytes:
    ents = {e.id: e.code for e in db.scalars(select(Entity))}
    sts = {s.id: s.name for s in db.scalars(select(Stakeholder))}
    # hidden suppliers are left out: empty cell, no dependency row
    visible = visible_supplier_ids(db, pr)
    sups = {s.id: s for s in db.scalars(select(Supplier).where(Supplier.id.in_(visible)))}
    svc_ids = {s.id for s in services}
    ast_ids = {a.id for a in assets}
    svc_assets: dict[int, list[int]] = {}
    for sid, aid in db.execute(select(service_assets.c.service_id, service_assets.c.asset_id)):
        svc_assets.setdefault(sid, []).append(aid)
    svc_sups: dict[int, list[int]] = {}
    for sid, supid in db.execute(select(service_suppliers.c.service_id, service_suppliers.c.supplier_id)):
        svc_sups.setdefault(sid, []).append(supid)
    svc_by_id = {s.id: s for s in services}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Services"
    _head(ws, [lab for lab, _s, _r in SERVICE_FIELDS.values()])
    for r, s in enumerate(services, start=2):
        vals = [
            s.code,
            s.name,
            ents.get(s.entity_id),
            SERVICE_KINDS.get(s.kind),
            s.description,
            CRITICALITY.get(s.criticality),
            "yes" if s.nis2_relevant else "no",
            sts.get(s.owner_id),
            s.rto_hours,
            s.rpo_hours,
            SERVICE_STATUS.get(s.status),
            ", ".join(s.tags or []),
            "; ".join(sups[i].code for i in svc_sups.get(s.id, []) if i in sups),
        ]
        for c, v in enumerate(vals, start=1):
            _safe(ws, r, c, v)
    wa = wb.create_sheet("Assets")
    _head(wa, [lab for lab, _s, _r in ASSET_FIELDS.values()])
    asset_services: dict[int, list[str]] = {}
    for sid, aids in svc_assets.items():
        for aid in aids:
            if sid in svc_by_id:
                asset_services.setdefault(aid, []).append(svc_by_id[sid].code)
    for r, a in enumerate(assets, start=2):
        vals = [
            a.code,
            a.external_id,
            a.name,
            ents.get(a.entity_id),
            a.site,
            a.asset_type,
            a.description,
            sts.get(a.owner_id),
            CIA_LEVELS.get(a.confidentiality),
            CIA_LEVELS.get(a.integrity),
            CIA_LEVELS.get(a.availability),
            CRITICALITY.get(a.criticality),
            a.ot_zone,
            a.ot_conduit,
            sups[a.supplier_id].name if a.supplier_id in sups else None,
            LIFECYCLE.get(a.lifecycle),
            a.location,
            a.vendor_product,
            a.crypto_algorithms,
            a.key_management_owner,
            ", ".join(a.tags or []),
            "; ".join(sorted(asset_services.get(a.id, []))),
        ]
        for c, v in enumerate(vals, start=1):
            _safe(wa, r, c, v)
    wd = wb.create_sheet("Dependencies")
    _head(wd, ["Service code", "Service", "Entity", "Depends on", "Code", "Name", "Supplier"])
    r = 2
    for s in services:
        for aid in svc_assets.get(s.id, []):
            if aid not in ast_ids:
                continue
            a = next(x for x in assets if x.id == aid)
            for c, v in enumerate(
                [
                    s.code,
                    s.name,
                    ents.get(s.entity_id),
                    "Asset",
                    a.code,
                    a.name,
                    sups[a.supplier_id].name if a.supplier_id in sups else None,
                ],
                start=1,
            ):
                _safe(wd, r, c, v)
            r += 1
        for supid in svc_sups.get(s.id, []):
            if supid in sups and s.id in svc_ids:
                sp = sups[supid]
                for c, v in enumerate(
                    [s.code, s.name, ents.get(s.entity_id), "Supplier", sp.code, sp.name, sp.name], start=1
                ):
                    _safe(wd, r, c, v)
                r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _head(ws, labels: list[str]) -> None:
    for c, h in enumerate(labels, start=1):
        cell = _safe(ws, 1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D6770")
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = max(12, min(40, len(h) + 6))
    ws.freeze_panes = "A2"


def count_rows(db: Session, model) -> int:
    return db.scalar(select(func.count(model.id))) or 0
