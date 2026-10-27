"""Plain-text extraction from uploaded files for search and AI retrieval."""

import csv
import email
import io
import json
import logging
import zipfile
from email import policy

log = logging.getLogger(__name__)
MAX_TEXT = 2_000_000  # characters kept per version


MAX_UNCOMPRESSED = 300 * 1024 * 1024
MAX_RATIO = 150


def zip_is_safe(data: bytes) -> bool:
    """Reject zip containers (docx/xlsx/pptx/odf/zip) that would decompress to excessive sizes."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            total = sum(i.file_size for i in infos)
            if total > MAX_UNCOMPRESSED or len(infos) > 10000:
                return False
            for i in infos:
                if (
                    i.compress_size
                    and i.file_size / i.compress_size > MAX_RATIO
                    and i.file_size > 10 * 1024 * 1024
                ):
                    return False
    except zipfile.BadZipFile:
        return False
    return True


def _pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            t = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - damaged page, keep going
            t = ""
        if t.strip():
            pages.append(f"[Page {i + 1}]\n{t}")
    return "\n\n".join(pages), len(reader.pages)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            txt = "".join(n.text or "" for n in block.iter() if n.tag.endswith("}t"))
            if txt.strip():
                parts.append(txt)
        elif tag == "tbl":
            for row in block.iter():
                if row.tag.endswith("}tr"):
                    cells = []
                    for cell in row.iter():
                        if cell.tag.endswith("}tc"):
                            cells.append(
                                "".join(n.text or "" for n in cell.iter() if n.tag.endswith("}t")).strip()
                            )
                    if any(cells):
                        parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            texts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        if any(t.strip() for t in texts):
            out.append(f"[Slide {i}]\n" + "\n".join(texts))
    return "\n\n".join(out), len(prs.slides)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[Sheet {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None and str(v).strip()]
            if vals:
                out.append(" | ".join(vals))
            if sum(len(x) for x in out) > MAX_TEXT:
                break
    return "\n".join(out)


def _html(data: bytes) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    for t in soup(["script", "style"]):
        t.decompose()
    return soup.get_text("\n")


def _odf(data: bytes) -> str:
    from bs4 import BeautifulSoup

    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("content.xml")
    return BeautifulSoup(xml, "html.parser").get_text("\n")


def _eml(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=policy.default)
    head = f"From: {msg['from']}\nTo: {msg['to']}\nDate: {msg['date']}\nSubject: {msg['subject']}\n\n"
    body = msg.get_body(preferencelist=("plain", "html"))
    content = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        content = _html(content.encode())
    return head + content


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> tuple[str | None, int | None, str]:
    """Returns (text, page_count, status) with status ok | empty | failed | unsupported."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    pages = None
    if ext in ("docx", "pptx", "xlsx", "xlsm", "odt", "ods", "odp") and not zip_is_safe(data):
        return None, None, "failed"
    try:
        if ext == "pdf":
            text, pages = _pdf(data)
        elif ext == "docx":
            text = _docx(data)
        elif ext == "pptx":
            text, pages = _pptx(data)
        elif ext in ("xlsx", "xlsm"):
            text = _xlsx(data)
        elif ext in ("html", "htm"):
            text = _html(data)
        elif ext in ("odt", "ods", "odp"):
            text = _odf(data)
        elif ext == "eml":
            text = _eml(data)
        elif ext == "json":
            text = json.dumps(json.loads(_decode(data)), indent=1, ensure_ascii=False)
        elif ext == "csv":
            rows = csv.reader(io.StringIO(_decode(data)))
            text = "\n".join(" | ".join(r) for r in rows)
        elif ext in ("txt", "md", "xml", "log", "yaml", "yml"):
            text = _decode(data)
        else:
            return None, None, "unsupported"
    except Exception as e:  # noqa: BLE001
        log.warning("text extraction failed for %s: %s", filename, e)
        return None, None, "failed"
    text = "\n".join(line.rstrip() for line in (text or "").splitlines())
    text = text.strip()[:MAX_TEXT]
    return (text or None), pages, ("ok" if text else "empty")
