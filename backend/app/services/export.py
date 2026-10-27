"""Export into the original Excel templates (formulas, formatting and drop-downs stay intact) and import of
filled-in workbooks."""

import io
import logging
import re
import warnings
from copy import copy
from datetime import date

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    ActionItem,
    Activity,
    Answer,
    Assessment,
    Entity,
    Gate,
    Project,
    RaciActivity,
    RaciCell,
    RaciColumn,
    RaidRecord,
    Requirement,
    Stakeholder,
    WorkPackage,
)
from app.services.scoring import recalc
from app.services.workbook import GAP_TEMPLATE, MAX_ENTITIES, PROJECT_TEMPLATE, parse_gap_workbook

log = logging.getLogger(__name__)


_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


class _SafeSheet:
    """Wraps a worksheet so user-supplied strings are never written as formulas (CSV/Excel injection)."""

    def __init__(self, ws):
        self.ws = ws

    def __setitem__(self, key, value):
        cell = self.ws[key]
        if isinstance(value, str) and value.startswith(_FORMULA_START):
            cell.value = value
            cell.data_type = "s"
        else:
            cell.value = value

    def __getitem__(self, key):
        return self.ws[key]

    def __getattr__(self, name):
        return getattr(self.ws, name)


def _load(path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path)


_BAD_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")
_SLOT_ROWS = 4  # entity rows the template Dashboard is laid out for
_CHECK = '=IFERROR(IF(COUNTA(INDIRECT("\'"&$O{r}&"\'!$A$5"))>0,"OK","Sheet not found"),"Sheet not found")'


def _sheet_title(code: str) -> str:
    return _BAD_TITLE_CHARS.sub("_", code)[:31]


def _copy_requirements(wb, src, title: str):
    """A copy of the catalogue sheet with everything copy_worksheet leaves behind."""
    ws = wb.copy_worksheet(src)
    ws.title = title
    ws.freeze_panes = src.freeze_panes
    ws.auto_filter.ref = src.auto_filter.ref
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines
    ws.sheet_view.zoomScale = src.sheet_view.zoomScale
    ws.protection = copy(src.protection)
    for dv in src.data_validations.dataValidation:
        ws.add_data_validation(copy(dv))
    for cf in src.conditional_formatting:
        for rule in cf.rules:
            ws.conditional_formatting.add(str(cf.sqref), copy(rule))
    return ws


def _fill_lists(wb, lists, entities: list[Entity], titles: list[str]) -> None:
    safe = _SafeSheet(lists)
    for i, (ent, title) in enumerate(zip(entities, titles, strict=True)):
        r = 6 + i
        for col, value in (
            ("F", ent.code),
            ("O", title),
            ("P", ent.finding_prefix),
        ):
            if r > 6:
                lists[f"{col}{r}"]._style = copy(lists[f"{col}6"]._style)
            safe[f"{col}{r}"] = value
        if r > 6:
            lists[f"Q{r}"]._style = copy(lists["Q6"]._style)
        lists[f"Q{r}"] = _CHECK.format(r=r)  # our own formula, so not through the string-forcing wrapper
    wb.defined_names["EntityList"].attr_text = f"Lists!$F$6:$F${5 + max(len(entities), 1)}"


def _ind(i: int, col: str) -> str:
    """The cells of one column on the sheet of the i-th entity, found through Lists!O."""
    return f'INDIRECT("\'"&Lists!$O${6 + i}&"\'!${col}$6:${col}$179")'


def _rebuild_dashboard(wb, n_entities: int) -> None:
    """Re-lay the template Dashboard (laid out for four entities) for any number of entities and point its
    formulas at the entity sheets the export created. Styles come from the template cells."""
    ws = wb["Dashboard"]
    k = max(n_entities, 1)
    d = k - _SLOT_ROWS
    snap = {
        (c.row, c.column): (c.value, copy(c._style))
        for row in ws.iter_rows(min_row=8, max_row=54, max_col=9)
        for c in row
    }
    heights = {r: ws.row_dimensions[r].height for r in range(8, 55)}
    old_formats = list(ws.conditional_formatting)
    for row in ws.iter_rows(min_row=8, max_row=60 + 2 * abs(d), max_col=10 + abs(d)):
        for c in row:
            c.value = None
            c._style = StyleArray()
        ws.row_dimensions[row[0].row].height = None

    def put(row: int, col: int, value, src_row: int, src_col: int | None = None) -> None:
        cell = ws.cell(row, col)
        cell._style = copy(snap[(src_row, src_col or col)][1])
        cell.value = value
        ws.row_dimensions[row].height = heights[src_row]

    def slot(i: int) -> int:  # entity rows beyond the template's four repeat the last one's style
        return min(i, _SLOT_ROWS - 1)

    # coverage by entity
    for r in (8, 9):
        for c in range(1, 9):
            put(r, c, snap[(r, c)][0], r)
    last = 9 + k
    for i in range(k):
        r = 10 + i
        values = [None] * 8
        if i < n_entities:
            values = [
                f"=Lists!$F${6 + i}",
                f'=COUNTIF({_ind(i, "I")},"In Scope")',
                f'=COUNTIFS({_ind(i, "I")},"In Scope",{_ind(i, "J")},"?*")',
                f'=IFERROR($C{r}/$B{r},"")',
                f'=IFERROR(AVERAGE({_ind(i, "N")}),"")',
                f'=COUNTIF({_ind(i, "O")},"Yes")',
                f'=COUNTIF({_ind(i, "J")},"No")+COUNTIF({_ind(i, "K")},"No")',
                f'=COUNTIF({_ind(i, "L")},"Not Tested")',
            ]
        for c, v in enumerate(values, 1):
            put(r, c, v, 10 + slot(i))
    r = last + 1
    totals = [
        snap[(14, 1)][0],
        f"=SUM($B$10:$B${last})",
        f"=SUM($C$10:$C${last})",
        f'=IFERROR($C{r}/$B{r},"")',
        f'=IFERROR(AVERAGE($E$10:$E${last}),"")',
        *(f"=SUM(${col}$10:${col}${last})" for col in "FGH"),
    ]
    for c, v in enumerate(totals, 1):
        put(r, c, v, 14)

    # average score per chapter: one column per entity, then group average, gap to target, work packages
    put(16 + d, 1, snap[(16, 1)][0], 16)
    hdr, first, last_ch = 17 + d, 18 + d, 34 + d
    group, gap, wp = 3 + k, 4 + k, 5 + k
    columns = [(1, 1), (2, 2), *((3 + i, 3 + slot(i)) for i in range(k)), (group, 7), (gap, 8), (wp, 9)]
    ent_cols = f"$C{{r}}:${get_column_letter(2 + k)}{{r}}"
    g = get_column_letter(group)
    for c, src in columns:
        label = snap[(17, src)][0] if c in (1, 2, group, gap, wp) else None
        if 3 <= c < 3 + n_entities:
            label = f"=Lists!$F${6 + c - 3}"
        put(hdr, c, label, 17, src)
        for r in range(first, last_ch + 1):
            put(r, c, None, r - d, src)
    for r in range(first, last_ch + 1):
        ws.cell(r, 1).value, ws.cell(r, 2).value = snap[(r - d, 1)][0], snap[(r - d, 2)][0]
        for i in range(n_entities):
            ws.cell(r, 3 + i).value = f'=IFERROR(AVERAGEIFS({_ind(i, "N")},{_ind(i, "B")},$A{r}),"")'
        ws.cell(r, group).value = f'=IFERROR(AVERAGE({ent_cols.format(r=r)}),"")'
        ws.cell(r, gap).value = f'=IF(${g}{r}="","",MAX(0,$B{r}-${g}{r}))'
        counts = "+".join(f'COUNTIFS({_ind(i, "O")},"Yes",{_ind(i, "B")},$A{r})' for i in range(n_entities))
        ws.cell(r, wp).value = f"={counts}" if counts else 0

    # answer distribution
    put(36 + d, 1, snap[(36, 1)][0], 36)
    for c in range(1, 6):
        put(37 + d, c, snap[(37, c)][0], 37)
    dist_first = 38 + d
    for i in range(k):
        r = dist_first + i
        values = [None] * 5
        if i < n_entities:
            values = [
                f"=Lists!$F${6 + i}",
                *(
                    f'=COUNTIFS({_ind(i, "I")},"In Scope",{_ind(i, "J")},"{answer}")'
                    for answer in ("Yes", "Partial", "No")
                ),
                f'=COUNTIF({_ind(i, "I")},"In Scope")-SUM($B{r}:$D{r})',
            ]
        for c, v in enumerate(values, 1):
            put(r, c, v, 38 + slot(i))
    dist_last = dist_first + k - 1

    # action plan indicators and scoring model, moved below the tables
    for (row, col), (value, _) in snap.items():
        if row >= 44:
            if col == 2 and isinstance(value, str) and value.startswith("="):
                value = _action_plan_formula(row, n_entities)
            put(row + 2 * d, col, value, row)

    ws.conditional_formatting = ConditionalFormattingList()
    for cf in old_formats:
        target = (
            f"C{first}:{g}{last_ch}"
            if str(cf.sqref).startswith("C")
            else f"{get_column_letter(gap)}{first}:{get_column_letter(gap)}{last_ch}"
        )
        for rule in cf.rules:
            ws.conditional_formatting.add(target, rule)
    for col in range(10, wp + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    _repoint_charts(ws, d, k, (hdr, first, last_ch), (37 + d, dist_first, dist_last), (group, gap))


def _action_plan_formula(row: int, n_entities: int) -> str:
    """Dashboard indicators over the Action Plan sheet; the template's column letters were one off."""
    plan = "'Action Plan'!"
    planned = f'{plan}$H$6:$H$205,"?*"'
    if row == 46:
        return f"=COUNTIF({planned})"
    if row in (47, 48, 49):
        priority = {47: "High", 48: "Medium", 49: "Low"}[row]
        return f'=COUNTIFS({planned},{plan}$I$6:$I$205,"{priority}")'
    if row in (50, 51, 52):
        status = {50: "Not Started", 51: "In Progress", 52: "Completed"}[row]
        return f'=COUNTIFS({planned},{plan}$N$6:$N$205,"{status}")'
    if row == 53:
        return f'=COUNTIF({plan}$O$6:$O$205,"Overdue")'
    required = "+".join(f'COUNTIF({_ind(i, "O")},"Yes")' for i in range(n_entities)) or "0"
    return f"={required}-COUNTIF({planned})"


def _repoint_charts(ws, d, k, chapters, answers, columns) -> None:
    radar, gap_chart, dist_chart = ws._charts
    head, first, last = chapters
    group, gap = columns

    def point(series, idx, col, rows, cat_rows):
        letter = get_column_letter(col)
        series.idx = series.order = idx
        series.val.numRef.f = f"Dashboard!${letter}${rows[1]}:${letter}${rows[2]}"
        series.val.numRef.numCache = None
        series.cat.strRef.f = f"Dashboard!$A${cat_rows[0]}:$A${cat_rows[1]}"
        series.cat.strRef.strCache = None
        series.tx.strRef.f = f"Dashboard!${letter}${rows[0]}"
        series.tx.strRef.strCache = None

    old = radar.series
    radar.series = [old[0], *(copy(old[1 + i % _SLOT_ROWS]) for i in range(k)), old[5]]
    for i, (series, col) in enumerate(zip(radar.series, [2, *range(3, 3 + k), group], strict=True)):
        point(series, i, col, chapters, (first, last))
    point(gap_chart.series[0], 0, gap, chapters, (first, last))
    for i, series in enumerate(dist_chart.series):
        point(series, i, 2 + i, answers, (answers[1], answers[2]))
    for chart in (radar, gap_chart, dist_chart):
        chart.anchor._from.col += d
        chart.anchor.to.col += d


def _apply_answers(ws, answers: dict) -> None:
    for r in range(6, ws.max_row + 1):
        num = ws[f"A{r}"].value
        a = answers.get(str(num)) if num else None
        if not a:
            continue
        ws[f"I{r}"] = a.scope
        ws[f"J{r}"] = a.implemented
        ws[f"K{r}"] = a.documented
        ws[f"L{r}"] = a.effectiveness
        ws[f"M{r}"] = a.evidence
        ws[f"P{r}"] = a.gap_description
        ws[f"Q{r}"] = a.recommended_action
        ws[f"R{r}"] = a.responsible.name if a.responsible else None
        ws[f"S{r}"] = a.target_date
        ws[f"T{r}"] = a.effort


def export_gap(db: Session, round_id: int, entity_ids: set[int] | None) -> bytes:
    """The template's `Requirements` sheet stays the catalogue; every entity gets its own copy with its
    answers, named by the entity code. Lists and the Dashboard are filled for exactly these entities."""
    wb = _load(GAP_TEMPLATE)
    lists = wb["Lists"]
    catalogue = wb["Requirements"]
    entities = [
        e
        for e in db.scalars(select(Entity).order_by(Entity.sort_order))
        if entity_ids is None or e.id in entity_ids
    ][:MAX_ENTITIES]
    names = [
        s.name
        for s in db.scalars(
            select(Stakeholder)
            .where(Stakeholder.is_active.is_(True))
            .order_by(Stakeholder.sort_order, Stakeholder.name)
        )
    ]
    safe_lists = _SafeSheet(lists)
    for i in range(6, 40):
        safe_lists[f"S{i}"] = names[i - 6] if i - 6 < len(names) else None
    titles = []
    for idx, ent in enumerate(entities):
        row = 6 + idx
        ws = _copy_requirements(wb, catalogue, _sheet_title(ent.code))
        titles.append(ws.title)
        ws["A2"].value = f"=Lists!$F${row}"
        for r in range(6, ws.max_row + 1):
            cell = ws[f"U{r}"]
            if isinstance(cell.value, str):
                cell.value = cell.value.replace("Lists!$P$6", f"Lists!$P${row}")
        asm = db.scalar(
            select(Assessment).where(Assessment.round_id == round_id, Assessment.entity_id == ent.id)
        )
        answers = (
            {
                a.requirement.number: a
                for a in db.scalars(select(Answer).where(Answer.assessment_id == asm.id))
            }
            if asm
            else {}
        )
        _apply_answers(_SafeSheet(ws), answers)
    _fill_lists(wb, lists, entities, titles)
    _rebuild_dashboard(wb, len(entities))
    # entity sheets follow the catalogue
    order = ["Dashboard", "Lists", "Requirements", *titles, "Action Plan", "Instructions"]
    wb._sheets = [wb[name] for name in order]
    ap = _SafeSheet(wb["Action Plan"])
    items = [
        i
        for i in db.scalars(select(ActionItem).order_by(ActionItem.code))
        if entity_ids is None or i.entity_id in entity_ids
    ]
    for n, it in enumerate(items):
        r = 6 + n
        if r > ap.max_row:
            break
        ap[f"B{r}"] = it.entity.code
        ap[f"C{r}"] = it.requirement.number
        ap[f"H{r}"] = it.title
        ap[f"I{r}"] = it.priority
        ap[f"J{r}"] = it.acceptance_criterion
        ap[f"K{r}"] = it.expected_evidence
        ap[f"L{r}"] = it.responsible.name if it.responsible else None
        ap[f"M{r}"] = it.target_date
        ap[f"N{r}"] = it.status
        ap[f"P{r}"] = it.comments
    for r in range(6 + len(items), ap.max_row + 1):
        for col in "BCHIJKLMNP":
            ap[f"{col}{r}"] = None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_project(db: Session, project_id: int) -> bytes:
    wb = _load(PROJECT_TEMPLATE)
    p = db.get(Project, project_id)
    ov = _SafeSheet(wb["Project Overview"])
    fields = {
        "Project Name": p.name,
        "Planned Start": p.planned_start,
        "Planned Completion": p.planned_finish,
        "Project Sponsor": p.sponsor,
        **{f["field"]: f["value"] for f in (p.overview or [])},
    }
    for r in range(6, ov.max_row + 1):
        f = ov[f"A{r}"].value
        if f in fields:
            ov[f"B{r}"] = fields[f]
    mp = _SafeSheet(wb["Master Plan"])
    acts = db.scalars(select(Activity).where(Activity.project_id == p.id).order_by(Activity.sort_order)).all()
    for n in range(6, mp.max_row + 1):
        for col in "ABCDEFGIJKMNO":
            mp[f"{col}{n}"] = None
    for n, a in enumerate(acts):
        r = 6 + n
        mp[f"A{r}"] = a.code
        mp[f"B{r}"] = a.phase.name if a.phase else None
        mp[f"C{r}"] = a.title
        mp[f"D{r}"] = a.owner.name if a.owner else None
        mp[f"E{r}"] = a.support
        mp[f"F{r}"] = a.planned_start
        mp[f"G{r}"] = a.planned_finish
        mp[f"I{r}"] = "; ".join(d.code for d in a.depends_on) or "None"
        mp[f"J{r}"] = a.deliverable
        mp[f"K{r}"] = a.gate.code if a.gate else None
        mp[f"M{r}"] = a.status
        mp[f"N{r}"] = a.priority
        mp[f"O{r}"] = a.progress
    gg = _SafeSheet(wb["Governance and Gates"])
    for n, g in enumerate(
        db.scalars(select(Gate).where(Gate.project_id == p.id).order_by(Gate.sort_order)).all()[:6]
    ):
        r = 6 + n
        gg[f"A{r}"] = g.code
        gg[f"B{r}"] = g.planned_date
        gg[f"C{r}"] = g.decision
        gg[f"D{r}"] = g.decision_owner.name if g.decision_owner else None
        gg[f"E{r}"] = g.required_input
        gg[f"F{r}"] = g.approval_criterion
        gg[f"G{r}"] = g.status
        gg[f"H{r}"] = g.decision_reference
    cols = db.scalars(
        select(RaciColumn).where(RaciColumn.project_id == p.id).order_by(RaciColumn.sort_order)
    ).all()[:11]
    rows = db.scalars(
        select(RaciActivity).where(RaciActivity.project_id == p.id).order_by(RaciActivity.sort_order)
    ).all()[:14]
    cells = {(c.activity_id, c.column_id): c.value for c in db.scalars(select(RaciCell))}
    letters = "BCDEFGHIJKL"
    for i, c in enumerate(cols):
        gg[f"{letters[i]}15"] = c.name
    for j, ra in enumerate(rows):
        gg[f"A{16 + j}"] = ra.name
        for i, c in enumerate(cols):
            gg[f"{letters[i]}{16 + j}"] = cells.get((ra.id, c.id))
    wp = _SafeSheet(wb["Work Packages"])
    for n, w in enumerate(
        db.scalars(
            select(WorkPackage).where(WorkPackage.project_id == p.id).order_by(WorkPackage.code)
        ).all()[:40]
    ):
        r = 6 + n
        wp[f"B{r}"] = ", ".join(f.code for f in w.findings) or None
        wp[f"C{r}"] = w.entity.name if w.entity else "Group / All Entities"
        wp[f"D{r}"] = w.measure_area
        wp[f"E{r}"] = w.title
        wp[f"F{r}"] = w.central_or_local
        wp[f"G{r}"] = w.priority
        wp[f"H{r}"] = (
            w.owner.display_name if w.owner else (w.owner_stakeholder.name if w.owner_stakeholder else None)
        )
        wp[f"I{r}"] = w.planned_start
        wp[f"J{r}"] = w.planned_finish
        wp[f"K{r}"] = w.acceptance_criterion
        wp[f"L{r}"] = w.expected_evidence
        wp[f"M{r}"] = w.dependency
        wp[f"N{r}"] = w.status
        wp[f"O{r}"] = w.progress
        wp[f"Q{r}"] = w.comments
    rd = _SafeSheet(wb["Risks Issues Decisions"])
    for n, x in enumerate(
        db.scalars(select(RaidRecord).where(RaidRecord.project_id == p.id).order_by(RaidRecord.id)).all()[:30]
    ):
        r = 6 + n
        for col, val in zip(
            "ABCDEFGHIJKLM",
            [
                x.code,
                x.type,
                x.date_raised,
                x.title,
                x.impact,
                x.probability,
                x.severity,
                x.owner.name if x.owner else None,
                x.mitigation,
                x.target_date,
                x.status,
                "Yes" if x.escalation_needed else "No",
                x.resolution,
            ],
            strict=True,
        ):
            rd[f"{col}{r}"] = val
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_gap(db: Session, data: bytes, round_id: int, pr, dry_run: bool = True) -> dict:
    """Import answers and action plan rows from a filled GAP workbook into the given round."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        parsed = parse_gap_workbook(io.BytesIO(data), read_only=True)
    stakeholders = {s.name.lower(): s for s in db.scalars(select(Stakeholder))}
    reqs = {r.number: r for r in db.scalars(select(Requirement))}
    ents = {e.code: e for e in db.scalars(select(Entity))}
    user_id = pr.id
    summary = {
        "entities": {},
        "action_items_new": 0,
        "unknown_entities": [],
        "unknown_responsible": set(),
        "skipped_locked": [],
        "skipped_no_permission": [],
    }
    fields = (
        "scope",
        "implemented",
        "documented",
        "effectiveness",
        "evidence",
        "gap_description",
        "recommended_action",
        "target_date",
        "effort",
    )
    valid = {
        "scope": {"In Scope", "Central at HQ", "Not Applicable"},
        "implemented": {"Yes", "Partial", "No", "N/A"},
        "documented": {"Yes", "Partial", "No", "N/A"},
        "effectiveness": {"Effective", "Partial", "Not Tested", "Ineffective", "N/A"},
        "effort": {"Low", "Medium", "High"},
    }
    for code, rows in parsed["answers"].items():
        ent = ents.get(code)
        if not ent:
            summary["unknown_entities"].append(code)
            continue
        if not pr.can("assessments.edit", ent.id):
            summary["skipped_no_permission"].append(code)
            continue
        asm = db.scalar(
            select(Assessment).where(Assessment.round_id == round_id, Assessment.entity_id == ent.id)
        )
        if not asm:
            continue
        if asm.locked and not pr.can("assessments.manage", ent.id):
            summary["skipped_locked"].append(code)
            continue
        answers = {
            a.requirement.number: a for a in db.scalars(select(Answer).where(Answer.assessment_id == asm.id))
        }
        changed = 0
        for row in rows:
            a = answers.get(row["number"])
            if not a:
                continue
            upd = {}
            for f in fields:
                v = row.get(f)
                if f in valid and v is not None and v not in valid[f]:
                    continue
                if v != getattr(a, f) and not (f == "scope" and v is None):
                    upd[f] = v
            resp = row.get("responsible")
            if resp:
                st = stakeholders.get(resp.lower())
                if st and st.id != a.responsible_id:
                    upd["responsible_id"] = st.id
                elif not st:
                    summary["unknown_responsible"].add(resp)
            if upd:
                changed += 1
                if not dry_run:
                    from app.models import AnswerRevision

                    diff = {
                        k: [
                            str(getattr(a, k)) if getattr(a, k) is not None else None,
                            str(v) if v is not None else None,
                        ]
                        for k, v in upd.items()
                    }
                    for k, v in upd.items():
                        setattr(a, k, v)
                    recalc(a, ent.finding_prefix)
                    a.updated_by_id = user_id
                    db.add(
                        AnswerRevision(
                            answer_id=a.id, changed_by_id=user_id, changes={"import": True, **diff}
                        )
                    )
        summary["entities"][code] = changed
    existing = {(i.entity_id, i.requirement_id) for i in db.scalars(select(ActionItem))}
    from app.api.util import next_code

    for ap in parsed["action_plan"]:
        ent, req = ents.get(ap["entity"] or ""), reqs.get(ap["requirement"])
        if not ent or not req or (ent.id, req.id) in existing:
            continue
        if not pr.can("findings.manage", ent.id):
            continue
        if ap["priority"] not in ("High", "Medium", "Low") or ap["status"] not in (
            "Not Started",
            "In Progress",
            "Completed",
            "Blocked",
            "Deferred",
        ):
            continue
        summary["action_items_new"] += 1
        existing.add((ent.id, req.id))
        if not dry_run:
            st = stakeholders.get((ap["responsible"] or "").lower())
            db.add(
                ActionItem(
                    code=next_code(db, ActionItem.code, "AP-"),
                    round_id=round_id,
                    entity_id=ent.id,
                    requirement_id=req.id,
                    title=ap["title"],
                    priority=ap["priority"],
                    acceptance_criterion=ap["acceptance_criterion"],
                    expected_evidence=ap["expected_evidence"],
                    responsible_id=st.id if st else None,
                    target_date=ap["target_date"],
                    status=ap["status"],
                    comments=ap["comments"],
                )
            )
            db.flush()
    if not dry_run:
        db.commit()
    summary["unknown_responsible"] = sorted(summary["unknown_responsible"])
    summary["dry_run"] = dry_run
    return summary


def export_filename(prefix: str) -> str:
    return f"{prefix}_{date.today().isoformat()}.xlsx"
