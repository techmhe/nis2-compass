"""Parsers for the two Excel templates (GAP Analysis questionnaire and Implementation Project Plan).

Used for the initial seed and for importing filled-in workbooks. Only literal cell values are read; formula
columns (score, WP required, finding ID, due status ...) are recalculated by the application.
"""

from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "seed" / "templates"
GAP_TEMPLATE = TEMPLATE_DIR / "NIS2_GAP_Analysis_Questionnaire_v2.xlsx"
PROJECT_TEMPLATE = TEMPLATE_DIR / "NIS2_Implementation_Project_Plan_v2.xlsx"
MAX_ENTITIES = 50  # Lists rows 6 to 55


class _Cell:
    __slots__ = ("value",)

    def __init__(self, value: Any = None):
        self.value = value


class _Grid:
    """A sheet read once into memory. Read-only worksheets only stream, so random cell access would rescan.

    Only non-empty rows are kept, and only columns A to T are read (the parsers use nothing further right). The
    declared sheet dimension is not trusted: an upload can declare a million rows or a one-cell sheet.
    """

    MAX_COL = 20

    def __init__(self, ws):
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
        self._rows: dict[int, tuple] = {}
        for r, row in enumerate(ws.iter_rows(max_col=self.MAX_COL, values_only=True), start=1):
            if any(v is not None for v in row):
                self._rows[r] = row
        self.max_row = max(self._rows, default=0)

    def __getitem__(self, ref: str) -> _Cell:
        letters = ref.rstrip("0123456789")
        row, col = int(ref[len(letters) :]), column_index_from_string(letters)
        cells = self._rows.get(row, ())
        return _Cell(cells[col - 1] if col <= len(cells) else None)


class _Book:
    def __init__(self, wb):
        self.sheetnames = wb.sheetnames
        self._wb = wb
        self._grids: dict[str, _Grid] = {}

    def __getitem__(self, name: str) -> _Grid:
        if name not in self._grids:
            self._grids[name] = _Grid(self._wb[name])
        return self._grids[name]


def _v(ws, cell: str) -> Any:
    val = ws[cell].value
    if isinstance(val, str):
        if val.startswith("="):
            return None
        val = val.strip()
        return val or None
    return val


def _d(val: Any) -> date | None:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        try:
            return date.fromisoformat(val[:10])
        except ValueError:
            return None
    return None


def _col(ws, col: str, start: int, end: int | None = None) -> list[Any]:
    out = []
    r = start
    while True:
        if end is not None and r > end:
            break
        v = _v(ws, f"{col}{r}")
        if v is None and end is None:
            break
        out.append(v)
        r += 1
    return out


def parse_gap_workbook(path: Path | str = GAP_TEMPLATE, read_only: bool = False) -> dict:
    """read_only streams the sheets instead of building the whole object model; use it for user uploads."""
    wb = _Book(openpyxl.load_workbook(path, data_only=False, read_only=read_only))
    lists = wb["Lists"]
    entities = []  # only exports list entities; the template ships none
    for r in range(6, 6 + MAX_ENTITIES):
        code = _v(lists, f"F{r}")
        if not code:
            break
        entities.append(
            {"name": code, "sheet": _v(lists, f"O{r}") or code, "prefix": _v(lists, f"P{r}") or code}
        )
    responsible = [x for x in _col(lists, "S", 6, 40) if x and not x.lower().startswith("replace these")]
    chapters = []
    dash = wb["Dashboard"]
    # the chapter table moves down with the number of entities in an export, so look for the labels
    targets = {}
    for r in range(1, dash.max_row + 1):
        label = _v(dash, f"A{r}")
        if label:
            targets.setdefault(label, dash[f"B{r}"].value)
    for label in _col(lists, "M", 6):
        num, _, title = label.partition(" ")
        t = targets.get(label)
        chapters.append(
            {"number": int(num), "title": title, "target": float(t) if isinstance(t, (int, float)) else 0.75}
        )

    first = wb["Requirements"]
    requirements = []
    for r in range(6, first.max_row + 1):
        num = _v(first, f"A{r}")
        if not num:
            continue
        requirements.append(
            {
                "number": str(num),
                "chapter": int(str(_v(first, f"B{r}")).split(" ")[0]),
                "section": _v(first, f"C{r}"),
                "text": _v(first, f"D{r}"),
                "check_points": _v(first, f"E{r}"),
                "legal_reference": _v(first, f"F{r}"),
                "iso_reference": _v(first, f"G{r}"),
                "is_core": (_v(first, f"H{r}") or "").lower() == "yes",
                "sort_order": r,
            }
        )

    answers: dict[str, list[dict]] = {}
    for ent in entities:
        if ent["sheet"] not in wb.sheetnames:
            continue
        ws = wb[ent["sheet"]]
        rows = []
        for r in range(6, ws.max_row + 1):
            num = _v(ws, f"A{r}")
            if not num:
                continue
            row = {
                "number": str(num),
                "scope": _v(ws, f"I{r}") or "In Scope",
                "implemented": _v(ws, f"J{r}"),
                "documented": _v(ws, f"K{r}"),
                "effectiveness": _v(ws, f"L{r}"),
                "evidence": _v(ws, f"M{r}"),
                "gap_description": _v(ws, f"P{r}"),
                "recommended_action": _v(ws, f"Q{r}"),
                "responsible": _v(ws, f"R{r}"),
                "target_date": _d(ws[f"S{r}"].value),
                "effort": _v(ws, f"T{r}"),
            }
            rows.append(row)
        answers[ent["name"]] = rows

    action_plan = []
    ap = wb["Action Plan"]
    for r in range(6, ap.max_row + 1):
        title = _v(ap, f"H{r}")
        if not title:
            continue
        action_plan.append(
            {
                "entity": _v(ap, f"B{r}"),
                "requirement": str(_v(ap, f"C{r}") or ""),
                "title": title,
                "priority": _v(ap, f"I{r}") or "Medium",
                "acceptance_criterion": _v(ap, f"J{r}"),
                "expected_evidence": _v(ap, f"K{r}"),
                "responsible": _v(ap, f"L{r}"),
                "target_date": _d(ap[f"M{r}"].value),
                "status": _v(ap, f"N{r}") or "Not Started",
                "comments": _v(ap, f"P{r}"),
            }
        )
    return {
        "entities": entities,
        "responsible": responsible,
        "chapters": chapters,
        "requirements": requirements,
        "answers": answers,
        "action_plan": action_plan,
    }


def parse_project_workbook(path: Path | str = PROJECT_TEMPLATE) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    ov = wb["Project Overview"]
    overview = []
    for r in range(6, ov.max_row + 1):
        f, v = _v(ov, f"A{r}"), ov[f"B{r}"].value
        if f:
            overview.append({"field": f, "value": v.date().isoformat() if isinstance(v, datetime) else v})
    lists = wb["Lists"]
    phases = [p for p in _col(lists, "C", 6, 30) if p]
    owners = [o for o in _col(lists, "F", 6, 40) if o]
    measure_areas = [m for m in _col(lists, "L", 6, 40) if m]
    project_entities = [e for e in _col(lists, "E", 6, 30) if e]

    gg = wb["Governance and Gates"]
    gates = []
    for r in range(6, 12):
        code = _v(gg, f"A{r}")
        if code:
            gates.append(
                {
                    "code": code,
                    "planned_date": _d(gg[f"B{r}"].value),
                    "decision": _v(gg, f"C{r}"),
                    "decision_owner": _v(gg, f"D{r}"),
                    "required_input": _v(gg, f"E{r}"),
                    "approval_criterion": _v(gg, f"F{r}"),
                    "status": _v(gg, f"G{r}") or "Not Started",
                    "decision_reference": _v(gg, f"H{r}"),
                }
            )
    raci_cols = []
    for c in "BCDEFGHIJKL":
        name = _v(gg, f"{c}15")
        if name:
            raci_cols.append((c, name))
    raci_rows = []
    for r in range(16, 31):
        act = _v(gg, f"A{r}")
        if not act or act == "Role Notes":
            break
        raci_rows.append({"activity": act, "cells": {name: _v(gg, f"{c}{r}") for c, name in raci_cols}})
    role_notes = []
    for r in range(34, gg.max_row + 1):
        role, note = _v(gg, f"A{r}"), _v(gg, f"B{r}")
        if role:
            role_notes.append({"role": role, "note": note})

    mp = wb["Master Plan"]
    activities = []
    for r in range(6, mp.max_row + 1):
        code = _v(mp, f"A{r}")
        if not code:
            continue
        progress = mp[f"O{r}"].value
        activities.append(
            {
                "code": code,
                "phase": _v(mp, f"B{r}"),
                "title": _v(mp, f"C{r}"),
                "owner": _v(mp, f"D{r}"),
                "support": _v(mp, f"E{r}"),
                "planned_start": _d(mp[f"F{r}"].value),
                "planned_finish": _d(mp[f"G{r}"].value),
                "dependency": _v(mp, f"I{r}"),
                "deliverable": _v(mp, f"J{r}"),
                "gate": _v(mp, f"K{r}"),
                "status": _v(mp, f"M{r}") or "Not Started",
                "priority": _v(mp, f"N{r}") or "Medium",
                "progress": float(progress) if isinstance(progress, (int, float)) else 0.0,
            }
        )

    wp = wb["Work Packages"]
    work_packages = []
    for r in range(6, wp.max_row + 1):
        title = _v(wp, f"E{r}")
        if not title:
            continue
        progress = wp[f"O{r}"].value
        work_packages.append(
            {
                "finding_ids": _v(wp, f"B{r}"),
                "entity": _v(wp, f"C{r}"),
                "measure_area": _v(wp, f"D{r}"),
                "title": title,
                "central_or_local": _v(wp, f"F{r}") or "Central",
                "priority": _v(wp, f"G{r}") or "Medium",
                "owner": _v(wp, f"H{r}"),
                "planned_start": _d(wp[f"I{r}"].value),
                "planned_finish": _d(wp[f"J{r}"].value),
                "acceptance_criterion": _v(wp, f"K{r}"),
                "expected_evidence": _v(wp, f"L{r}"),
                "dependency": _v(wp, f"M{r}"),
                "status": _v(wp, f"N{r}") or "Not Started",
                "progress": float(progress) if isinstance(progress, (int, float)) else 0.0,
                "comments": _v(wp, f"Q{r}"),
            }
        )

    rd = wb["Risks Issues Decisions"]
    raid = []
    for r in range(6, rd.max_row + 1):
        code = _v(rd, f"A{r}")
        if not code:
            continue
        raid.append(
            {
                "code": code,
                "type": _v(rd, f"B{r}") or "Risk",
                "date_raised": _d(rd[f"C{r}"].value),
                "title": _v(rd, f"D{r}") or "",
                "impact": _v(rd, f"E{r}"),
                "probability": _v(rd, f"F{r}"),
                "severity": _v(rd, f"G{r}"),
                "owner": _v(rd, f"H{r}"),
                "mitigation": _v(rd, f"I{r}"),
                "target_date": _d(rd[f"J{r}"].value),
                "status": _v(rd, f"K{r}") or "Open",
                "escalation_needed": (_v(rd, f"L{r}") or "").lower() == "yes",
                "resolution": _v(rd, f"M{r}"),
            }
        )
    return {
        "overview": overview,
        "phases": phases,
        "owners": owners,
        "measure_areas": measure_areas,
        "entities": project_entities,
        "gates": gates,
        "raci_columns": [n for _, n in raci_cols],
        "raci_rows": raci_rows,
        "role_notes": role_notes,
        "activities": activities,
        "work_packages": work_packages,
        "raid": raid,
    }
