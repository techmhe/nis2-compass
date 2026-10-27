"""One-off: remove the original group's structure from the two Excel templates.

Kept for reproducibility. It expects the per-entity workbooks as they were before the first-run setup
(`git show 1bee5ca:templates/<file>`); a workbook that is already neutral is left alone.

    backend/.venv/bin/python tools/templates/neutralize.py

GAP workbook: the first entity sheet becomes `Requirements` (answer columns cleared), the other entity sheets
and the entity list on `Lists` go, the action plan ships empty and its lookups cover up to 50 entities. The
Dashboard keeps its four-slot layout as the style source; `export_gap` rebuilds it for the real entities.
Project workbook: per-entity activities become one "... for every entity" activity each, entity names leave
the overview, lists, RACI matrix and role notes.
Both: author metadata is cleared. The result is copied to `templates/` and `backend/app/seed/templates/`.
"""

import shutil
import warnings
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment

REPO = Path(__file__).resolve().parents[2]
TARGETS = (REPO / "templates", REPO / "backend/app/seed/templates")
GAP = "NIS2_GAP_Analysis_Questionnaire_v2.xlsx"
PROJECT = "NIS2_Implementation_Project_Plan_v2.xlsx"
MAX_ENTITIES = 50
LAST_ENTITY_ROW = 5 + MAX_ENTITIES  # Lists rows 6 .. 55


def _load(path: Path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path)


def _clear_properties(wb) -> None:
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""


def _set(ws, coord: str, expected_start: str, new: str) -> None:
    cell = ws[coord]
    assert isinstance(cell.value, str) and cell.value.startswith(expected_start), (
        coord,
        cell.value,
    )
    cell.value = new


def _clear_chart_caches(wb) -> None:
    """Series titles, categories and values carry cached copies of the cells they point at (entity names)."""
    for ws in wb.worksheets:
        for chart in ws._charts:
            for series in chart.series:
                for ref in (series.tx and series.tx.strRef, series.cat and series.cat.strRef):
                    if ref is not None:
                        ref.strCache = None
                if series.cat is not None and series.cat.numRef is not None:
                    series.cat.numRef.numCache = None
                if series.val is not None and series.val.numRef is not None:
                    series.val.numRef.numCache = None


def neutralize_gap(path: Path) -> None:
    wb = _load(path)
    if "Requirements" in wb.sheetnames:
        return
    lists = wb["Lists"]
    # the per-entity sheets are named in Lists!O6:O9; the first one becomes the catalogue sheet
    old_sheets = [lists[f"O{r}"].value for r in range(6, 10)]
    ws = wb[old_sheets[0]]
    ws.title = "Requirements"
    ws["A2"].value = "Requirement catalogue"
    for r in range(6, ws.max_row + 1):
        for col in "JKLMPQRST":
            ws[f"{col}{r}"].value = None
    for name in old_sheets[1:]:
        del wb[name]

    note = lists["O11"].value
    lists.unmerge_cells("O11:Q14")
    blank = copy(lists["O20"]._style)
    for r in range(6, 15):
        for col in "FOPQ":
            lists[f"{col}{r}"].value = None
    for r in range(11, 15):
        for col in "OPQ":
            lists[f"{col}{r}"]._style = copy(blank)
    # the note moved out of the way of the entity rows, which may now run down to row 55
    lists["U6"].value = note
    lists["U6"].alignment = Alignment(wrap_text=True, vertical="top")
    lists.merge_cells("U6:Z14")
    lists.column_dimensions["U"].width = 14
    wb.defined_names["EntityList"].attr_text = f"Lists!$F$6:$F${LAST_ENTITY_ROW}"

    _clear_chart_caches(wb)

    ap = wb["Action Plan"]
    for r in range(6, ap.max_row + 1):
        for col in "BCHIJKLMNP":
            ap[f"{col}{r}"].value = None
        for col in "DEFGQ":
            v = ap[f"{col}{r}"].value
            if isinstance(v, str):
                ap[f"{col}{r}"].value = v.replace(
                    "$O$6:$O$9", f"$O$6:$O${LAST_ENTITY_ROW}"
                ).replace("$F$6:$F$9", f"$F$6:$F${LAST_ENTITY_ROW}")

    ins = wb["Instructions"]
    _set(
        ins,
        "B7",
        "Identify gaps",
        (
            "Identify gaps against the technical and methodological requirements of Article 21(2) NIS2 and "
            "convert them into manageable work packages across the in-scope legal entities of the group."
        ),
    )
    _set(
        ins,
        "B12",
        "IT operations",
        (
            'IT operations of the subsidiaries are often managed centrally by HQ. Set Scope to "Central at HQ" '
            "on the entity sheets where a control is operated by HQ; the row is then excluded from that "
            "entity's coverage and score."
        ),
    )
    _set(ins, "A19", old_sheets[0], "Requirements / entity sheets")
    _set(
        ins,
        "B19",
        "One assessment",
        (
            "Requirements is the catalogue of all 174 requirements. An export adds one assessment sheet per "
            "legal entity, named by its entity code, with the same 174 requirements so the entities stay "
            "comparable. The tab codes are short keys; the full names are kept in the application and shown in "
            "column F of the Lists sheet."
        ),
    )
    _set(
        ins,
        "B42",
        "Generated for every",
        (
            "Generated for every requirement that needs a work package, using the entity's finding prefix "
            "and the requirement number."
        ),
    )
    _set(
        ins,
        "B60",
        "Maintained in one place",
        (
            "Maintained in the application and written to the Lists sheet by the export, one row per entity "
            "starting at row 6. Column F is the display name shown on the Dashboard, in the title of each "
            "assessment sheet and in the Action Plan drop-down. Column P is the Finding ID prefix. Column O "
            "holds the tab name and must match the actual sheet tab; if you rename a tab, update column O and "
            "the check in column Q switches back to OK."
        ),
    )
    _clear_properties(wb)
    wb.active = 0
    wb.save(path)


def _merge_activities(mp) -> None:
    """Replace the per-entity rows P2.2-P2.5 and P3.1-P3.8 with one 'for every entity' activity each."""
    cols = "ABCDEFGIJKMNO"  # H, L, P are formulas
    rows = {
        r: {c: mp[f"{c}{r}"].value for c in cols}
        for r in range(6, mp.max_row + 1)
        if mp[f"A{r}"].value
    }
    old = {v["A"]: v for v in rows.values()}
    acts = [
        v
        for v in rows.values()
        if v["A"] not in {"P2.3", "P2.4", "P2.5"} | {f"P3.{i}" for i in range(3, 9)}
    ]
    by_code = {a["A"]: a for a in acts}
    by_code["P2.2"].update(
        C="Issue evidence and interview request for every entity",
        E="Central IT; Local IT; MSP+IT; OT; HR; Procurement; DPO; Advisers; Local Contact",
        G=old["P2.5"]["G"],
    )
    by_code["P3.1"].update(
        C="Conduct remote assessment for every entity (174 requirements)",
        E="Central IT; Local IT; MSP+IT; OT; HR; Procurement; DPO; Advisers; Local Contact",
        G=old["P3.7"]["G"],
        I="P2.2",
    )
    by_code["P3.2"].update(
        C="Validate findings for every entity",
        E="Central IT; Local IT; MSP+IT; Advisers",
        F=old["P3.2"]["F"],
        G=old["P3.8"]["G"],
        I="P3.1",
    )
    by_code["P4.1"]["I"] = "P3.2"
    formulas = {c: mp[f"{c}6"].value for c in "HLP"}
    for i, r in enumerate(range(6, 41)):
        act = acts[i] if i < len(acts) else dict.fromkeys(cols)
        for c in cols:
            mp[f"{c}{r}"].value = act[c]
        for c, f in formulas.items():
            mp[f"{c}{r}"].value = Translator(f, origin=f"{c}6").translate_formula(
                f"{c}{r}"
            )


def neutralize_project(path: Path) -> None:
    wb = _load(path)
    ov = wb["Project Overview"]
    if not any(c.value == "P2.5" for c in wb["Master Plan"]["A"]):
        return  # the per-entity activities are already merged
    ov["B12"].value = "HQ and all subsidiaries"
    _merge_activities(wb["Master Plan"])

    lists = wb["Lists"]
    for r in range(7, 11):  # keep "Group / All Entities" in E6
        lists[f"E{r}"].value = None

    gg = wb["Governance and Gates"]
    assert str(gg["A20"].value).startswith("Finding validation ") and str(gg["A21"].value).startswith(
        "Finding validation "
    )
    for r in range(21, 30):  # one validation row replaces the German and the Dutch one
        for c in "ABCDEFGHIJKL":
            gg[f"{c}{r}"].value = gg[f"{c}{r + 1}"].value
    gg["A20"].value = "Finding validation"
    _set(
        gg,
        "B40",
        "Provides controls and evidence for HQ",
        (
            "Provides controls and evidence for HQ and for the subsidiaries that rely on central IT"
        ),
    )
    _set(
        gg,
        "B41",
        "Provides controls, evidence",
        (
            "Provides controls, evidence and remediation for subsidiaries without internal IT"
        ),
    )
    ins = wb["Instructions"]
    _set(
        ins,
        "B52",
        "The GAP questionnaire",
        ins["B52"].value.replace("for each of the four entities", "for each entity"),
    )
    _clear_properties(wb)
    wb.active = 0
    wb.save(path)


def main() -> None:
    source = TARGETS[0]
    neutralize_gap(source / GAP)
    neutralize_project(source / PROJECT)
    for other in TARGETS[1:]:
        for name in (GAP, PROJECT):
            shutil.copyfile(source / name, other / name)


if __name__ == "__main__":
    main()
